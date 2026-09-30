#!/usr/bin/env python3
"""Build the detailed legacy-QC iNKT reproduction deck.

The script creates:
  * a 38-page, 16:9 PDF with vector text;
  * a PowerPoint file containing the same rendered pages;
  * one high-resolution PNG per page;
  * a manifest with source paths, titles, and SHA-256 hashes.

The deck deliberately distinguishes exact count-level QC reconstruction from
the downstream Scanpy reanalysis.  A complete audited analysis run supplies
DE/pathway results, while an explicit embedding run and PC sweep supply every
UMAP used in the rebuilt deck.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import textwrap
from io import BytesIO
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable, Sequence
from zipfile import ZipFile

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[3]
DEFAULT_ANALYSIS_RUN = ROOT / "output/iNKT_legacy_ppt_qc_runs/20260818_125231"
DEFAULT_EMBEDDING_RUN = ROOT / "output/iNKT_legacy_ppt_qc_runs/20260825_095422"
DEFAULT_OUT_DIR = ROOT / "output/iNKT_reproduction_deck"

RUN = DEFAULT_ANALYSIS_RUN
EMBEDDING_RUN = DEFAULT_EMBEDDING_RUN
PRE = EMBEDDING_RUN / "preprocess"
EXT = RUN / "extended"
COMPARE = RUN / "comparison_to_legacy_ppt"
SWEEP = DEFAULT_OUT_DIR / "umap_pc_sweep"
SELECTED_FIGURES = SWEEP / "selected_figures"
STD = SWEEP / "selected_signatures/standardized_legacy"
DOC = ROOT / "docs/audits/inkt_legacy_ppt_vs_legacy_qc_rerun_content.md"
LEGACY_PPT = ROOT / "input/iNKT/iNKT.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5
PNG_DPI = 180

# Theme: a cleaner continuation of the blue/orange Office palette used by the
# legacy deck.  Tissue and condition colors remain semantically stable.
NAVY = "#102A43"
NAVY_2 = "#243B53"
BLUE = "#4472C4"
BLUE_2 = "#2F80ED"
TEAL = "#1F8A8A"
TEAL_LIGHT = "#DDF3F2"
ORANGE = "#ED7D31"
GOLD = "#C99520"
PURPLE = "#7A5AF8"
RED = "#C94C4C"
GREEN = "#2E8B57"
INK = "#172B3A"
MID = "#52616B"
MUTED = "#7B8794"
LIGHT = "#E6ECF2"
PALE = "#F4F7FA"
WHITE = "#FFFFFF"
BLACK = "#111827"

TISSUE = {
    "Bone marrow": "#D9B300",
    "Spleen": "#D62728",
    "Thymus": "#2455FF",
}

DEFAULT_FOOTER = (
    "Legacy QC reconstructed at the cell/gene-count level; downstream analysis is a current reanalysis."
)
DE_FOOTER = (
    "Exploratory cell-level Wilcoxon comparison; FDR <= 0.05 and |logFC| >= 0.25 unless noted; "
    "one biological sample per tissue x condition."
)


@dataclass
class DeckContext:
    out_dir: Path
    slides_dir: Path
    pdf_path: Path
    pptx_path: Path
    manifest_path: Path
    qc: dict
    preprocess_summary: dict
    best_matches: list[dict[str, str]]
    driver_retention: list[dict[str, str]]
    standardized_calls: list[dict[str, str]]
    sweep_summary: dict
    sweep_metrics: list[dict[str, str]]


def configure_sources(analysis_run: Path, embedding_run: Path, pc_sweep_dir: Path) -> None:
    global RUN, EMBEDDING_RUN, PRE, EXT, COMPARE, SWEEP, SELECTED_FIGURES, STD
    RUN = analysis_run.resolve()
    EMBEDDING_RUN = embedding_run.resolve()
    PRE = EMBEDDING_RUN / "preprocess"
    EXT = RUN / "extended"
    COMPARE = RUN / "comparison_to_legacy_ppt"
    SWEEP = pc_sweep_dir.resolve()
    SELECTED_FIGURES = SWEEP / "selected_figures"
    STD = SWEEP / "selected_signatures/standardized_legacy"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def read_json(path: Path) -> dict:
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)


def manifest_path(path: Path) -> str:
    resolved = path.resolve()
    try:
        return str(resolved.relative_to(ROOT))
    except ValueError:
        return str(resolved)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_inputs() -> None:
    required = [
        DOC,
        LEGACY_PPT,
        RUN / "pipeline_status.txt",
        EMBEDDING_RUN / "qc_validation.json",
        PRE / "summary.json",
        COMPARE / "legacy_ppt_de_best_matches.csv",
        COMPARE / "legacy_ppt_pathway_driver_retention.csv",
        PRE / "figures/qc_violin_by_sample.png",
        PRE / "figures/qc_scatter_counts_genes_sample.png",
        PRE / "figures/highly_variable_genes.png",
        PRE / "figures/pca_variance_ratio.png",
        PRE / "figures/dotplot_inkt_markers_by_cluster.png",
        PRE / "figures/cluster_percent_by_sample.png",
        SWEEP / "summary.json",
        SWEEP / "tables/umap_pc_sweep_metrics.csv",
        SWEEP / "figures/umap_pc_sweep_tissue_cluster.png",
        SWEEP / "figures/umap_pc_sweep_metrics.png",
        SWEEP / "figures/legacy_reference_200pc.png",
        SELECTED_FIGURES / "umap_selected_sample_condition_tissue_cluster.png",
        SELECTED_FIGURES / "umap_selected_condition_cluster_top_markers.png",
        SELECTED_FIGURES / "trajectory_selected_paga_cluster_graph.png",
        SELECTED_FIGURES / "trajectory_selected_umap_cluster_dpt.png",
        SELECTED_FIGURES / "trajectory_selected_umap_dpt_pseudotime.png",
        SELECTED_FIGURES / "trajectory_selected_paga_compare_umap_pseudotime.png",
        EXT / "signatures/signature_scores/cluster_subtype_evidence.png",
        STD / "standardized_cluster_calls.csv",
        STD / "standardized_cluster_heatmap.png",
        STD / "standardized_inkt1_umap_violin.png",
        STD / "standardized_inkt2_umap_violin.png",
        STD / "standardized_inkt17_umap_violin.png",
        STD / "summary.json",
        EXT / "signatures/gene_sets/signature_jaccard_heatmap.png",
        EXT / "signatures/marker_validation/marker_expression_condition_tissue_dotplot.png",
        EXT / "signatures/marker_validation/marker_t2_ctrl_effect_heatmap.png",
        EXT / "de_pathway/overlap_pathway/current_tissue_vs_cluster_overlap_jaccard.png",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing required inputs:\n" + "\n".join(missing))

    status = (RUN / "pipeline_status.txt").read_text(encoding="utf-8")
    if "overall=0" not in status:
        raise RuntimeError("The audited pipeline did not finish with overall=0")

    qc = read_json(EMBEDDING_RUN / "qc_validation.json")
    assert qc["status"] == "passed"
    assert qc["shape"] == [15532, 10670]
    assert qc["highly_variable_genes"] == 3000
    assert qc["n_clusters_leiden_res_0_5"] == 8
    assert qc["sample_counts"] == {
        "Ctrl_BM": 3379,
        "Ctrl_Spleen": 3422,
        "Ctrl_Thymus": 1372,
        "T2_BM": 2731,
        "T2_Spleen": 3371,
        "T2_Thymus": 1257,
    }

    sweep = read_json(SWEEP / "summary.json")
    assert sweep["candidate_pcs"] == [50, 100, 200]
    assert sweep["selected_pc"] in sweep["candidate_pcs"]
    assert sweep["parameters"] == {
        "metric": "euclidean",
        "min_dist": 0.5,
        "n_neighbors": 20,
        "pca_mask": "3,000 highly_variable genes",
        "pca_solver": "arpack",
        "random_state": 0,
        "spread": 1.0,
    }
    selected_pc = int(sweep["selected_pc"])
    selected_required = [
        SWEEP / f"figures/umap_npcs_{selected_pc:03d}.png",
        SWEEP / "inkt_selected_umap.h5ad",
    ]
    selected_missing = [str(path) for path in selected_required if not path.exists()]
    if selected_missing:
        raise FileNotFoundError("Missing selected-PC inputs:\n" + "\n".join(selected_missing))
    assert "fixed audited" in sweep["cluster_boundary_policy"].lower()
    assert "no validated discrete" in sweep["subtype_policy"].lower()
    assert sweep["trajectory"]["status"] == "recomputed_from_selected_graph"
    assert sweep["trajectory"]["root_cluster"] == "7"
    assert sweep["artifacts"]["selected_h5ad_sha256"] == sha256(
        SWEEP / "inkt_selected_umap.h5ad"
    )
    sweep_metrics = read_csv(SWEEP / "tables/umap_pc_sweep_metrics.csv")
    assert [int(row["n_pcs"]) for row in sweep_metrics] == [50, 100, 200]
    for row in sweep_metrics:
        for key, value in row.items():
            if key != "n_pcs" and not math.isfinite(float(value)):
                raise ValueError(f"Non-finite sweep metric: n_pcs={row['n_pcs']} {key}={value}")

    standardized = read_json(STD / "summary.json")
    assert standardized["task"] == "standardized-legacy"
    assert standardized["bootstrap_replicates"] == 10000
    calls = {
        row["cluster"]: row["supported_call"]
        for row in read_csv(STD / "standardized_cluster_calls.csv")
    }
    assert calls == {
        "0": "legacy-list inkt1-enriched",
        "1": "mixed",
        "2": "unclassified",
        "3": "unclassified",
        "4": "legacy-list inkt1-enriched",
        "5": "legacy-list inkt17-enriched",
        "6": "legacy-list inkt1-enriched",
        "7": "unclassified",
    }


def make_context(out_dir: Path) -> DeckContext:
    slides_dir = out_dir / "slides"
    slides_dir.mkdir(parents=True, exist_ok=True)
    return DeckContext(
        out_dir=out_dir,
        slides_dir=slides_dir,
        pdf_path=out_dir / "iNKT_legacy_QC_reproduction_detailed.pdf",
        pptx_path=out_dir / "iNKT_legacy_QC_reproduction_detailed.pptx",
        manifest_path=out_dir / "deck_manifest.json",
        qc=read_json(EMBEDDING_RUN / "qc_validation.json"),
        preprocess_summary=read_json(PRE / "summary.json"),
        best_matches=read_csv(COMPARE / "legacy_ppt_de_best_matches.csv"),
        driver_retention=read_csv(COMPARE / "legacy_ppt_pathway_driver_retention.csv"),
        standardized_calls=read_csv(STD / "standardized_cluster_calls.csv"),
        sweep_summary=read_json(SWEEP / "summary.json"),
        sweep_metrics=read_csv(SWEEP / "tables/umap_pc_sweep_metrics.csv"),
    )


def canvas(background: str = WHITE):
    fig = plt.figure(figsize=(SLIDE_W, SLIDE_H), facecolor=background)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, SLIDE_W)
    ax.set_ylim(SLIDE_H, 0)
    ax.axis("off")
    ax.set_facecolor(background)
    return fig, ax


def rounded_rect(
    ax,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    face: str = WHITE,
    edge: str = LIGHT,
    lw: float = 0.8,
    radius: float = 0.08,
    alpha: float = 1.0,
    zorder: int = 1,
):
    patch = FancyBboxPatch(
        (x, y),
        w,
        h,
        boxstyle=f"round,pad=0.015,rounding_size={radius}",
        linewidth=lw,
        edgecolor=edge,
        facecolor=face,
        alpha=alpha,
        zorder=zorder,
    )
    ax.add_patch(patch)
    return patch


def add_text(
    ax,
    x: float,
    y: float,
    text: str,
    *,
    size: float = 12,
    color: str = INK,
    weight: str = "normal",
    ha: str = "left",
    va: str = "top",
    linespacing: float = 1.18,
    zorder: int = 10,
    family: str = "DejaVu Sans",
    style: str = "normal",
    rotation: float = 0,
):
    return ax.text(
        x,
        y,
        text,
        fontsize=size,
        color=color,
        fontweight=weight,
        ha=ha,
        va=va,
        linespacing=linespacing,
        zorder=zorder,
        family=family,
        style=style,
        rotation=rotation,
    )


def wrap(text: str, width: int) -> str:
    paragraphs = []
    for paragraph in text.split("\n"):
        if not paragraph:
            paragraphs.append("")
        else:
            paragraphs.append(
                textwrap.fill(
                    paragraph,
                    width=width,
                    break_long_words=False,
                    break_on_hyphens=False,
                )
            )
    return "\n".join(paragraphs)


def header(
    ax,
    title: str,
    *,
    section: str,
    slide_no: int,
    subtitle: str | None = None,
    footer: str = DEFAULT_FOOTER,
):
    ax.add_patch(Rectangle((0, 0), SLIDE_W, 0.10, facecolor=TEAL, edgecolor="none"))
    rounded_rect(ax, 0.52, 0.33, 1.52, 0.30, face=TEAL_LIGHT, edge=TEAL_LIGHT, radius=0.15)
    add_text(ax, 1.28, 0.48, section.upper(), size=8.5, color=TEAL, weight="bold", ha="center", va="center")
    # Keep long, descriptive reproduction titles on one line without clipping.
    # Shorter legacy-style titles remain at the full 25 pt scale.
    title_size = min(25.0, max(18.0, 25.0 * 62.0 / max(62, len(title))))
    add_text(ax, 0.52, 0.79, title, size=title_size, color=NAVY, weight="bold", va="center")
    if subtitle:
        add_text(ax, 0.54, 1.14, subtitle, size=10.5, color=MID, va="top")
    ax.plot([0.52, 12.81], [1.38, 1.38], color=LIGHT, linewidth=0.8)
    ax.plot([0.52, 12.81], [7.14, 7.14], color=LIGHT, linewidth=0.7)
    add_text(ax, 0.54, 7.31, footer, size=6.7, color=MUTED, va="center")
    add_text(ax, 12.79, 7.31, f"{slide_no:02d}", size=7.5, color=MUTED, weight="bold", ha="right", va="center")


def section_label(ax, x: float, y: float, text: str, color: str = TEAL):
    add_text(ax, x, y, text.upper(), size=8.3, color=color, weight="bold", va="center")


def metric_card(
    ax,
    x: float,
    y: float,
    w: float,
    h: float,
    value: str,
    label: str,
    *,
    accent: str = TEAL,
    note: str | None = None,
):
    rounded_rect(ax, x, y, w, h, face=WHITE, edge=LIGHT, radius=0.10)
    ax.add_patch(Rectangle((x, y), 0.08, h, facecolor=accent, edgecolor="none", zorder=3))
    add_text(ax, x + 0.25, y + 0.22, value, size=23, color=accent, weight="bold")
    add_text(ax, x + 0.25, y + 0.72, label, size=9.4, color=INK, weight="bold")
    if note:
        add_text(ax, x + 0.25, y + 1.03, wrap(note, max(20, int(w * 14))), size=7.2, color=MUTED)


def bullet_list(
    ax,
    x: float,
    y: float,
    items: Sequence[str],
    *,
    width_chars: int = 48,
    size: float = 10.5,
    color: str = INK,
    bullet_color: str = TEAL,
    line_gap: float = 0.12,
    line_height: float = 0.25,
):
    cursor = y
    for item in items:
        wrapped = wrap(item, width_chars)
        n_lines = wrapped.count("\n") + 1
        ax.scatter([x], [cursor + 0.075], s=16, c=[bullet_color], marker="o", zorder=10)
        add_text(ax, x + 0.18, cursor, wrapped, size=size, color=color)
        cursor += n_lines * line_height + line_gap
    return cursor


def callout(
    ax,
    x: float,
    y: float,
    w: float,
    h: float,
    title: str,
    body: str,
    *,
    accent: str = TEAL,
    face: str = PALE,
    title_size: float = 10,
    body_size: float = 8.5,
    width_chars: int = 55,
):
    rounded_rect(ax, x, y, w, h, face=face, edge=LIGHT, radius=0.10)
    ax.add_patch(Rectangle((x, y), 0.06, h, facecolor=accent, edgecolor="none"))
    add_text(ax, x + 0.22, y + 0.18, title, size=title_size, color=accent, weight="bold")
    add_text(ax, x + 0.22, y + 0.52, wrap(body, width_chars), size=body_size, color=INK)


def simple_table(
    ax,
    x: float,
    y: float,
    w: float,
    h: float,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    *,
    col_widths: Sequence[float] | None = None,
    header_color: str = NAVY,
    header_text: str = WHITE,
    body_size: float = 8.5,
    header_size: float = 8.4,
    align: Sequence[str] | None = None,
    row_colors: Sequence[str] | None = None,
):
    n_cols = len(headers)
    n_rows = len(rows)
    if col_widths is None:
        col_widths = [1 / n_cols] * n_cols
    total = sum(col_widths)
    widths = [w * value / total for value in col_widths]
    if align is None:
        align = ["left"] + ["center"] * (n_cols - 1)
    header_h = min(0.42, h * 0.18)
    row_h = (h - header_h) / max(1, n_rows)
    cursor_x = x
    for idx, (head, cell_w) in enumerate(zip(headers, widths)):
        ax.add_patch(Rectangle((cursor_x, y), cell_w, header_h, facecolor=header_color, edgecolor=WHITE, linewidth=0.8))
        ha = "center" if align[idx] != "left" else "left"
        tx = cursor_x + cell_w / 2 if ha == "center" else cursor_x + 0.10
        add_text(ax, tx, y + header_h / 2, head, size=header_size, color=header_text, weight="bold", ha=ha, va="center")
        cursor_x += cell_w
    for ridx, row in enumerate(rows):
        cursor_x = x
        fill = row_colors[ridx] if row_colors and ridx < len(row_colors) else (WHITE if ridx % 2 == 0 else PALE)
        for cidx, (value, cell_w) in enumerate(zip(row, widths)):
            ax.add_patch(Rectangle((cursor_x, y + header_h + ridx * row_h), cell_w, row_h, facecolor=fill, edgecolor=LIGHT, linewidth=0.55))
            ha = "center" if align[cidx] != "left" else "left"
            tx = cursor_x + cell_w / 2 if ha == "center" else cursor_x + 0.10
            add_text(
                ax,
                tx,
                y + header_h + ridx * row_h + row_h / 2,
                str(value),
                size=body_size,
                color=INK,
                ha=ha,
                va="center",
            )
            cursor_x += cell_w


def trim_white(image: Image.Image, threshold: int = 248, pad: int = 8) -> Image.Image:
    rgb = image.convert("RGB")
    background = Image.new("RGB", rgb.size, (255, 255, 255))
    diff = ImageChops.difference(rgb, background).convert("L")
    diff = diff.point(lambda px: 255 if px < 255 - threshold else 0)
    bbox = diff.getbbox()
    if not bbox:
        return rgb
    left, top, right, bottom = bbox
    left = max(0, left - pad)
    top = max(0, top - pad)
    right = min(rgb.width, right + pad)
    bottom = min(rgb.height, bottom + pad)
    return rgb.crop((left, top, right, bottom))


def add_pil_image(
    ax,
    image: Image.Image,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    trim: bool = True,
    border: bool = True,
    label: str | None = None,
):
    image = image.convert("RGB")
    if trim:
        image = trim_white(image)
    ratio = image.width / image.height
    box_ratio = w / h
    if ratio >= box_ratio:
        draw_w = w
        draw_h = w / ratio
    else:
        draw_h = h
        draw_w = h * ratio
    draw_x = x + (w - draw_w) / 2
    draw_y = y + (h - draw_h) / 2
    if border:
        rounded_rect(ax, x, y, w, h, face=WHITE, edge=LIGHT, radius=0.06, zorder=0)
    ax.imshow(image, extent=(draw_x, draw_x + draw_w, draw_y + draw_h, draw_y), zorder=4)
    if label:
        label_w = min(max(1.15, 0.075 * len(label) + 0.38), max(1.15, w - 0.16))
        rounded_rect(
            ax,
            x + 0.08,
            y + 0.08,
            label_w,
            0.28,
            face=NAVY,
            edge=NAVY,
            radius=0.12,
            zorder=6,
        )
        add_text(
            ax,
            x + 0.08 + label_w / 2,
            y + 0.22,
            label,
            size=7.5,
            color=WHITE,
            weight="bold",
            ha="center",
            va="center",
            zorder=7,
        )


def add_image(
    ax,
    path: Path,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    trim: bool = True,
    border: bool = True,
    label: str | None = None,
):
    add_pil_image(
        ax,
        Image.open(path),
        x,
        y,
        w,
        h,
        trim=trim,
        border=border,
        label=label,
    )


def add_legacy_ppt_media(
    ax,
    member: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    label: str | None = None,
):
    with ZipFile(LEGACY_PPT) as archive:
        image = Image.open(BytesIO(archive.read(f"ppt/media/{member}"))).copy()
    add_pil_image(ax, image, x, y, w, h, trim=True, border=True, label=label)


def progress_bar(
    ax,
    x: float,
    y: float,
    w: float,
    value: float,
    *,
    color: str = TEAL,
    label: str | None = None,
    height: float = 0.15,
):
    rounded_rect(ax, x, y, w, height, face=LIGHT, edge=LIGHT, radius=height / 2)
    rounded_rect(ax, x, y, max(0.01, w * value), height, face=color, edge=color, radius=height / 2)
    if label:
        add_text(ax, x + w + 0.12, y + height / 2, label, size=8.2, color=color, weight="bold", va="center")


def arrow(ax, x1: float, y1: float, x2: float, y2: float, color: str = MUTED, lw: float = 1.4):
    ax.add_patch(
        FancyArrowPatch(
            (x1, y1),
            (x2, y2),
            arrowstyle="-|>",
            mutation_scale=12,
            linewidth=lw,
            color=color,
            zorder=5,
        )
    )


def gene_chip(ax, x: float, y: float, text: str, *, color: str, width: float | None = None):
    if width is None:
        width = 0.34 + 0.083 * len(text)
    rounded_rect(ax, x, y, width, 0.31, face=WHITE, edge=color, lw=1.0, radius=0.13)
    add_text(ax, x + width / 2, y + 0.155, text, size=7.4, color=color, weight="bold", ha="center", va="center")
    return width


def stat_pair(ax, x: float, y: float, left: tuple[str, str], right: tuple[str, str], *, colors=(TEAL, BLUE)):
    metric_card(ax, x, y, 2.72, 1.27, left[0], left[1], accent=colors[0])
    metric_card(ax, x + 2.94, y, 2.72, 1.27, right[0], right[1], accent=colors[1])


def plot_recovery_bars(ax, x: float, y: float, w: float, h: float):
    labels = ["Thymus", "Bone marrow", "Spleen"]
    recovered = [31, 35, 44]
    totals = [34, 39, 50]
    colors = [TISSUE[label] for label in labels]
    rounded_rect(ax, x, y, w, h, face=WHITE, edge=LIGHT, radius=0.08)
    section_label(ax, x + 0.25, y + 0.28, "Legacy core recovered")
    for idx, (label, rec, total, color) in enumerate(zip(labels, recovered, totals, colors)):
        yy = y + 0.72 + idx * 0.83
        add_text(ax, x + 0.27, yy, label, size=9.3, color=INK, weight="bold")
        progress_bar(ax, x + 1.52, yy + 0.03, w - 2.47, rec / total, color=color, label=f"{rec}/{total}")
        add_text(ax, x + 0.27, yy + 0.30, f"{100 * rec / total:.1f}% recovered; all retained signs concordant", size=7.0, color=MUTED)


def title_slide(ctx: DeckContext, slide_no: int):
    fig, ax = canvas(NAVY)
    ax.add_patch(Rectangle((0, 0), 0.13, SLIDE_H, facecolor=TEAL, edgecolor="none"))
    for x, y, r, c, a in [
        (10.8, 1.0, 1.95, BLUE, 0.18),
        (11.7, 4.9, 2.35, ORANGE, 0.14),
        (8.9, 6.6, 1.55, PURPLE, 0.14),
        (8.2, 2.8, 0.95, TEAL, 0.20),
    ]:
        ax.add_patch(plt.Circle((x, y), r, facecolor=c, edgecolor="none", alpha=a))
    add_text(ax, 0.82, 0.70, "iNKT SINGLE-CELL REANALYSIS", size=11, color="#7DD3D0", weight="bold")
    add_text(ax, 0.82, 1.47, "Exact legacy QC reconstruction", size=31, color=WHITE, weight="bold")
    add_text(ax, 0.82, 2.13, "and detailed downstream reproduction", size=29, color=WHITE, weight="bold")
    add_text(
        ax,
        0.84,
        3.05,
        "Ctrl vs T2/CML across bone marrow, spleen, and thymus",
        size=15,
        color="#DCE8F2",
    )
    rounded_rect(ax, 0.82, 4.10, 10.30, 1.30, face="#173B55", edge="#31556D", radius=0.12)
    metrics = [("15,532", "cells"), ("10,670", "genes"), ("6", "samples"), ("8", "current clusters")]
    for idx, (value, label) in enumerate(metrics):
        xx = 1.13 + idx * 2.50
        add_text(ax, xx, 4.42, value, size=23, color=WHITE, weight="bold")
        add_text(ax, xx, 4.90, label, size=9, color="#AFC5D5", weight="bold")
        if idx < len(metrics) - 1:
            ax.plot([xx + 1.95, xx + 1.95], [4.30, 5.14], color="#31556D", linewidth=0.8)
    add_text(
        ax,
        0.84,
        6.24,
        (
            f"Analysis {RUN.name}  |  embedding {EMBEDDING_RUN.name}  |  "
            f"selected {ctx.sweep_summary['selected_pc']} PCs  |  count-level QC audit passed"
        ),
        size=8.5,
        color="#9FB7C8",
    )
    add_text(ax, 12.70, 7.12, f"{slide_no:02d}", size=8, color="#9FB7C8", ha="right", va="center")
    return fig


def slide_02(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Exact reconstruction of the legacy QC cohort", section="Material & QC", slide_no=n)
    samples = [
        ("Ctrl_BM", "3,998", "3,379", "84.52%"),
        ("Ctrl_Spleen", "3,829", "3,422", "89.37%"),
        ("Ctrl_Thymus", "2,101", "1,372", "65.30%"),
        ("T2_BM", "3,131", "2,731", "87.22%"),
        ("T2_Spleen", "3,760", "3,371", "89.65%"),
        ("T2_Thymus", "1,639", "1,257", "76.69%"),
        ("Total", "18,458", "15,532", "84.15%"),
    ]
    simple_table(
        ax,
        0.55,
        1.66,
        5.75,
        4.82,
        ["Sample", "Raw cells", "Post-QC", "Retained"],
        samples,
        col_widths=[1.8, 1.2, 1.2, 1.1],
        body_size=8.8,
        row_colors=[WHITE, PALE, WHITE, PALE, WHITE, PALE, TEAL_LIGHT],
    )
    section_label(ax, 6.73, 1.70, "Reconstructed filter order")
    steps = [
        ("18,458 raw cells", "all six 10x samples", BLUE),
        ("10,670 genes", "min_cells = 100 applied first", TEAL),
        ("15,532 cells", "200 <= n_genes < 2,500; mt < 5%", ORANGE),
    ]
    for idx, (value, note, color) in enumerate(steps):
        yy = 2.05 + idx * 1.12
        rounded_rect(ax, 6.72, yy, 5.82, 0.82, face=WHITE, edge=color, lw=1.2, radius=0.10)
        add_text(ax, 7.00, yy + 0.18, value, size=14, color=color, weight="bold")
        add_text(ax, 9.05, yy + 0.25, note, size=8.6, color=MID)
        if idx < 2:
            arrow(ax, 9.63, yy + 0.86, 9.63, yy + 1.07, color=MUTED)
    callout(
        ax,
        6.72,
        5.63,
        5.82,
        0.82,
        "Audit result",
        "Cell count, gene count, and all 6 post-QC sample counts match the legacy presentation exactly.",
        accent=GREEN,
        face="#EEF8F2",
        width_chars=76,
    )
    return fig


def slide_03(ctx: DeckContext, n: int):
    candidates = [int(value) for value in ctx.sweep_summary["candidate_pcs"]]
    selected_pc = int(ctx.sweep_summary["selected_pc"])
    parameters = ctx.sweep_summary["parameters"]
    fig, ax = canvas()
    header(ax, "HVG selection and PCA representation", section="Feature space", slide_no=n)
    add_image(ax, PRE / "figures/highly_variable_genes.png", 0.55, 1.63, 7.05, 4.82, label="3,000 HVGs")
    add_image(ax, PRE / "figures/pca_variance_ratio.png", 7.88, 1.63, 4.88, 3.65, label="PCA")
    callout(
        ax,
        7.88,
        5.48,
        4.88,
        0.94,
        "Nested PCA sensitivity design",
        (
            f"One {max(candidates)}-PC PCA was computed from 3,000 HVGs; neighbors/UMAP used "
            f"{', '.join(str(value) for value in candidates)} PCs. Selected {selected_pc} PCs "
            f"with {parameters['n_neighbors']} neighbors."
        ),
        accent=BLUE,
        face="#EEF4FC",
        width_chars=70,
        body_size=7.5,
    )
    return fig


def slide_04(ctx: DeckContext, n: int):
    candidates = [int(value) for value in ctx.sweep_summary["candidate_pcs"]]
    parameters = ctx.sweep_summary["parameters"]
    fig, ax = canvas()
    header(
        ax,
        f"UMAP sensitivity across {', '.join(str(value) for value in candidates)} PCs",
        section="Embedding",
        slide_no=n,
    )
    add_image(
        ax,
        SWEEP / "figures/umap_pc_sweep_tissue_cluster.png",
        0.55,
        1.53,
        12.20,
        4.05,
        trim=False,
        label="PC SWEEP",
    )
    callout(
        ax,
        0.55,
        5.73,
        12.20,
        0.73,
        "Controlled comparison",
        (
            f"Only n_pcs changes; n_neighbors={parameters['n_neighbors']}, "
            f"min_dist={parameters['min_dist']:g}, metric={parameters['metric']}, and "
            f"random_state={parameters['random_state']} are fixed. Audited Leiden labels define the same boundaries in every panel."
        ),
        accent=TEAL,
        face=TEAL_LIGHT,
        width_chars=170,
        title_size=8.5,
        body_size=7.0,
    )
    return fig


def slide_pc_selection(ctx: DeckContext, n: int):
    selected_pc = int(ctx.sweep_summary["selected_pc"])
    selection_score = float(ctx.sweep_summary["selection_score"])
    selected_metrics = next(
        row for row in ctx.sweep_metrics if int(row["n_pcs"]) == selected_pc
    )
    tissue_best = max(ctx.sweep_metrics, key=lambda row: float(row["tissue_silhouette"]))
    tissue_delta = float(tissue_best["tissue_silhouette"]) - float(
        selected_metrics["tissue_silhouette"]
    )
    fig, ax = canvas()
    header(
        ax,
        f"Selection result: {selected_pc} PCs balances separation, stability, and legacy geometry",
        section="Embedding",
        slide_no=n,
    )
    add_image(
        ax,
        SWEEP / "figures/umap_pc_sweep_metrics.png",
        0.55,
        1.52,
        12.20,
        2.34,
        trim=True,
        label="METRICS",
    )
    add_image(
        ax,
        SWEEP / f"figures/umap_npcs_{selected_pc:03d}.png",
        0.55,
        4.05,
        3.55,
        2.32,
        trim=True,
        label=f"SELECTED {selected_pc}",
    )
    add_image(
        ax,
        SWEEP / "figures/legacy_reference_200pc.png",
        4.30,
        4.05,
        3.55,
        2.32,
        trim=True,
        label="LEGACY 200",
    )
    callout(
        ax,
        8.05,
        4.05,
        4.70,
        2.32,
        "Recommendation boundary",
        (
            f"{selected_pc} PCs leads Leiden separation ({float(selected_metrics['leiden_silhouette']):.3f}), "
            f"trustworthiness ({float(selected_metrics['trustworthiness']):.3f}), and legacy centroid "
            f"concordance ({float(selected_metrics['legacy_centroid_concordance']):.3f}). The best competing "
            f"tissue silhouette improves by only {tissue_delta:.3f}. Weighted rank={selection_score:.3f} "
            "is an audit aid, not an effect size. "
            "Fixed audited Leiden boundaries were overlaid and were not recomputed. "
            "No validated discrete iNKT1, iNKT2, or NKT17 cell labels exist; continuous "
            "program scores support sensitivity assessment but are not forced subtype calls."
        ),
        accent=GREEN,
        face="#EEF8F2",
        width_chars=58,
        body_size=7.2,
    )
    return fig


def slide_05(ctx: DeckContext, n: int):
    selected_pc = int(ctx.sweep_summary["selected_pc"])
    fig, ax = canvas()
    header(
        ax,
        f"Eight fixed Leiden states on the selected {selected_pc}-PC embedding",
        section="Clustering",
        slide_no=n,
    )
    add_image(
        ax,
        SELECTED_FIGURES / "umap_selected_condition_cluster_top_markers.png",
        0.55,
        1.62,
        7.45,
        4.82,
        label="SELECTED UMAP",
    )
    add_image(
        ax,
        SELECTED_FIGURES / "trajectory_selected_paga_cluster_graph.png",
        8.22,
        1.62,
        4.53,
        3.30,
        label="SELECTED PAGA",
    )
    callout(
        ax,
        8.22,
        5.12,
        4.53,
        1.30,
        "Use conservative language",
        "Current c0-c7 labels were held fixed across the PC sweep. PAGA uses the selected graph and shows state connectivity; it is not a developmental lineage graph.",
        accent=PURPLE,
        face="#F4F0FF",
        width_chars=55,
    )
    return fig


def slide_06(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Unified iNKT1 enrichment is strongest in c4 and c6", section="iNKT1 markers", slide_no=n)
    add_image(ax, STD / "standardized_inkt1_umap_violin.png", 0.55, 1.62, 8.05, 4.82, trim=False, label="ONE SCALE")
    metric_card(ax, 8.86, 1.72, 1.85, 1.25, "67.15%", "Ctrl Xcl1+", accent=BLUE)
    metric_card(ax, 10.86, 1.72, 1.85, 1.25, "68.11%", "T2 Xcl1+", accent=ORANGE)
    callout(
        ax,
        8.86,
        3.18,
        3.85,
        1.38,
        "Unified score result",
        "c4 (+0.72 SD) and c6 (+0.81 SD) are strongest; c0 is modest (+0.20). c1 is mixed, while c2/c3 are not positively enriched.",
        accent=GREEN,
        face="#EEF8F2",
        width_chars=48,
    )
    callout(
        ax,
        8.86,
        4.83,
        3.85,
        1.58,
        "Legacy Xcl1 claim corrected",
        "Xcl1 is broadly expressed in both groups. Descriptive global log2FC is only +0.166 and varies by tissue; it is not T2-specific.",
        accent=RED,
        face="#FFF0F0",
        width_chars=48,
    )
    return fig


def slide_07(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "iNKT2 signal is strongest in c5; c7 is unclassified", section="iNKT2 markers", slide_no=n)
    add_image(ax, STD / "standardized_inkt2_umap_violin.png", 0.55, 1.62, 8.05, 4.82, trim=False, label="ONE SCALE")
    metric_card(ax, 8.86, 1.72, 1.85, 1.25, "+1.15 SD", "c5 iNKT2", accent=ORANGE)
    metric_card(ax, 10.86, 1.72, 1.85, 1.25, "-1.62 SD", "c7 iNKT2", accent=PURPLE)
    callout(
        ax,
        8.86,
        3.18,
        3.85,
        1.42,
        "c5 is a mixed program",
        "c5 is also higher for the legacy iNKT17 list (+1.54 SD). A positive iNKT2 score therefore does not establish an iNKT2 subtype.",
        accent=ORANGE,
        face="#FFF5EC",
        width_chars=48,
    )
    callout(
        ax,
        8.86,
        4.88,
        3.85,
        1.42,
        "c7 claim withdrawn",
        "All three literature-list scores in c7 are below the dataset-wide mean. The earlier raw-score winner was method/source-dependent.",
        accent=RED,
        face="#FFF0F0",
        width_chars=48,
    )
    return fig


def slide_08(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "c5 is legacy-list iNKT17-enriched, but identity is incomplete", section="iNKT17 markers", slide_no=n)
    add_image(ax, STD / "standardized_inkt17_umap_violin.png", 0.55, 1.62, 8.05, 4.82, trim=False, label="ONE SCALE")
    metric_card(ax, 8.86, 1.72, 1.85, 1.25, "+1.54 SD", "c5 iNKT17", accent=RED)
    metric_card(ax, 10.86, 1.72, 1.85, 1.25, "+0.39 SD", "vs iNKT2", accent=GOLD, note="margin CI +0.24 to +0.54")
    callout(
        ax,
        8.86,
        3.18,
        3.85,
        1.48,
        "Marker coverage limits identity",
        "Il17a, Il17f, and Ccr6 are unavailable after strict legacy gene filtering. Rorc is retained but sparse; broad activation genes also contribute.",
        accent=ORANGE,
        face="#FFF5EC",
        width_chars=48,
    )
    callout(
        ax,
        8.86,
        4.92,
        3.85,
        1.40,
        "Permitted conclusion",
        "c5 is enriched for the legacy literature iNKT17 list. This is a continuous program result, not proof of a discrete iNKT17 cell subtype.",
        accent=TEAL,
        face=TEAL_LIGHT,
        width_chars=48,
    )
    return fig


def slide_09(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "One common scale replaces forced subtype winners", section="Unified scoring", slide_no=n)
    add_image(ax, STD / "standardized_cluster_heatmap.png", 0.55, 1.58, 8.12, 2.82, trim=False, label="COMMON SCALE")
    add_text(ax, 0.62, 4.58, "Legacy raw-score maps (formula and normalization undocumented)", size=8.3, color=NAVY, weight="bold")
    add_legacy_ppt_media(ax, "image20.png", 0.55, 4.78, 2.55, 1.55)
    add_legacy_ppt_media(ax, "image24.png", 3.33, 4.78, 2.55, 1.55)
    add_legacy_ppt_media(ax, "image28.png", 6.11, 4.78, 2.55, 1.55)
    callout(
        ax,
        8.94,
        1.62,
        3.78,
        1.36,
        "Unified current ruler",
        "Each gene is z-scored across 15,532 cells; genes are averaged within a legacy list; each list is then standardized to mean 0 and SD 1.",
        accent=BLUE,
        face="#EEF4FC",
        width_chars=47,
    )
    callout(
        ax,
        8.94,
        3.24,
        3.78,
        1.35,
        "Coverage",
        "Measured genes: iNKT1 25/26, iNKT2 76/86, iNKT17 30/36. Missing genes are excluded, not treated as zero expression.",
        accent=TEAL,
        face=TEAL_LIGHT,
        width_chars=47,
    )
    callout(
        ax,
        8.94,
        4.85,
        3.78,
        1.48,
        "Comparison boundary",
        "Old maps use unequal 0-400-like ranges and an unknown formula. Compare spatial patterns only; current SD values cannot be converted into legacy raw units.",
        accent=ORANGE,
        face="#FFF5EC",
        width_chars=47,
    )
    return fig


def slide_10(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Legacy global DE directions are retained within the tested HVG universe", section="Global DE", slide_no=n, footer=DE_FOOTER)
    metric_card(ax, 0.55, 1.70, 3.55, 1.40, "16", "Pass current default threshold", accent=GREEN, note="FDR <= 0.05 and |logFC| >= 0.25")
    metric_card(ax, 4.30, 1.70, 3.55, 1.40, "4", "FDR-only legacy genes", accent=GOLD, note="tested and significant, smaller effect")
    metric_card(ax, 8.05, 1.70, 4.70, 1.40, "19", "Not tested in global HVG ranking", accent=MUTED, note="not equivalent to nonsignificant")
    simple_table(
        ax,
        0.55,
        3.45,
        7.65,
        2.80,
        ["Representative gene", "Current logFC", "Legacy direction"],
        [
            ("Iglc2", "+5.89", "up"),
            ("Slc15a2", "-2.40", "down"),
            ("Eps8l1", "+1.15", "up"),
            ("Hspa1b", "+0.90", "up"),
            ("Fos / Jun", "+0.77 / +0.73", "up"),
            ("Sfpq / Dnaja1", "-0.70 / -0.66", "down"),
        ],
        col_widths=[2.4, 2.0, 2.2],
        body_size=8.4,
    )
    callout(
        ax,
        8.47,
        3.45,
        4.28,
        2.80,
        "Correct denominator",
        "20/20 tested legacy genes are FDR-significant and directionally concordant. Global DE ranks only 3,000 HVGs, whereas tissue and cluster DE test all 10,670 retained genes. Do not report 16/39 as the sole recovery rate.",
        accent=TEAL,
        face=TEAL_LIGHT,
        width_chars=52,
        body_size=8.6,
    )
    return fig


def slide_11(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Tissue-stratified DE strongly recovers the legacy core", section="Tissue DE", slide_no=n, footer=DE_FOOTER)
    plot_recovery_bars(ax, 0.55, 1.65, 6.05, 3.55)
    simple_table(
        ax,
        6.88,
        1.65,
        5.87,
        3.55,
        ["Tissue", "Current DEGs", "T2 up / down", "Legacy recovery"],
        [
            ("Thymus", "352", "195 / 157", "31 / 34"),
            ("Bone marrow", "718", "350 / 368", "35 / 39"),
            ("Spleen", "864", "356 / 508", "44 / 50"),
        ],
        col_widths=[1.5, 1.25, 1.55, 1.55],
        body_size=9.0,
        row_colors=["#F4F0FF", "#EEF4FC", "#FFF5EC"],
    )
    callout(
        ax,
        0.55,
        5.50,
        12.20,
        0.87,
        "QC sensitivity matters",
        "Thymus changed from 1,638 DEGs (1,596 up / 42 down) under permissive QC to 352 (195 up / 157 down) under Legacy-QC: the apparent response is now smaller and directionally balanced.",
        accent=ORANGE,
        face="#FFF5EC",
        width_chars=160,
    )
    return fig


def slide_12(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Tissue identity dominates cluster composition", section="Composition", slide_no=n)
    add_image(ax, PRE / "figures/cluster_percent_by_sample.png", 0.55, 1.62, 8.35, 4.82, label="BY SAMPLE")
    anchor_rows = [
        ("c0", "4,828", "95.8% BM", "45.5%"),
        ("c3", "5,677", "93.7% spleen", "47.2%"),
        ("c6", "2,379", "97.3% thymus", "47.2%"),
        ("c7", "92", "96.7% thymus", "47.8%"),
    ]
    simple_table(
        ax,
        9.15,
        1.62,
        3.60,
        3.10,
        ["Cluster", "N", "Tissue", "T2"],
        anchor_rows,
        col_widths=[0.8, 0.8, 1.55, 0.85],
        body_size=7.8,
    )
    callout(
        ax,
        9.15,
        4.98,
        3.60,
        1.44,
        "No condition-exclusive cluster",
        "T2 fractions span only 45.5%-54.4% across c0-c7. Condition effects appear within tissue states rather than as a separate T2 island.",
        accent=GREEN,
        face="#EEF8F2",
        width_chars=43,
    )
    return fig


def response_slide(
    n: int,
    title: str,
    section: str,
    tissue: str,
    mapping: str,
    recovered: str,
    concordance: str,
    sizes: str,
    current_de: str,
    shared_genes: Sequence[str],
    missing_genes: Sequence[str],
    interpretation: str,
    *,
    accent: str,
):
    fig, ax = canvas()
    header(ax, title, section=section, slide_no=n, footer=DE_FOOTER)
    metric_card(ax, 0.55, 1.68, 3.62, 1.40, recovered, "Legacy genes recovered", accent=accent, note=concordance)
    metric_card(ax, 4.37, 1.68, 3.62, 1.40, current_de, "Current significant DEGs", accent=TEAL, note=sizes)
    metric_card(ax, 8.19, 1.68, 4.56, 1.40, tissue, "Tissue-constrained mapping", accent=TISSUE.get(tissue, PURPLE), note=mapping)
    rounded_rect(ax, 0.55, 3.38, 7.35, 2.75, face=WHITE, edge=LIGHT, radius=0.10)
    section_label(ax, 0.82, 3.68, "Representative recovered genes", accent)
    x = 0.82
    y = 4.02
    for gene in shared_genes:
        chip_w = 0.42 + 0.085 * len(gene)
        if x + chip_w > 7.60:
            x = 0.82
            y += 0.49
        gene_chip(ax, x, y, gene, color=accent, width=chip_w)
        x += chip_w + 0.12
    add_text(ax, 0.82, 5.35, "Below current threshold", size=8.2, color=MUTED, weight="bold")
    add_text(ax, 2.55, 5.35, ", ".join(missing_genes), size=8.2, color=RED)
    callout(
        ax,
        8.16,
        3.38,
        4.59,
        2.75,
        "Interpretation",
        interpretation,
        accent=accent,
        face=PALE,
        width_chars=56,
        body_size=8.7,
    )
    return fig


def driver_slide(
    n: int,
    title: str,
    section: str,
    main_metric: str,
    mapping: str,
    gene_groups: Sequence[tuple[str, Sequence[str], str]],
    explanation: str,
    caution: str,
    *,
    accent: str,
):
    fig, ax = canvas()
    header(ax, title, section=section, slide_no=n, footer=DE_FOOTER)
    metric_card(ax, 0.55, 1.67, 3.70, 1.42, main_metric, "Legacy drivers retained", accent=accent, note=mapping)
    callout(ax, 4.52, 1.67, 8.23, 1.42, "Module-level conclusion", explanation, accent=TEAL, face=TEAL_LIGHT, width_chars=102)
    box_y = 3.40
    for idx, (label, genes, color) in enumerate(gene_groups):
        xx = 0.55 + idx * 4.12
        rounded_rect(ax, xx, box_y, 3.84, 1.85, face=WHITE, edge=color, lw=1.0, radius=0.10)
        section_label(ax, xx + 0.24, box_y + 0.30, label, color)
        add_text(ax, xx + 0.24, box_y + 0.66, wrap("  ".join(genes), 33), size=9.0, color=INK, weight="bold")
    callout(ax, 0.55, 5.55, 12.20, 0.82, "Do not over-interpret pathway labels", caution, accent=ORANGE, face="#FFF5EC", width_chars=160, body_size=8.2)
    return fig


def slide_13(ctx: DeckContext, n: int):
    return response_slide(
        n,
        "Legacy c1 thymus response is retained in current c6",
        "Legacy c1 DE",
        "Thymus",
        "legacy c1 -> current c6 (post-hoc program match)",
        "26 / 30 (86.7%)",
        "26/26 directionally concordant",
        "140 up / 164 down; 1,087 T2 / 1,227 Ctrl",
        "304",
        ["Hspa1a", "Hspa1b", "Fos", "Dusp1", "Tmsb10", "Cd52", "Hspa8", "Dnaja1", "Ccl5", "Cxcr6"],
        ["Fau", "Rplp1", "mt-Co3", "mt-Cytb"],
        "The thymus-dominant activation/proteostasis response is highly retained. This validates a response program, not the numerical identity of legacy c1.",
        accent=PURPLE,
    )


def slide_14(ctx: DeckContext, n: int):
    return driver_slide(
        n,
        "All eight legacy c1 pathway drivers are retained in current c6",
        "Legacy c1 pathways",
        "8 / 8",
        "current c6 thymus",
        [
            ("Activation / IEG", ["Fos", "Dusp1"], ORANGE),
            ("Proteostasis", ["Hspa1a", "Hspa1b", "Hspa8", "Dnaja1"], TEAL),
            ("Tissue / mitochondrial", ["Ccl5", "Cox8a"], PURPLE),
        ],
        "The complete driver module recurs after exact QC alignment, improving over the earlier permissive-QC result (6/8).",
        "The 38 legacy terms were not recomputed with the original Reactome/KEGG library. Opposing Hspa1a/b versus Hspa8/Dnaja1 directions imply remodeling, not blanket heat-shock activation.",
        accent=PURPLE,
    )


def slide_15(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Legacy c2 response is represented by current c4/c3", section="Legacy c2 DE", slide_no=n, footer=DE_FOOTER)
    callout(ax, 0.55, 1.65, 5.95, 3.10, "Bone marrow: compact and clear", "Current c4 BM: 223 T2 / 244 Ctrl; 83 DEGs (56 up / 27 down). Recovery is 17/18 (94.4%), with 17/17 sign agreement; only Ppp1r12a is below threshold.", accent=BLUE, face="#EEF4FC", width_chars=72, body_size=9.0)
    callout(ax, 6.80, 1.65, 5.95, 3.10, "Spleen: split across states", "Maximum coverage: current c3 spleen, 19/27 recovered. Compact alternative: current c4 spleen, 18/27 with higher Jaccard. The two candidates share 14 recovered genes but each also contributes distinct genes.", accent=ORANGE, face="#FFF5EC", width_chars=72, body_size=9.0)
    simple_table(
        ax,
        0.55,
        5.02,
        12.20,
        1.30,
        ["Legacy unit", "Primary current candidate", "Recovery", "Alternative", "Interpretation"],
        [
            ("c2 bone marrow", "c4-BM", "17/18", "none needed", "compact match"),
            ("c2 spleen", "c3-spleen", "19/27", "c4-spleen 18/27", "boundary split"),
        ],
        col_widths=[2.0, 2.4, 1.3, 2.4, 3.2],
        body_size=8.3,
    )
    return fig


def slide_16(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Legacy c2 pathway drivers are recoverable but distributed", section="Legacy c2 pathways", slide_no=n, footer=DE_FOOTER)
    metric_card(ax, 0.55, 1.70, 3.62, 1.38, "4 / 4", "BM drivers in c4", accent=BLUE, note="Cox8a, Dnaja1, Hspa8, Ndufa4")
    metric_card(ax, 4.36, 1.70, 3.62, 1.38, "7 / 7", "Spleen drivers in c3", accent=ORANGE, note="maximum-coverage state")
    metric_card(ax, 8.17, 1.70, 4.58, 1.38, "4 / 7", "Spleen drivers in c4", accent=GOLD, note="higher-specificity alternative")
    callout(ax, 0.55, 3.42, 5.93, 2.72, "Current curated ORA", "c4-BM supports Direct TCR activation and Residency. c3-spleen strongly supports both programs, whereas c4-spleen has no FDR < 0.05 curated term.", accent=TEAL, face=TEAL_LIGHT, width_chars=72, body_size=9.0)
    callout(ax, 6.80, 3.42, 5.95, 2.72, "Why two candidates can both be useful", "Driver completeness favors broad c3 in spleen; DEG-set specificity favors compact c4. The divergence is evidence that the legacy response was redistributed, not that one algorithm found the uniquely correct cluster.", accent=PURPLE, face="#F4F0FF", width_chars=72, body_size=9.0)
    return fig


def slide_17(ctx: DeckContext, n: int):
    return response_slide(
        n,
        "Legacy c5 BM response is strongly recovered in current c0",
        "Legacy c5 DE",
        "Bone marrow",
        "legacy c5 -> current c0 (BM-dominant anchor)",
        "53 / 61 (86.9%)",
        "53/53 directionally concordant",
        "274 up / 318 down; 2,090 T2 / 2,533 Ctrl",
        "592",
        ["Fos", "Jun", "Junb", "Dusp1", "Hspa1a", "Hspa1b", "Hspa8", "Dnaja1", "Pink1", "Tomm6", "Ndufa13"],
        ["Actg1", "Fau", "Ptpn18", "Rplp1/2", "mt-Co2/3", "mt-Cytb"],
        "This broad BM anchor recovers the core activation/proteostasis/mitochondrial response. Because c0 has 592 DEGs, it may consolidate several legacy BM states.",
        accent=BLUE,
    )


def slide_18(ctx: DeckContext, n: int):
    return driver_slide(
        n,
        "A conserved BM activation-proteostasis-mitochondrial response",
        "Legacy c5 pathways",
        "15 / 15",
        "legacy c5 drivers in current c0-BM",
        [
            ("Immediate-early", ["Fos", "Jun", "Dusp1"], ORANGE),
            ("Chaperone / stress", ["Hspa1a/b", "Hspa8", "Hsph1", "Dnaja1"], TEAL),
            ("Respiratory chain", ["Cox6b1", "Cox7a2/b", "Cox8a", "Ndufa2/7/13"], BLUE),
        ],
        "Current c0-BM also shows T2-up Direct TCR activation and Residency enrichment at FDR ~10^-6.",
        "Legacy 'Prion disease' labels are driven by shared mitochondrial/chaperone genes. Describe the recovered module, not a neurodegenerative disease mechanism; the two curated sets also share IEG genes.",
        accent=BLUE,
    )


def slide_19(ctx: DeckContext, n: int):
    return response_slide(
        n,
        "Legacy c6 spleen has a compact current c5 match",
        "Legacy c6 DE",
        "Spleen",
        "legacy c6 -> current c5 (compact response match)",
        "22 / 29 (75.9%)",
        "22/22 directionally concordant",
        "44 up / 29 down; 387 T2 / 267 Ctrl",
        "73",
        ["Jun", "Junb", "Ndufv3", "Tmsb10", "Hspa8", "Dnaja1", "Hsp90ab1", "Dynll1"],
        ["Cox8a", "H3f3b", "mt-Co3", "mt-Nd4", "Rplp2", "Uba52", "Ubb"],
        "Response similarity and subtype identity are separate. On the unified legacy-list scale, c5 carries both iNKT2 (+1.15 SD) and iNKT17 (+1.54 SD) programs; treat it as mixed.",
        accent=ORANGE,
    )


def slide_20(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "The legacy c6 HSP/IEG module is largely retained", section="Legacy c6 pathways", slide_no=n, footer=DE_FOOTER)
    metric_card(ax, 0.55, 1.70, 3.80, 1.40, "5 / 6", "Drivers in compact c5-spleen", accent=ORANGE, note="Cox8a is below threshold")
    metric_card(ax, 4.58, 1.70, 3.80, 1.40, "6 / 6", "Drivers in broad c3-spleen", accent=TEAL, note="higher coverage; lower specificity")
    metric_card(ax, 8.60, 1.70, 4.15, 1.40, "FDR ~0.005", "Direct TCR / Residency in c5", accent=GREEN)
    rounded_rect(ax, 0.55, 3.45, 12.20, 2.70, face=WHITE, edge=LIGHT, radius=0.10)
    add_text(ax, 0.84, 3.76, "Bidirectional response pattern", size=11, color=NAVY, weight="bold")
    add_text(ax, 1.05, 4.35, "IEG activation", size=9.3, color=ORANGE, weight="bold")
    arrow(ax, 2.58, 4.43, 4.08, 4.43, color=ORANGE, lw=2)
    add_text(ax, 4.28, 4.43, "Jun / Junb up", size=11, color=ORANGE, weight="bold", va="center")
    add_text(ax, 1.05, 5.25, "Constitutive chaperones", size=9.3, color=TEAL, weight="bold")
    arrow(ax, 3.05, 5.33, 4.08, 5.33, color=TEAL, lw=2)
    add_text(ax, 4.28, 5.33, "Hspa8 / Dnaja1 / Hsp90ab1 down", size=11, color=TEAL, weight="bold", va="center")
    callout(ax, 8.38, 4.05, 3.98, 1.50, "Interpretation", "A remodeling program is more accurate than 'heat-shock pathway activation'. The original 13 pathway terms were not recomputed.", accent=PURPLE, face="#F4F0FF", width_chars=48)
    return fig


def slide_21(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Legacy c7 BM signal is recoverable but not uniquely mappable", section="Legacy c7 DE", slide_no=n, footer=DE_FOOTER)
    metric_card(ax, 0.55, 1.68, 3.76, 1.42, "3 / 4", "Recovered in c0-BM", accent=BLUE, note="broad 592-DEG anchor")
    metric_card(ax, 4.56, 1.68, 3.76, 1.42, "3 / 4", "Recovered in c5-BM", accent=ORANGE, note="compact 39-DEG candidate")
    metric_card(ax, 8.57, 1.68, 4.18, 1.42, "c5 mixed", "Subtype evidence incomplete", accent=GOLD, note="legacy-list iNKT17 enrichment; core markers missing")
    simple_table(
        ax,
        0.55,
        3.42,
        5.80,
        2.60,
        ["Legacy gene", "Current status", "Interpretation"],
        [
            ("Cox7c", "recovered", "mitochondrial"),
            ("Gm10076", "recovered", "general response"),
            ("Sfpq", "recovered", "RNA processing"),
            ("Rplp2", "below threshold", "ribosomal"),
        ],
        col_widths=[1.5, 1.8, 2.5],
        body_size=8.5,
    )
    callout(ax, 6.65, 3.42, 6.10, 2.60, "Why the legacy subtype label is not retained", "Four mitochondrial/ribosomal/RNA-processing genes cannot establish subtype identity. Unified scoring highlights a c5 iNKT17-related program, but Il17a/Ccr6 are unavailable and c5 also has a strong iNKT2 score; retain a mixed-program interpretation.", accent=RED, face="#FFF0F0", width_chars=75, body_size=9.0)
    return fig


def slide_22(ctx: DeckContext, n: int):
    return response_slide(
        n,
        "Legacy c8 spleen response is highly covered by current c3",
        "Legacy c8 DE",
        "Spleen",
        "legacy c8 -> current c3 (spleen-dominant anchor)",
        "73 / 84 (86.9%)",
        "73/73 directionally concordant",
        "275 up / 460 down; 2,557 T2 / 2,762 Ctrl",
        "735",
        ["Fos", "Jun", "Junb", "Jund", "Dusp1", "Hspa1b", "Hspa8", "Hsp90ab1", "Ndufa13", "Xcl1", "Ccnd2"],
        ["Actb", "Fau", "mt-Co2/3", "Myl6", "Ptpn18", "Rplp1/2", "Tmsb4x", "Trac", "Xist"],
        "This is high coverage by a broad spleen response, not proof of high specificity or identity preservation. On the unified legacy-list scale, current c3 is near baseline for all three subtype programs and remains unclassified.",
        accent=ORANGE,
    )


def slide_23(ctx: DeckContext, n: int):
    return driver_slide(
        n,
        "All 20 legacy c8 pathway drivers are retained in current c3",
        "Legacy c8 pathways",
        "20 / 20",
        "current c3 spleen",
        [
            ("Activation", ["Fos", "Jun", "Jund", "Dusp1", "Ppp1r15a"], ORANGE),
            ("Proteostasis", ["Hspa8", "Hsph1", "Hsp90ab1", "Dnaja1"], TEAL),
            ("Mitochondrial", ["Cox6b1/7a2/7b/8a", "Ndufa2/4/7/13", "Uqcrb"], BLUE),
        ],
        "Current c3-spleen strongly enriches Direct TCR activation and Residency among T2-up genes (FDR ~10^-6).",
        "The original Prion/Measles/Estrogen/MAPK labels reflect shared driver genes and were not recomputed using the original library. Report the driver module rather than literal disease mechanisms.",
        accent=ORANGE,
    )


def slide_24(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Legacy c9 response is distributed across current states", section="Legacy c9 DE", slide_no=n, footer=DE_FOOTER)
    callout(ax, 0.55, 1.65, 5.95, 3.18, "Bone marrow", "Maximum coverage: c0-BM recovers 9/11 with 9/9 sign agreement. Compact c4-BM recovers 8/11. The missing genes are primarily Actb/Rplp2 or H2afj/Tma7/Rplp2, depending on candidate.", accent=BLUE, face="#EEF4FC", width_chars=72, body_size=9.0)
    callout(ax, 6.80, 1.65, 5.95, 3.18, "Spleen", "Maximum coverage: c3-spleen recovers 19/28. Compact alternatives c4 and c5 recover 12/28 and 13/28. c4 has stronger set enrichment; c5 has more complete pathway-driver recovery.", accent=ORANGE, face="#FFF5EC", width_chars=72, body_size=9.0)
    callout(ax, 0.55, 5.15, 12.20, 1.05, "Conservative conclusion", "The legacy cross-tissue response has been absorbed into several current BM/spleen states. This supports cluster consolidation and boundary redrawing rather than a one-to-one rename.", accent=PURPLE, face="#F4F0FF", width_chars=150, body_size=9.0)
    return fig


def slide_25(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "The legacy c9 chaperone axis is retained across states", section="Legacy c9 pathways", slide_no=n, footer=DE_FOOTER)
    metric_card(ax, 0.55, 1.70, 3.55, 1.42, "2 / 2", "BM drivers in c0 or c4", accent=BLUE, note="Dnaja1, Hspa8")
    metric_card(ax, 4.30, 1.70, 3.55, 1.42, "4 / 4", "Spleen drivers in c3", accent=ORANGE, note="Dnaja1, Hspa8, Hsp90ab1, Jun")
    metric_card(ax, 8.05, 1.70, 4.70, 1.42, "4 / 4", "Spleen drivers in c5", accent=TEAL, note="compact alternative")
    rounded_rect(ax, 0.55, 3.48, 12.20, 2.52, face=WHITE, edge=LIGHT, radius=0.10)
    add_text(ax, 0.85, 3.83, "Stable chaperone axis", size=12, color=NAVY, weight="bold")
    gene_chip(ax, 0.85, 4.35, "Dnaja1", color=TEAL, width=1.10)
    gene_chip(ax, 2.15, 4.35, "Hspa8", color=TEAL, width=1.00)
    gene_chip(ax, 3.35, 4.35, "Hsp90ab1", color=TEAL, width=1.25)
    gene_chip(ax, 4.80, 4.35, "Jun", color=ORANGE, width=0.78)
    callout(ax, 6.18, 4.00, 6.15, 1.35, "Interpretation limit", "Two-to-four recurring drivers cannot establish cluster identity. The legacy page also lacked a clear BM/spleen label for its lower table, so the revised deck makes tissue scope explicit.", accent=GOLD, face="#FFF8E8", width_chars=76)
    return fig


def slide_26(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Detectable response depends on test status and cell count", section="Zero DEG strata", slide_no=n, footer=DE_FOOTER)
    rows = [
        ("c1", "32/46 | tested: 0", "45/36 | tested: 0", "3 cells | skipped"),
        ("c2", "145/233 | DE", "74/45 | tested: 0", "50/64 | tested: 0"),
        ("c5", "181/228 | DE", "387/267 | DE", "46/20 | tested: 0"),
        ("c7", "0 cells | skipped", "3 cells | skipped", "44/45 | tested: 0"),
    ]
    simple_table(
        ax,
        0.55,
        1.66,
        12.20,
        3.65,
        ["Current cluster", "Bone marrow T2/Ctrl", "Spleen T2/Ctrl", "Thymus T2/Ctrl"],
        rows,
        col_widths=[1.6, 3.2, 3.2, 3.2],
        body_size=8.3,
    )
    callout(ax, 0.55, 5.60, 5.90, 0.78, "Tested + zero DEG", "A completed test found no genes passing the threshold; this is not proof of zero biological effect.", accent=GOLD, face="#FFF8E8", width_chars=72, body_size=7.8)
    callout(ax, 6.75, 5.60, 6.00, 0.78, "Skipped", "At least one condition had <20 cells. Never combine this category with a tested zero result.", accent=RED, face="#FFF0F0", width_chars=72, body_size=7.8)
    return fig


def slide_27(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Legacy tissue DEG cores are recovered; anchors show balanced overlap", section="DEG overlap", slide_no=n, footer=DE_FOOTER)
    plot_recovery_bars(ax, 0.55, 1.65, 5.45, 3.55)
    add_image(ax, EXT / "de_pathway/overlap_pathway/current_tissue_vs_cluster_overlap_jaccard.png", 6.28, 1.65, 6.47, 3.55, label="JACCARD")
    simple_table(
        ax,
        0.55,
        5.45,
        12.20,
        0.92,
        ["Anchor", "Anchor DEGs", "Shared with tissue", "Up Jaccard", "Down Jaccard", "Anchor-only"],
        [
            ("c0-BM", "592", "489", "0.629", "0.566", "103"),
            ("c3-spleen", "735", "600", "0.656", "0.566", "135"),
            ("c6-thymus", "304", "261", "0.667", "0.655", "43"),
        ],
        col_widths=[1.7, 1.7, 2.4, 1.7, 1.8, 1.5],
        body_size=7.3,
        header_size=7.2,
    )
    return fig


def slide_28(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Legacy response programs map to current tissue anchors", section="Response validation", slide_no=n, footer=DE_FOOTER)
    mappings = [
        ("legacy c1 thymus", "current c6 thymus", "26 / 30", "86.7%", PURPLE, "Klf6, Cxcr6, Rhob"),
        ("legacy c5 BM", "current c0 BM", "53 / 61", "86.9%", BLUE, "Tomm6, Pink1, Fus"),
        ("legacy c8 spleen", "current c3 spleen", "73 / 84", "86.9%", ORANGE, "Atp5j2, Actg1, Xcl1, Ccnd2"),
    ]
    for idx, (legacy, current, recovered, pct, color, genes) in enumerate(mappings):
        xx = 0.55 + idx * 4.13
        rounded_rect(ax, xx, 1.70, 3.84, 3.92, face=WHITE, edge=color, lw=1.2, radius=0.12)
        add_text(ax, xx + 0.25, 2.00, legacy, size=10.5, color=MID, weight="bold")
        arrow(ax, xx + 1.05, 2.46, xx + 2.78, 2.46, color=color, lw=2)
        add_text(ax, xx + 0.25, 2.80, current, size=12.2, color=color, weight="bold")
        add_text(ax, xx + 0.25, 3.45, recovered, size=24, color=color, weight="bold")
        add_text(ax, xx + 2.63, 3.56, pct, size=10, color=color, weight="bold")
        add_text(ax, xx + 0.25, 4.22, "Representative recovered genes", size=7.7, color=MUTED, weight="bold")
        add_text(ax, xx + 0.25, 4.58, wrap(genes, 32), size=9.0, color=INK)
        add_text(ax, xx + 0.25, 5.15, "All retained signs concordant", size=7.8, color=GREEN, weight="bold")
    callout(ax, 0.55, 5.74, 12.20, 0.72, "Key distinction", "These are post-hoc, tissue-constrained response-program matches—not independent validation of legacy cluster identity.", accent=RED, face="#FFF0F0", width_chars=170, body_size=7.8)
    return fig


def slide_29(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Gene-level response hubs recur; the interaction network was not reproduced", section="Legacy network", slide_no=n, footer=DE_FOOTER)
    metric_card(ax, 0.55, 1.68, 3.55, 1.42, "5 / 10", "Legacy hubs in c6-thymus", accent=PURPLE)
    metric_card(ax, 4.30, 1.68, 3.55, 1.42, "7 / 10", "Legacy hubs in c0-BM", accent=BLUE)
    metric_card(ax, 8.05, 1.68, 4.70, 1.42, "9 / 10", "Legacy hubs in c3-spleen", accent=ORANGE)
    rounded_rect(ax, 0.55, 3.45, 5.72, 2.60, face=WHITE, edge=LIGHT, radius=0.10)
    section_label(ax, 0.83, 3.77, "Shared across all three anchors")
    for idx, gene in enumerate(["Ndufa13", "Cox6c", "Atp5e", "Hspa1b"]):
        gene_chip(ax, 0.85 + (idx % 2) * 2.45, 4.22 + (idx // 2) * 0.65, gene, color=TEAL, width=1.55)
    callout(ax, 6.58, 3.45, 6.17, 2.60, "What was not reproduced", "The legacy interaction database, species mapping, confidence cutoff, background, edge definition, degree, and topology were not documented. PAGA is a cell-state graph and cannot substitute for a gene interaction network.", accent=RED, face="#FFF0F0", width_chars=76, body_size=9.0)
    return fig


def slide_30(ctx: DeckContext, n: int):
    selected_pc = int(ctx.sweep_summary["selected_pc"])
    trajectory = ctx.sweep_summary["trajectory"]
    root_cluster = str(trajectory["root_cluster"])
    cluster_medians = {
        str(cluster): float(value)
        for cluster, value in trajectory["cluster_median_dpt"].items()
    }
    other_medians = [
        value for cluster, value in cluster_medians.items() if cluster != root_cluster
    ]
    fig, ax = canvas()
    header(
        ax,
        f"Selected {selected_pc}-PC topology; DPT is a root-versus-rest ordering",
        section="Trajectory",
        slide_no=n,
    )
    add_image(
        ax,
        SELECTED_FIGURES / "trajectory_selected_umap_cluster_dpt.png",
        0.55,
        1.62,
        8.15,
        4.82,
        label="SELECTED UMAP + DPT",
    )
    metric_card(
        ax,
        8.95,
        1.68,
        3.80,
        1.36,
        f"c{root_cluster}",
        "Audited root cluster",
        accent=PURPLE,
        note=(
            f"{trajectory['root_cluster_n_cells']} cells; "
            f"{100 * float(trajectory['root_cluster_dominant_tissue_fraction']):.1f}% "
            f"{str(trajectory['root_cluster_dominant_tissue']).replace('_', ' ')}"
        ),
    )
    metric_card(
        ax,
        8.95,
        3.24,
        3.80,
        1.36,
        f"{cluster_medians[root_cluster]:.3f}",
        f"c{root_cluster} median DPT",
        accent=PURPLE,
        note="recomputed on selected graph",
    )
    metric_card(
        ax,
        8.95,
        4.80,
        3.80,
        1.36,
        f"{min(other_medians):.3f}-{max(other_medians):.3f}",
        "Other cluster medians",
        accent=RED,
        note="selected-graph root sensitivity",
    )
    return fig


def slide_31(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Within-tissue T2 responses outperform 'tissue-specific markers'", section="Marker panels", slide_no=n, footer=DE_FOOTER)
    add_image(ax, EXT / "signatures/marker_validation/marker_t2_ctrl_effect_heatmap.png", 0.55, 1.62, 6.25, 4.82, label="EFFECTS")
    add_image(ax, EXT / "signatures/marker_validation/marker_expression_condition_tissue_dotplot.png", 7.05, 1.62, 5.70, 3.25, label="FRACTION + MEAN")
    simple_table(
        ax,
        7.05,
        5.08,
        5.70,
        1.30,
        ["Legacy panel", "Recovered", "Rate"],
        [("Thymus", "8 / 10", "80.0%"), ("Bone marrow", "16 / 22", "72.7%"), ("Spleen", "20 / 27", "74.1%")],
        col_widths=[2.2, 1.5, 1.5],
        body_size=7.8,
    )
    return fig


def slide_32(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "A four-gene cross-tissue signal is retained; rare signals remain fragile", section="Cross-tissue candidates", slide_no=n, footer=DE_FOOTER)
    simple_table(
        ax,
        0.55,
        1.66,
        8.15,
        3.95,
        ["Gene", "BM logFC", "Spleen logFC", "Thymus logFC", "Direction"],
        [
            ("Iglc2", "+5.851", "+5.456", "+25.767", "up in all"),
            ("Slc15a2", "-2.734", "-2.613", "-1.784", "down in all"),
            ("Xlr", "-3.332", "-3.410", "-2.476", "down in all"),
            ("3830403N18Rik", "-2.074", "-2.538", "-3.827", "down in all"),
        ],
        col_widths=[2.2, 1.45, 1.55, 1.55, 1.65],
        body_size=8.5,
        row_colors=["#EEF8F2", PALE, WHITE, PALE],
    )
    callout(ax, 8.98, 1.66, 3.77, 1.72, "Not retained at current threshold", "P2rx7 (thymus), Mttp (thymus), and St3gal3 (BM) do not pass the current tissue-DE threshold.", accent=RED, face="#FFF0F0", width_chars=46, body_size=8.3)
    callout(ax, 8.98, 3.68, 3.77, 1.93, "Rare-cell signals", "S100a8/a9 pass in spleen but are expressed in <=1.02% of cells. Large logFC can reflect near-zero control means, rare cells, ambient RNA, or composition shifts.", accent=ORANGE, face="#FFF5EC", width_chars=46, body_size=8.3)
    callout(ax, 0.55, 5.72, 12.20, 0.74, "Interpretation", "Directionally stable does not mean high-frequency or mechanistically dominant; always report fraction expressing with logFC and FDR.", accent=TEAL, face=TEAL_LIGHT, width_chars=170, body_size=7.5)
    return fig


def slide_33(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Curated activation recurs; legacy terms were not recomputed", section="Tissue pathways", slide_no=n, footer=DE_FOOTER)
    simple_table(
        ax,
        0.55,
        1.63,
        7.55,
        4.82,
        ["Tissue / direction", "Curated term", "Overlap", "FDR"],
        [
            ("BM up", "Direct TCR activation", "6/11", "5.76e-6"),
            ("BM up", "Residency", "9/36", "1.02e-5"),
            ("Spleen up", "Direct TCR activation", "6/11", "6.37e-6"),
            ("Spleen up", "Residency", "9/36", "1.18e-5"),
            ("Thymus up", "iNKT17 literature", "6/30", "1.71e-4"),
            ("Thymus up", "Direct TCR activation", "3/11", "0.00356"),
            ("Thymus down", "iNKT1 Wang", "7/44", "3.80e-5"),
            ("Thymus down", "iNKT1 in-house", "6/35", "6.31e-5"),
        ],
        col_widths=[2.0, 2.8, 1.2, 1.25],
        body_size=7.7,
        header_size=7.7,
    )
    callout(ax, 8.40, 1.63, 4.35, 1.45, "What is supported", "BM and spleen repeatedly show T2-up Direct TCR/Residency overlap; thymus shows a tissue-dependent mixture of activation and iNKT1-down genes.", accent=GREEN, face="#EEF8F2", width_chars=53, body_size=8.2)
    callout(ax, 8.40, 3.36, 4.35, 1.45, "What is not supported", "A thymus iNKT17 overlap does not imply more iNKT17 cells. Standardized continuous program enrichment is not a cell count or proof of discrete subtype conversion.", accent=RED, face="#FFF0F0", width_chars=53, body_size=8.2)
    callout(ax, 8.40, 5.08, 4.35, 1.37, "Legacy term boundary", "Original PAGER/Reactome/KEGG terms were recovered from the deck/workbooks but not recomputed with the same library, version, species mapping, or background.", accent=ORANGE, face="#FFF5EC", width_chars=53, body_size=8.0)
    return fig


def slide_34(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "What this reproduction supports", section="Evidence hierarchy", slide_no=n)
    columns = [
        ("STRONG", GREEN, ["Exact count-level QC cohort reconstruction", "Stable tissue-dominated structure", "88.0%-91.2% tissue-DE core recovery", "Complete major driver-module recovery"]),
        ("SUPPORTED WITH CAUTION", GOLD, ["Relative iNKT1 enrichment in c0/c4/c6", "Mixed iNKT2/iNKT17 programs in c5", "Activation/proteostasis/mitochondrial remodeling", "Post-hoc response-program matches"]),
        ("NOT ESTABLISHED", RED, ["One-to-one legacy/current cluster identity", "c7 as a stable iNKT2 subtype", "Discrete iNKT17 subtype identity", "Lineage or population-level T2 effects"]),
    ]
    for idx, (label, color, bullets) in enumerate(columns):
        xx = 0.55 + idx * 4.13
        rounded_rect(ax, xx, 1.70, 3.84, 4.60, face=WHITE, edge=color, lw=1.4, radius=0.12)
        rounded_rect(ax, xx + 0.20, 1.96, 3.44, 0.48, face=color, edge=color, radius=0.18)
        add_text(ax, xx + 1.92, 2.20, label, size=9.4, color=WHITE, weight="bold", ha="center", va="center")
        bullet_list(ax, xx + 0.43, 2.83, bullets, width_chars=36, size=9.2, bullet_color=color, line_height=0.24, line_gap=0.24)
    return fig


def slide_35(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Reproducibility package and required claim language", section="Handoff", slide_no=n)
    metric_card(ax, 0.55, 1.68, 3.72, 1.40, "overall = 0", "All pipeline stages passed", accent=GREEN)
    metric_card(ax, 4.50, 1.68, 3.72, 1.40, "6 / 6", "Sample counts matched", accent=TEAL)
    metric_card(ax, 8.45, 1.68, 4.30, 1.40, "1 per group", "Biological samples per tissue", accent=RED, note="cell-level inference only")
    simple_table(
        ax,
        0.55,
        3.42,
        7.10,
        2.70,
        ["Artifact", "Location"],
        [
            ("Pipeline status", "pipeline_status.txt"),
            ("QC validation", "qc_validation.json"),
            ("Current processed object", "preprocess/inkt_scanpy_tutorial_processed.h5ad"),
            ("Legacy/current matches", "comparison_to_legacy_ppt/*.csv"),
            ("Curated ORA", "extended/de_pathway/overlap_pathway/*.csv"),
        ],
        col_widths=[2.15, 4.50],
        body_size=7.5,
    )
    callout(ax, 7.95, 3.42, 4.80, 1.25, "Required wording", "'Legacy QC was reconstructed exactly at the cell- and gene-count level.' Do not claim complete replication of the legacy downstream analysis.", accent=TEAL, face=TEAL_LIGHT, width_chars=59, body_size=8.0)
    callout(ax, 7.95, 4.90, 4.80, 1.22, "Next analyses before publication", "Recover biological replicate IDs; align full-gene global DE; obtain original GMT/library versions; run biologically distinct DPT root sensitivity analyses.", accent=ORANGE, face="#FFF5EC", width_chars=59, body_size=8.0)
    return fig


def slide_36(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Sensitivity: raw score_genes values are not a common ruler", section="Appendix", slide_no=n)
    add_image(ax, EXT / "signatures/signature_scores/cluster_subtype_evidence.png", 0.55, 1.62, 8.12, 2.95, trim=False, label="RAW SCORES")
    simple_table(
        ax,
        0.55,
        4.83,
        8.12,
        1.52,
        ["Source", "iNKT1 genes / ctrl", "iNKT2 genes / ctrl", "iNKT17 genes / ctrl"],
        [
            ("Literature", "25 / 50", "76 / 76", "30 / 50"),
            ("In-house", "35 / 50", "16 / 50", "15 / 50"),
            ("Wang 2022", "44 / 50", "36 / 50", "31 / 50"),
        ],
        col_widths=[1.65, 2.05, 2.05, 2.05],
        body_size=7.6,
        header_size=7.0,
    )
    callout(
        ax,
        8.94,
        1.62,
        3.78,
        1.66,
        "Why this moved to the appendix",
        "Gene-set sizes and matched-control references differ. Raw values answer a sensitivity question but cannot serve as calibrated subtype probabilities or a shared intensity scale.",
        accent=RED,
        face="#FFF0F0",
        width_chars=47,
    )
    callout(
        ax,
        8.94,
        3.58,
        3.78,
        1.36,
        "c7 is method-dependent",
        "The raw argmax favored iNKT2, whereas all three unified literature-list scores are below zero. The main text therefore leaves c7 unclassified.",
        accent=PURPLE,
        face="#F4F0FF",
        width_chars=47,
    )
    callout(
        ax,
        8.94,
        5.22,
        3.78,
        1.13,
        "Interpretation rule",
        "Use these panels only to assess source sensitivity; do not transfer the raw winner labels into biological subtype names.",
        accent=TEAL,
        face=TEAL_LIGHT,
        width_chars=47,
        body_size=8.0,
    )
    return fig


def slide_37(ctx: DeckContext, n: int):
    fig, ax = canvas()
    header(ax, "Sensitivity: subtype lists have limited cross-source overlap", section="Appendix", slide_no=n)
    add_image(ax, EXT / "signatures/gene_sets/signature_jaccard_heatmap.png", 0.55, 1.62, 6.00, 4.82, label="SET OVERLAP")
    add_image(ax, EXT / "signatures/gene_sets/subtype_signature_membership.png", 6.82, 1.62, 5.93, 4.82, label="MEMBERSHIP")
    return fig


SLIDES: list[tuple[str, Callable[[DeckContext, int], object]]] = [
    ("iNKT single-cell reanalysis", title_slide),
    ("Exact reconstruction of the legacy QC cohort", slide_02),
    ("HVG selection and PCA representation", slide_03),
    ("UMAP sensitivity across 50, 100, and 200 PCs", slide_04),
    ("UMAP PC selection and recommendation", slide_pc_selection),
    ("Eight current Leiden states and exploratory graph topology", slide_05),
    ("Unified iNKT1 enrichment is strongest in c4 and c6", slide_06),
    ("iNKT2 signal is strongest in c5; c7 is unclassified", slide_07),
    ("c5 is legacy-list iNKT17-enriched, but identity is incomplete", slide_08),
    ("One common scale replaces forced subtype winners", slide_09),
    ("Legacy global DE directions are retained within the tested HVG universe", slide_10),
    ("Tissue-stratified DE strongly recovers the legacy core", slide_11),
    ("Tissue identity dominates cluster composition", slide_12),
    ("Legacy c1 thymus response is retained in current c6", slide_13),
    ("All eight legacy c1 pathway drivers are retained in current c6", slide_14),
    ("Legacy c2 response is represented by current c4/c3", slide_15),
    ("Legacy c2 pathway drivers are recoverable but distributed", slide_16),
    ("Legacy c5 BM response is strongly recovered in current c0", slide_17),
    ("A conserved BM activation-proteostasis-mitochondrial response", slide_18),
    ("Legacy c6 spleen has a compact current c5 match", slide_19),
    ("The legacy c6 HSP/IEG module is largely retained", slide_20),
    ("Legacy c7 BM signal is recoverable but not uniquely mappable", slide_21),
    ("Legacy c8 spleen response is highly covered by current c3", slide_22),
    ("All 20 legacy c8 pathway drivers are retained in current c3", slide_23),
    ("Legacy c9 response is distributed across current states", slide_24),
    ("The legacy c9 chaperone axis is retained across states", slide_25),
    ("Detectable response depends on test status and cell count", slide_26),
    ("Legacy tissue DEG cores are recovered; anchors show balanced overlap", slide_27),
    ("Legacy response programs map to current tissue anchors", slide_28),
    ("Gene-level response hubs recur; the interaction network was not reproduced", slide_29),
    ("Tissue topology is robust; DPT is a root-versus-rest ordering", slide_30),
    ("Within-tissue T2 responses outperform tissue-specific markers", slide_31),
    ("A four-gene cross-tissue signal is retained; rare signals remain fragile", slide_32),
    ("Curated activation recurs; legacy terms were not recomputed", slide_33),
    ("What this reproduction supports", slide_34),
    ("Reproducibility package and required claim language", slide_35),
    ("Sensitivity: raw score_genes values are not a common ruler", slide_36),
    ("Sensitivity: subtype lists have limited cross-source overlap", slide_37),
]


def render(ctx: DeckContext) -> list[dict]:
    records: list[dict] = []
    with PdfPages(ctx.pdf_path, metadata={"Title": "iNKT legacy-QC detailed reproduction", "Author": "Codex"}) as pdf:
        for slide_no, (title, build) in enumerate(SLIDES, start=1):
            fig = build(ctx, slide_no)
            png_path = ctx.slides_dir / f"slide_{slide_no:02d}.png"
            fig.savefig(png_path, dpi=PNG_DPI, facecolor=fig.get_facecolor(), edgecolor="none")
            pdf.savefig(fig, facecolor=fig.get_facecolor(), edgecolor="none")
            plt.close(fig)
            records.append(
                {
                    "slide": slide_no,
                    "title": title,
                    "png": str(png_path.relative_to(ROOT)),
                    "png_sha256": sha256(png_path),
                }
            )
    return records


def build_pptx(ctx: DeckContext, records: Sequence[dict]) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W)
    presentation.slide_height = Inches(SLIDE_H)
    blank = presentation.slide_layouts[6]
    # Remove the default first slide only if a template happened to create one.
    while presentation.slides:
        slide_id = presentation.slides._sldIdLst[0]
        presentation.part.drop_rel(slide_id.rId)
        del presentation.slides._sldIdLst[0]
    for record in records:
        slide = presentation.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(ROOT / record["png"]),
            0,
            0,
            width=Inches(SLIDE_W),
            height=Inches(SLIDE_H),
        )
    presentation.core_properties.title = "iNKT legacy-QC detailed reproduction"
    presentation.core_properties.subject = "Exact count-level legacy QC reconstruction and downstream response-program reanalysis"
    presentation.core_properties.author = "Codex"
    presentation.save(ctx.pptx_path)


def write_manifest(ctx: DeckContext, records: Sequence[dict]) -> None:
    sweep_artifacts = ctx.sweep_summary["artifacts"]
    manifest = {
        "title": "iNKT legacy-QC detailed reproduction",
        "n_slides": len(records),
        "slide_size_inches": [SLIDE_W, SLIDE_H],
        "audited_run": manifest_path(RUN),
        "analysis_run": manifest_path(RUN),
        "embedding_run": manifest_path(EMBEDDING_RUN),
        "pc_sweep": {
            "directory": manifest_path(SWEEP),
            "summary": manifest_path(SWEEP / "summary.json"),
            "summary_sha256": sha256(SWEEP / "summary.json"),
            "candidate_pcs": ctx.sweep_summary["candidate_pcs"],
            "selected_pc": ctx.sweep_summary["selected_pc"],
            "selection_score": ctx.sweep_summary["selection_score"],
            "parameters": ctx.sweep_summary["parameters"],
            "cluster_boundary_policy": ctx.sweep_summary["cluster_boundary_policy"],
            "subtype_policy": ctx.sweep_summary["subtype_policy"],
            "trajectory_policy": ctx.sweep_summary["trajectory_policy"],
            "trajectory": ctx.sweep_summary["trajectory"],
            "selected_h5ad": sweep_artifacts["selected_h5ad"],
            "selected_h5ad_sha256": sweep_artifacts["selected_h5ad_sha256"],
        },
        "legacy_ppt": manifest_path(LEGACY_PPT),
        "content_spec": manifest_path(DOC),
        "pc_selection_report": manifest_path(ctx.out_dir / "UMAP_PC_selection_report.md"),
        "outputs": {
            "pdf": manifest_path(ctx.pdf_path),
            "pdf_sha256": sha256(ctx.pdf_path),
            "pptx": manifest_path(ctx.pptx_path),
            "pptx_sha256": sha256(ctx.pptx_path),
        },
        "claim_boundary": (
            "Exact reconstruction applies to post-QC cell/gene counts and all six sample counts; "
            "downstream embedding, clustering, DE, signatures, pathways, and trajectory are a current reanalysis."
        ),
        "statistical_warning": "One biological sample per tissue x condition; cell-level comparisons are exploratory.",
        "slides": list(records),
    }
    ctx.manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--analysis-run-dir",
        type=Path,
        default=DEFAULT_ANALYSIS_RUN,
        help="Complete audited run supplying DE, signatures, pathways, and comparison tables",
    )
    parser.add_argument(
        "--embedding-run-dir",
        type=Path,
        default=DEFAULT_EMBEDDING_RUN,
        help="Legacy-QC preprocess run supplying QC, HVG, and PCA inputs",
    )
    parser.add_argument(
        "--pc-sweep-dir",
        type=Path,
        default=DEFAULT_OUT_DIR / "umap_pc_sweep",
        help="Completed UMAP PC-sweep directory containing summary.json",
    )
    parser.add_argument(
        "--out-dir",
        type=Path,
        default=DEFAULT_OUT_DIR,
        help="Output directory (default: output/iNKT_reproduction_deck)",
    )
    args = parser.parse_args()
    configure_sources(args.analysis_run_dir, args.embedding_run_dir, args.pc_sweep_dir)
    validate_inputs()
    ctx = make_context(args.out_dir.resolve())
    records = render(ctx)
    build_pptx(ctx, records)
    write_manifest(ctx, records)
    print(f"slides={len(records)}")
    print(f"pdf={ctx.pdf_path}")
    print(f"pptx={ctx.pptx_path}")
    print(f"manifest={ctx.manifest_path}")


if __name__ == "__main__":
    main()
