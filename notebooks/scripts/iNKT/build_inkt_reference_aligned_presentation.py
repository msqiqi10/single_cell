"""Render existing iNKT analyses in the corresponding reference figure formats.
No differential expression, enrichment, embedding or clustering is recomputed.
"""
from __future__ import annotations
import os
for k in ('OMP_NUM_THREADS','OPENBLAS_NUM_THREADS','MKL_NUM_THREADS','NUMEXPR_NUM_THREADS'):
    os.environ[k]='1'
os.environ['CUDA_VISIBLE_DEVICES']=''
import hashlib,json,textwrap
from pathlib import Path
import numpy as np
import pandas as pd
import anndata as ad
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import Ellipse
from scipy.ndimage import distance_transform_edt
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader
from inkt_palette import SAMPLE_PALETTE,SAMPLE_DISPLAY_ORDER

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'output/iNKT_presentation_alignment_20260915'
M=ROOT/'output/iNKT_meeting_followup_20260905'
D=ROOT/'output/iNKT_discovery_20260905'
C=ROOT/'output/iNKT_reproduction_deck/20260830_C5_paper_Fig3DEF_followup'
R=ROOT/'output/iNKT_revision_responses_20260906'
ORDER=['C0','C1','C2','C3','C4','C5-1','C5-2','C6','C7']
KEY='cluster_c5_split_20260830'
BLUE='#12658c';RED='#c70039';GREY='#cecece'
TISSUES=['bone_marrow','spleen','thymus']
NAMES={'bone_marrow':'Bone marrow','spleen':'Spleen','thymus':'Thymus'}
COLORS=dict(zip(ORDER,['#31a8d3','#ee9920','#087a63','#9467bd','#e5b53c','#4485ab','#d56b73','#67a547','#787878']))
SOURCES={};PAGES=[];CHECKS=[];EXPORTS=[]
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':12,'axes.titlesize':15,'axes.labelsize':12,'xtick.labelsize':10,'ytick.labelsize':10,'pdf.fonttype':42,'svg.fonttype':'none','axes.spines.top':False,'axes.spines.right':False})

def record(p):
    p=Path(p)
    key=str(p.relative_to(ROOT))
    if key not in SOURCES:SOURCES[key]=hashlib.sha256(p.read_bytes()).hexdigest()
    return p

def read(p):return pd.read_csv(record(p),low_memory=False)
def save_table(d,name):
    p=OUT/'tables'/f'{name}.csv';d.to_csv(p,index=False);EXPORTS.append(str(p.relative_to(OUT)));return p

def check(value,label):
    if not value:raise AssertionError(label)
    CHECKS.append(label)

def sf(x):
    if pd.isna(x):return 'NA'
    if x==0:return '<machine limit'
    return f'{x:.2g}'

def short(uid):return uid.replace('cluster_tissue__','').replace('tissue__','').replace('__',' / ').replace('bone_marrow','BM').replace('spleen','Spl').replace('thymus','Thy')
def title(fig,title,subtitle,source,note):
    fig.text(.035,.95,title,fontsize=23,weight='bold',va='top')
    fig.text(.035,.897,subtitle,fontsize=12,va='top',color='#444444')
    fig.text(.035,.078,note,fontsize=10,va='top')
    fig.text(.035,.023,source,fontsize=8.5,color='#555555')
    fig.text(.965,.023,str(len(PAGES)+1),ha='right',fontsize=9)

def new(title_text,subtitle,source,note='Exploratory cell-level results; one sample label per tissue × condition.'):
    fig=plt.figure(figsize=(16,9),facecolor='white');title(fig,title_text,subtitle,source,note)
    fig._page_info={'title':title_text,'subtitle':subtitle,'reference':source,'note':note}
    return fig

def finish(fig,pdf):
    idx=len(PAGES)+1;png=OUT/'slides'/f'{idx:02}.png';svg=OUT/'slides'/f'{idx:02}.svg'
    fig.savefig(png,dpi=140);fig.savefig(svg);pdf.savefig(fig)
    PAGES.append(dict(page=idx,png=str(png.relative_to(OUT)),svg=str(svg.relative_to(OUT)),**fig._page_info))
    print(f'{idx:02} {fig._page_info["title"]}',flush=True);plt.close(fig)

def table(ax,headers,rows,widths=None,fontsize=11):
    ax.axis('off');t=ax.table(cellText=rows,colLabels=headers,cellLoc='left',colLoc='left',colWidths=widths,bbox=[0,0,1,1]);t.auto_set_font_size(False);t.set_fontsize(fontsize)
    for (r,c),cell in t.get_celld().items():
        cell.set_edgecolor('#dddddd');cell.set_linewidth(.5)
        if r==0:cell.set_facecolor('#ededed');cell.set_text_props(weight='bold')
        elif r%2==0:cell.set_facecolor('#f8f8f8')
    return t

def umap(ax,xy,mask=None,values=None,labels=None,cmap='viridis',vmin=None,vmax=None,size=2):
    if mask is None:mask=np.ones(len(xy),bool)
    ax.scatter(xy[:,0],xy[:,1],s=1,c='#e8e8e8',rasterized=True,linewidths=0)
    if values is not None:
        vals=np.asarray(values);idx=np.flatnonzero(mask);idx=idx[np.argsort(vals[idx])]
        im=ax.scatter(xy[idx,0],xy[idx,1],s=size,c=vals[idx],cmap=cmap,vmin=vmin,vmax=vmax,rasterized=True,linewidths=0)
    elif labels is not None:
        im=None
        for label,color in labels.items():
            m=mask & (np.asarray(label[1]) if isinstance(label,tuple) else np.zeros(len(xy),bool))
    else:im=ax.scatter(xy[mask,0],xy[mask,1],s=size,color=BLUE,rasterized=True,linewidths=0)
    ax.set_xticks([]);ax.set_yticks([]);ax.set_xlabel('UMAP1',fontsize=9);ax.set_ylabel('UMAP2',fontsize=9);ax.set_aspect('equal',adjustable='box')
    return im

def expr(a,g):
    x=a.raw[:,g].X
    return np.asarray(x.toarray() if hasattr(x,'toarray') else x).ravel()

def overview(pdf):
    fig=new('iNKT: results in the reference presentation formats','Existing analyses from July–September 2026 • presentation revision: 15 September 2026','References: original iNKT.pptx (33 slides); Kuznetsova et al., Blood Advances 2025, Figure 3 and Tables S3/S5.',
            'Same analytical question → same figure grammar. Different algorithms, populations and thresholds are stated explicitly.')
    rows=[['Sample/QC, markers, DEG and pathway pages','Original iNKT deck','Counts + plots; gene dotplots + numerical tables'],['Cluster composition','Blood Fig. 3C','Two 100% stacked-bar plots, each with its own denominator'],['Differential expression','Blood Fig. 3D','Volcano: log2FC vs −log10 nominal P; blue/red direction'],['Up-DEG pathway enrichment','Blood Fig. 3E','Single score column, −ln(P), with driver-gene callouts'],['Shared upregulated genes','Blood Fig. 3F','Four-set Venn; identical saved membership and exact counts'],['Signed GSEA; batch/state sensitivity','Additional analyses','Separate figures; NES is not the paper enrichment score']]
    table(fig.add_axes([.04,.28,.92,.52]),['Analysis','Reference','Presentation used here'],rows,[.29,.2,.51],12)
    fig.text(.05,.19,'Our cells are iNKT; the Blood paper studies NK cells. Current C# labels are not paper K# or legacy c# identities.',fontsize=14)
    finish(fig,pdf)

def qc_and_embedding(a,pdf):
    o=a.obs;xy=np.asarray(a.obsm['X_umap'])
    fig=new('Material, dataset and QC','Six samples • 15,532 retained cells • 10,670 measured genes','Original iNKT deck, slide 2; current object: 20260905 scored_base.h5ad.',
            'QC counts match the legacy deck; downstream UMAP, Leiden and gene statistics remain a reanalysis.')
    rows=[]
    raw={'Ctrl_Thymus':2101,'Ctrl_BM':3998,'Ctrl_Spleen':3829,'T2_Thymus':1639,'T2_BM':3131,'T2_Spleen':3760}
    for s in SAMPLE_DISPLAY_ORDER:rows.append([s,raw[s],int(o['sample'].astype(str).eq(s).sum())])
    table(fig.add_axes([.04,.3,.38,.5]),['Sample','Input cells','Retained'],rows,[.5,.25,.25],12)
    for j,(key,lab) in enumerate([('n_genes_by_counts','Detected genes'),('total_counts','UMI counts'),('pct_counts_mt','Mitochondrial %')]):
        ax=fig.add_axes([.49+j*.165,.34,.135,.4]);ax.violinplot(o[key].to_numpy(),showmedians=True);ax.set_xticks([]);ax.set_title(lab,fontsize=13)
    fig.text(.49,.24,'Gene min_cells = 100, then\n200 ≤ detected genes < 2,500; mitochondrial fraction < 5%.',fontsize=13)
    finish(fig,pdf)
    fig=new('Clusters and tissue composition','Fixed 50-PC UMAP; C5 refined to C5-1 and C5-2','Original iNKT slides 4–5/12; Blood Fig. 3B is t-SNE/K-means, so axes and labels are not copied.',
            'Original sample hues retained on the left; Blood-style blue/red encodes condition on the right. Coordinates are unchanged.')
    for j,kind in enumerate(['sample',KEY,'condition']):
        ax=fig.add_axes([.05+j*.325,.26,.28,.53]);umap(ax,xy,np.zeros(len(o),bool))
        palette=SAMPLE_PALETTE if kind=='sample' else COLORS if kind==KEY else {'Ctrl':BLUE,'T2':RED}
        for k,c in palette.items():
            m=o[kind].astype(str).eq(k).to_numpy();ax.scatter(xy[m,0],xy[m,1],s=2,color=c,label=k,rasterized=True,linewidths=0)
            if kind==KEY and m.any():
                pos=np.median(xy[m],axis=0);ax.text(*pos,k,fontsize=10,weight='bold',bbox=dict(facecolor='white',alpha=.75,edgecolor='none',pad=1))
        ax.set_title(['Sample','Current clusters','Condition'][j]);ax.legend(loc='upper center',bbox_to_anchor=(.5,-.1),ncol=3,fontsize=8,markerscale=3,frameon=False)
    finish(fig,pdf)

def marker_pages(a,pdf):
    xy=np.asarray(a.obsm['X_umap']);o=a.obs;coverage=read(M/'tables/state_signature_coverage.csv').set_index('signature')
    for sig in ['iNKT1','iNKT2','iNKT17']:
        genes=str(coverage.loc[sig,'genes']).split(';');vals=o[sig].to_numpy()
        fig=new(f'{sig} marker program','Score map, cluster distributions, individual markers and expression dotplot','Original iNKT slides 6–8; marker panel and saved continuous scores from the 20260905 state analysis.',
                'Current score_genes units differ from legacy summed-marker scores; these panels do not establish discrete subtype identities.')
        ax=fig.add_axes([.035,.49,.23,.32]);im=umap(ax,xy,values=vals,vmin=np.min(vals),vmax=np.max(vals));ax.set_title(f'{sig} score',fontsize=14);fig.colorbar(im,ax=ax,fraction=.04,pad=.025)
        ax=fig.add_axes([.045,.22,.23,.18]);data=[vals[o[KEY].astype(str).eq(c)] for c in ORDER];ax.violinplot(data,showmedians=True,showextrema=False);ax.set_xticks(range(1,10),ORDER,rotation=55);ax.set_ylabel('Module score',fontsize=10)
        for j,g in enumerate(genes):
            x=expr(a,g);ax=fig.add_axes([.31+(j%3)*.125,.56-(j//3)*.29,.115,.22]);im=umap(ax,xy,values=x,vmin=0,vmax=max(float(x.max()),.01));ax.set_title(g,fontsize=12);ax.set_xlabel('');ax.set_ylabel('');fig.colorbar(im,ax=ax,fraction=.05,pad=.02).ax.tick_params(labelsize=7)
        ax=fig.add_axes([.76,.37,.205,.36]);rows=[]
        for iy,g in enumerate(genes):
            x=expr(a,g)
            for ix,c in enumerate(ORDER):
                z=x[o[KEY].astype(str).eq(c).to_numpy()];rows.append({'gene':g,'cluster':c,'fraction':float(np.mean(z>0)),'mean_log1p':float(np.mean(z))})
        d=pd.DataFrame(rows);lim=max(float(d.mean_log1p.max()),.01)
        for row in d.itertuples():im=ax.scatter(ORDER.index(row.cluster),genes.index(row.gene),s=220*row.fraction,c=[row.mean_log1p],vmin=0,vmax=lim,cmap='viridis',edgecolor='none')
        ax.set_yticks(range(len(genes)),genes);ax.set_xticks(range(9),ORDER,rotation=55);ax.invert_yaxis();ax.set_title('Marker expression',fontsize=14);cax=fig.add_axes([.78,.275,.165,.012]);fig.colorbar(im,cax=cax,orientation='horizontal',label='Mean log1p (all cells)');cax.tick_params(labelsize=8)
        for n in [.1,.5,1]:ax.scatter([],[],s=220*n,c='#777777',label=f'{n:.0%}')
        ax.legend(title='Fraction detected',loc='upper center',bbox_to_anchor=(.5,-.45),ncol=3,fontsize=8,frameon=False)
        if sig=='iNKT17':fig.text(.32,.19,'Ccr6 is absent from the filtered feature set.\nMissing markers are not shown as zero expression.',fontsize=11)
        save_table(d,f'marker_{sig}');finish(fig,pdf)

def frequency_pages(pdf):
    d=read(M/'tables/frequency_reconciliation_1.csv')
    for tissue in TISSUES:
        q=d[d.tissue.eq(tissue)].set_index('cluster').reindex(ORDER)
        fig=new(f'{NAMES[tissue]}: cluster and condition frequencies','C  •  The two denominators from Blood Figure 3C, presented side by side','Blood Fig. 3C; saved 20260905 frequency_reconciliation_1.csv.',
                'Captured-cell proportions only. No mouse-level error bars or significance are inferred from individual cells.')
        left=fig.add_axes([.08,.3,.30,.48]);right=fig.add_axes([.49,.3,.45,.48]);bot=np.zeros(2)
        for c in ORDER:
            val=q.loc[c,['control_cluster_frequency_pct','tumor_cluster_frequency_pct']].to_numpy(float);left.bar([0,1],val,bottom=bot,color=COLORS[c],label=c,width=.58);bot+=val
        check(np.allclose(bot,100),f'{tissue} condition-normalized stacks sum to 100')
        totals=[int(q.control_count.sum()),int(q.tumor_count.sum())]
        left.set_xticks([0,1],[f'Control\nn={totals[0]:,}',f'Tumor (T2)\nn={totals[1]:,}']);left.set_ylim(0,108);left.set_ylabel('% of iNKT cells in condition');left.set_title('Cluster frequency within condition')
        for i in range(2):left.text(i,103,f'n={totals[i]:,}',ha='center',fontsize=10)
        left.legend(ncol=5,loc='upper center',bbox_to_anchor=(.5,-.23),fontsize=8,frameon=False)
        ctrl=q.control_share_within_tissue_cluster_pct.to_numpy(float);tum=q.tumor_share_within_tissue_cluster_pct.to_numpy(float);valid=q.control_count.add(q.tumor_count).to_numpy()>0
        check(np.allclose(ctrl[valid]+tum[valid],100),f'{tissue} cluster-normalized stacks sum to 100')
        right.bar(np.arange(9),ctrl,color=BLUE,label='Control',width=.64);right.bar(np.arange(9),tum,bottom=ctrl,color=RED,label='Tumor (T2)',width=.64)
        for i,c in enumerate(ORDER):
            total=int(q.loc[c,'control_count']+q.loc[c,'tumor_count']);right.text(i,102,str(total) if total else 'NA',ha='center',fontsize=9)
        right.set_xticks(np.arange(9),ORDER,rotation=35);right.set_ylim(0,112);right.set_ylabel('% of cells in cluster');right.set_title('Condition share within cluster');right.legend(loc='upper center',bbox_to_anchor=(.5,-.2),ncol=2,frameon=False,fontsize=10)
        fig.text(.07,.115,'Left: n(cluster, condition) / N(condition)     |     Right: n(condition, cluster) / N(cluster)',fontsize=13)
        save_table(q.reset_index(),f'Fig3C_{tissue}');finish(fig,pdf)

def selected_de(d,n=14):
    sig=d[(d.pvals_adj<=.05)&(d.logfoldchanges.abs()>=.25)].copy()
    up=sig[sig.logfoldchanges>0].sort_values(['pvals_adj','gene']).head(n//2)
    down=sig[sig.logfoldchanges<0].sort_values(['pvals_adj','gene']).head(n//2)
    return pd.concat([up,down]).drop_duplicates('gene')

def gene_page(a,uid,pdf,reference):
    d=read(M/'de'/f'{uid}.csv.gz');sel=selected_de(d);n=len(sel);check(n>0,f'DEG display nonempty {uid}')
    global_de=read(M/'de/global.csv.gz').set_index('gene');sel=sel.copy();sel['passes_global_same_rule']=[bool(global_de.loc[g,'pvals_adj']<=.05 and abs(global_de.loc[g,'logfoldchanges'])>=.25) for g in sel.gene]
    fig=new(f'T2 vs. Control DEGs: {short(uid)}','Expression/detection dotplot alongside the exact same genes in a signed DEG table',reference,
            'Displayed: 7 strongest FDR-ranked genes per direction, when available; full gene table linked in the source index.')
    ax=fig.add_axes([.055,.31,.23,.47]);xy=np.asarray(a.obsm['X_umap']);m=np.ones(a.n_obs,bool)
    if uid.startswith('tissue__'):m=a.obs.tissue.astype(str).eq(uid.split('__')[1]).to_numpy()
    if uid.startswith('cluster_tissue__'):
        _,c,t=uid.split('__');m=(a.obs.tissue.astype(str).eq(t)&a.obs[KEY].astype(str).eq(c)).to_numpy()
    umap(ax,xy,np.zeros(a.n_obs,bool))
    for cond,col in [('Ctrl',BLUE),('T2',RED)]:
        z=m&a.obs.condition.astype(str).eq(cond).to_numpy();ax.scatter(xy[z,0],xy[z,1],color=col,s=4,rasterized=True,linewidths=0,label=f'{cond} (n={z.sum():,})')
    ax.legend(loc='upper center',bbox_to_anchor=(.5,-.15),fontsize=9,frameon=False);ax.set_title('Compared cells',fontsize=14)
    ax=fig.add_axes([.365,.24,.12,.56]);genes=sel.gene.tolist();vmax=float(sel[['mean_log1p_expression_control','mean_log1p_expression_tumor']].to_numpy().max())
    for y,row in enumerate(sel.itertuples()):
        for x,cond in enumerate(['control','tumor']):
            im=ax.scatter(x,y,s=220*getattr(row,'pct_expressing_'+cond),c=[getattr(row,'mean_log1p_expression_'+cond)],cmap='viridis',vmin=0,vmax=vmax,edgecolor='none')
    ax.set_xticks([0,1],['Ctrl','T2']);ax.set_yticks(range(n),[g+(' *' if not bool(sel.iloc[i].passes_global_same_rule) else '') for i,g in enumerate(genes)]);ax.set_ylim(n-.5,-.5);ax.set_xlim(-.55,1.55);ax.set_title('Expression',fontsize=13)
    cax=fig.add_axes([.365,.175,.12,.014]);fig.colorbar(im,cax=cax,orientation='horizontal',label='Mean log1p (all cells)');cax.tick_params(labelsize=8)
    for z in [.1,.5,1]:ax.scatter([],[],s=220*z,color='#777777',label=f'{z:.0%}')
    ax.legend(loc='upper left',bbox_to_anchor=(1.05,1.09),ncol=3,fontsize=7,title='Fraction detected',title_fontsize=8,frameon=False)
    rows=[[r.gene,f'{r.scores:+.2f}',f'{r.logfoldchanges:+.3f}',sf(r.pvals),sf(r.pvals_adj)] for r in sel.itertuples()]
    table(fig.add_axes([.55,.24,.41,.56]),['Gene','Score','log2FC','P','BH FDR'],rows,[.24,.16,.18,.21,.21],10.5)
    fig.text(.055,.115,'Selection: BH FDR ≤ 0.05, |log2FC| ≥ 0.25.  * Not passing the same rule in the global comparison; not an interaction test.',fontsize=10.5)
    save_table(sel,f'DEG_{uid}');finish(fig,pdf)

def legacy_pathway_pages(pdf):
    d=read(M/'tables/PPT_pathway_validation_every_row.csv')
    specs=[('c1_thymus','cluster_tissue__C6__thymus',14),('c5_bone_marrow','cluster_tissue__C0__bone_marrow',18),('c6_spleen','cluster_tissue__C5-1__spleen',20)]
    for old,uid,slide in specs:
        q=d[d.legacy_unit.eq(old)&d.current_unit.eq(uid)&d.sensitivity.eq('full')].copy()
        q['best_current_fdr']=q[['robust_Tumor_fdr','robust_Control_fdr']].min(axis=1);q=q.sort_values(['best_current_fdr','ppt_row']);show=q.head(8)
        fig=new(f'Pathways: {short(uid)}',f'Legacy {old} pathway vocabulary; current full-member ORA with explicit up/down direction',f'Original iNKT slide {slide}; 20260905 PPT_pathway_validation_every_row.csv.',
                'Legacy/current cluster correspondence is a response comparison, not an identity claim; pathway versions and thresholds differ.')
        rows=[]
        for r in show.itertuples():
            def gene_preview(value):
                genes=str(value).split(';') if pd.notna(value) else []
                return ', '.join(genes[:3]) + (f' +{len(genes)-3} more' if len(genes)>3 else '') if genes else 'none'
            drivers='T2: '+gene_preview(r.robust_Tumor_genes)+'\nCtrl: '+gene_preview(r.robust_Control_genes)
            name=r.resolved_name if pd.notna(r.resolved_name) else r.legacy_name
            rows.append([textwrap.fill(str(name),38),str(r.legacy_source).replace('_','\n'),sf(r.legacy_pFDR),sf(r.robust_Tumor_fdr),sf(r.robust_Control_fdr),drivers])
        table(fig.add_axes([.035,.22,.93,.59]),['Pathway / reaction','Legacy source','Old FDR','T2-up FDR','Ctrl-up FDR','Current overlapping genes'],rows,[.29,.13,.09,.09,.09,.31],9.5)
        fig.text(.04,.15,'Shown: 8 strongest current min(up-FDR, down-FDR), retaining original names; first 3 drivers per direction shown; complete records in exported table.',fontsize=10)
        save_table(q,f'legacy_pathways_{old}');finish(fig,pdf)

def volcano_page(pdf):
    units=['cluster_tissue__C0__bone_marrow','cluster_tissue__C6__thymus','cluster_tissue__C5-1__spleen','cluster_tissue__C5-2__bone_marrow']
    fig=new('D  Differential expression: Control vs. Tumor','Paper axes and blue/red direction; current tissue × cluster comparisons','Blood Fig. 3D; current DE tables (20260905). Paper combines K1+K2; these panels do not claim that population equivalence.',
            'Color rule retained: gene BH FDR ≤ 0.05 and |log2FC| ≥ 0.25. Vertical display thresholds are therefore ±0.25.')
    for j,uid in enumerate(units):
        d=read(M/'de'/f'{uid}.csv.gz');x=d.logfoldchanges.to_numpy();rawy=-np.log10(d.pvals.clip(lower=np.nextafter(0.,1.)).to_numpy());y=np.minimum(rawy,100)
        color=np.where((d.pvals_adj<=.05)&(x>=.25),RED,np.where((d.pvals_adj<=.05)&(x<=-.25),BLUE,'#c5c5c5'))
        ax=fig.add_axes([.075+(j%2)*.48,.57-(j//2)*.36,.4,.25]);ax.scatter(x,y,s=4,c=color,linewidths=0,rasterized=True);ax.axvline(-.25,ls=':',color='#555555');ax.axvline(.25,ls=':',color='#555555');ax.axhline(-np.log10(.05),ls=':',color='#555555')
        ax.set_xlim(-5,5);ax.set_ylim(0,107);ax.set_xlabel('log2(T2 / Ctrl)');ax.set_ylabel('−log10(nominal P)');ax.set_title(short(uid),fontsize=13)
        for side in [-1,1]:
            s=d[(d.logfoldchanges*side>.25)&(d.pvals_adj<=.05)].sort_values('pvals_adj').head(3)
            for k,r in enumerate(s.itertuples()):
                ax.annotate(r.gene,xy=(np.clip(r.logfoldchanges,-5,5),min(-np.log10(max(r.pvals,np.nextafter(0.,1.))),100)),xytext=(side*(2.4+.2*k),91-17*k),fontsize=9,ha='center',arrowprops={'arrowstyle':'-','lw':.6})
        ax.text(.03,.94,'Control',color=BLUE,transform=ax.transAxes,fontsize=9);ax.text(.97,.94,'Tumor',color=RED,transform=ax.transAxes,ha='right',fontsize=9)
        q=d.copy();q['display_y']=y;q['display_y_capped']=rawy>100;q['display_x_outside']=abs(x)>5;save_table(q,f'Fig3D_{uid}')
    fig.text(.075,.13,'Common axes: |log2FC| ≤ 5; −log10(P) capped at 100. All uncapped values and display flags are exported.',fontsize=10.5)
    finish(fig,pdf)

def enrichment_pages(pdf):
    s3=read(M/'sources/Blood_TableS3.csv');err=float(np.max(abs(s3['Enrichment score']+np.log(s3['P-value']))));check(err<.0001,'all 48 paper enrichment scores match -ln(P) to published rounding')
    d=read(M/'tables/Blood_TableS3_all48_pathway_validation.csv');paper=s3.set_index('Description');terms=s3[s3.Description.isin(d[d.in_Fig3E].paper_term.unique())].Description.tolist();check(len(terms)==17,'all 17 Figure3E terms retained in paper order')
    specs=['cluster_tissue__C0__bone_marrow','cluster_tissue__C5-1__spleen','cluster_tissue__C6__thymus']
    for uid in specs:
        q=d[d.unit_id.eq(uid)&d.sensitivity.eq('full')].set_index('paper_term').reindex(terms).copy();q['paper_score']=paper.loc[terms,'Enrichment score'];q['current_score']=-np.log(q.Tumor_ORA_p.clip(lower=np.nextafter(0.,1.)))
        fig=new(f'E  Paper pathways: {short(uid)}','Single enrichment-score columns with actual driver genes; all 17 reference pathways remain visible','Blood Fig. 3E/Table S3; our ORA: nominal DEG P ≤ 0.05, linear FC ≥ 1.5; frozen KEGG_2019_Mouse.',
                'Score = −ln(nominal enrichment P), verified against Table S3; * pathway BH FDR ≤ 0.05. Scale does not remove method/library differences.')
        ax=fig.add_axes([.055,.19,.60,.62]);ax.set_xlim(0,1);ax.set_ylim(16.6,-.6);ax.axis('off')
        cm=plt.get_cmap('Reds').copy();cm.set_bad('#dddddd');vmax=max(24,float(q.current_score.max()))
        for i,(term,r) in enumerate(q.iterrows()):
            ax.text(.58,i,term,ha='right',va='center',fontsize=11)
            for x,col,score,fdr in [(.62,'paper',r.paper_score,r.paper_fdr),(.77,'current',r.current_score,r.Tumor_ORA_fdr)]:
                ax.add_patch(plt.Rectangle((x-.023,i-.46),.046,.92,color=cm(score/vmax) if np.isfinite(score) else '#dddddd',ec='white',lw=.4))
                ax.text(x+.035,i,('* ' if pd.notna(fdr) and fdr<=.05 else '')+('NA' if not np.isfinite(score) else f'{score:.1f}'),fontsize=8,va='center')
        ax.text(.62,-1.05,'Paper NK',ha='center',fontsize=11);ax.text(.77,-1.05,'Our iNKT',ha='center',fontsize=11)
        ax2=fig.add_axes([.70,.19,.265,.62]);ax2.set_ylim(16.6,-.6);ax2.set_xlim(0,1);ax2.axis('off');ax2.text(0,-1.05,'Our T2-up overlapping genes',fontsize=11)
        for i,r in enumerate(q.itertuples()):
            g=str(r.Tumor_ORA_drivers) if pd.notna(r.Tumor_ORA_drivers) else '—'
            if len(g)>66:g=g[:63]+'…'
            ax2.text(0,i,g.replace(';',', '),fontsize=8.5,color=RED if g!='—' else '#888888',va='center')
        cax=fig.add_axes([.60,.12,.20,.012]);fig.colorbar(plt.cm.ScalarMappable(norm=plt.Normalize(0,vmax),cmap=cm),cax=cax,orientation='horizontal',label='Enrichment score −ln(P)');cax.tick_params(labelsize=8)
        save_table(q.reset_index(),f'Fig3E_reference_{uid}');finish(fig,pdf)
    for uid in specs:
        d=read(M/'enrichment'/f'{uid}__full__KEGG_Mouse_2019__ORA_paper.csv.gz');q=d[d.direction.eq('Tumor')].sort_values(['fdr','pvalue','term']).head(10).copy();q['score']=-np.log(q.pvalue.clip(lower=np.nextafter(0.,1.)))
        fig=new(f'E  Current top pathways: {short(uid)}','Full-library top 10 T2-up ORA pathways, using the same score-column and gene-callout format','Blood Fig. 3E presentation adapted to current data; exact tested KEGG terms retained.',
                'Top 10 by pathway BH FDR, then nominal P and term name; rows without * are not FDR-significant. No post-hoc pathway renaming.')
        ax=fig.add_axes([.04,.22,.57,.57]);ax.set_xlim(0,1);ax.set_ylim(9.6,-.6);ax.axis('off');vmax=max(float(q.score.max()),1)
        for i,r in enumerate(q.itertuples()):
            ax.text(.78,i,textwrap.fill(r.term,46),ha='right',va='center',fontsize=11)
            ax.add_patch(plt.Rectangle((.81,i-.45),.045,.9,color=plt.cm.Reds(r.score/vmax),ec='white'))
            ax.text(.87,i,('* ' if r.fdr<=.05 else '')+f'{r.score:.1f}',fontsize=10,va='center')
        ax2=fig.add_axes([.62,.22,.34,.57]);ax2.axis('off');ax2.set_xlim(0,1);ax2.set_ylim(9.6,-.6)
        for i,r in enumerate(q.itertuples()):ax2.text(0,i,textwrap.fill(str(r.overlap_genes).replace(';',', '),57),fontsize=10,va='center',color=RED)
        ax.text(.83,-1,'−ln(P)',ha='center',fontsize=12);ax2.text(0,-1,'T2-up driver genes',fontsize=12)
        save_table(q,f'Fig3E_top10_{uid}');finish(fig,pdf)

def venn_geometry():
    y,x=np.mgrid[0:1:600j,0:1:600j];params=[(cx,cy,.72,.4,a) for (cx,cy),a in zip([(.35,.4),(.45,.5),(.55,.5),(.65,.4)],[135,135,45,45])];regions=np.zeros(x.shape,int)
    for j,(cx,cy,w,h,a) in enumerate(params):
        t=np.deg2rad(a);u=(x-cx)*np.cos(t)+(y-cy)*np.sin(t);v=-(x-cx)*np.sin(t)+(y-cy)*np.cos(t);regions+=((u/(w/2))**2+(v/(h/2))**2<=1)*(2**j)
    coords={}
    for mask in range(1,16):
        check(np.any(regions==mask),f'four-ellipse exact region {mask} exists');pos=np.unravel_index(np.argmax(distance_transform_edt(regions==mask)),regions.shape);coords[mask]=(x[pos],y[pos])
    return params,coords

def venn_pages(pdf):
    params,coords=venn_geometry();colors=['#5966a3','#e3bd45','#68b95f','#cf5c56'];sets=['C5-1__bone_marrow','C5-1__spleen','C5-2__bone_marrow','C5-2__spleen']
    fig=new('F  Shared T2-up genes across state × tissue','Four-set Venn diagrams with exact counts from the existing membership tables','Blood Fig. 3F; original C5 overlap memberships from 20260830; geometry is schematic, not area-proportional.',
            'All-four intersection = 0 in both rules. This does not establish absence of shared biology; set membership depends on thresholds.')
    for j,rule in enumerate(['paper','robust']):
        d=read(C/'tables'/f'20260830_Fig3F_C5split_BM_Spleen_membership_{rule}.csv');mask=sum(d[s].astype(int)*(2**i) for i,s in enumerate(sets));counts=mask.value_counts().to_dict();check(sum(counts.values())==len(d),f'{rule} Venn counts partition union')
        ax=fig.add_axes([.025+j*.49,.24,.47,.53]);ax.set_aspect('equal');ax.set_xlim(-.05,1.05);ax.set_ylim(0,1);ax.axis('off')
        for k,p in enumerate(params):ax.add_patch(Ellipse(p[:2],p[2],p[3],angle=p[4],facecolor=colors[k],edgecolor='#333333',alpha=.4,lw=1.3))
        for b,(x,y) in coords.items():ax.text(x,y,str(counts.get(b,0)),ha='center',va='center',fontsize=12,weight='bold' if b==15 else 'normal')
        for k,(x,y) in enumerate([(.15,.12),(.23,.83),(.77,.83),(.85,.12)]):
            ax.text(x,y,sets[k].replace('__',' / ').replace('bone_marrow','BM').replace('spleen','Spl')+f'\n{int(d[sets[k]].sum())} genes',ha='center',va='center',color=colors[k],fontsize=11,weight='bold')
        ax.set_title('Paper rule: P < 0.05, log2FC ≥ 1' if rule=='paper' else 'Sensitivity: BH FDR ≤ 0.05, log2FC ≥ 0.25',fontsize=13)
        export=pd.DataFrame([{'bitmask':b,'sets':' & '.join(sets[k] for k in range(4) if b&(2**k)),'count':counts.get(b,0),'label_x':coords[b][0],'label_y':coords[b][1]} for b in range(1,16)]);save_table(export,f'Fig3F_exact_regions_{rule}')
        original=read(C/'tables'/f'20260830_Fig3F_C5split_BM_Spleen_summary_{rule}.csv');check(int(original[original.metric.eq('all_four')].n_genes.iloc[0])==counts.get(15,0),f'{rule} all-four agrees with original summary')
        for k,s in enumerate(sets):check(sum(counts.get(b,0) for b in range(1,16) if b&(2**k))==int(d[s].sum()),f'{rule} Venn reconstructs {s}')
    fig.text(.04,.18,'Cell counts (Ctrl/T2): C5-1 BM 124/85; C5-1 Spl 239/345; C5-2 BM 104/96; C5-2 Spl 28/42.',fontsize=12)
    fig.text(.04,.13,'Paper K1/K2 and our C5-1/C5-2 are different populations. Thymus is excluded here because C5-2 Ctrl n=11.',fontsize=11)
    finish(fig,pdf)

def image_page(path,title_text,subtitle,reference,pdf,note='Additional analysis; retain its original metric and interpretation.'):
    fig=new(title_text,subtitle,reference,note);ax=fig.add_axes([.035,.13,.93,.72]);ax.imshow(Image.open(record(path)));ax.axis('off');finish(fig,pdf)

def integration_page(pdf):
    fig=new('Batch sensitivity: show the actual UMAPs','Original graph, Harmony by tissue and Harmony by sample','R01 from the 20260906 report; embedding objects and metrics generated on 20260905.',
            'Integration increases mixing but changes condition/cluster structure. An independently identified technical batch is unavailable.')
    for i,name in enumerate(['original','harmony_tissue','harmony_sample']):
        p=record(M/'objects'/f'{name}.h5ad');a=ad.read_h5ad(p,backed='r');xy=np.asarray(a.obsm['X_umap']);o=a.obs.copy();a.file.close()
        ax=fig.add_axes([.045+i*.325,.3,.28,.5]);umap(ax,xy,np.zeros(len(o),bool))
        for cond,color in [('Ctrl',BLUE),('T2',RED)]:
            m=o.condition.astype(str).eq(cond).to_numpy();ax.scatter(xy[m,0],xy[m,1],s=2,color=color,label=cond,rasterized=True,linewidths=0)
        ax.set_title(name.replace('_',' ').title());ax.legend(loc='upper center',bbox_to_anchor=(.5,-.12),ncol=2,fontsize=10,frameon=False)
    fig.text(.045,.16,'All panels colored by the same condition labels; UMAP coordinates differ between fits. Expression DE still uses uncorrected counts.',fontsize=12)
    finish(fig,pdf)

def final_status(pdf):
    fig=new('What is comparable, and what remains different','Presentation changes preserve the measured quantity and the population being compared','Original iNKT deck and Blood main/supplementary figures reviewed; see presentation_audit.md for the full mapping.',
            'Current priority: BM C4 expression-state candidates and Spl C3 CCT/TriC-related expression; independent validation remains pending.')
    rows=[['Matched presentation','Dotplot + DEG table; two composition denominators; volcano axes; ORA score strip; exact four-set Venn.'],['Different computation','Current UMAP/Leiden vs paper t-SNE/K-means; current module scores vs legacy summed scores.'],['Different pathway statistics','ORA uses −ln(P); signed GSEA uses NES. Displaying a term does not imply BH FDR significance.'],['Different populations','iNKT vs paper NK; C# vs legacy c# or K#. No one-to-one identity assumed.'],['Not performed / unsupported','RNA velocity, paper qPCR/flow/function assays, PAGER interaction network replication; no replacement figures invented.'],['Current state of completion','R01–R13 remains as documented: velocity input gap, incomplete state coverage, old-library limitations and no animal-level replication.']]
    table(fig.add_axes([.045,.25,.91,.55]),['Category','Interpretation'],[[x,textwrap.fill(y,100)] for x,y in rows],[.23,.77],12)
    finish(fig,pdf)

def write_audit():
    text='''# Reference presentation audit and fixes — 2026-09-15

The original `input/iNKT/iNKT.pptx` (33 slides), the supplied Blood Advances paper (Figure 3 visually, main figure legends and relevant results text), and its local supplementary methods/Table S3 were reviewed. This revision uses existing saved analyses and redraws their presentation; no DE, enrichment, clustering, embeddings or velocity was rerun.

## Main findings and corrections

| Reference | Reference format and question | Previous mismatch | Revised presentation |
|---|---|---|---|
| iNKT slides 2–5 | Input/QC, dimensionality reduction and clusters | Results spread across narrative reports | Sample count table, QC distributions, PC comparison and tissue/cluster maps |
| iNKT slides 6–9 | Score UMAP + violin + per-gene UMAP + gene×cluster dotplot | Signature views and summaries separated | Same panel family for current iNKT1/2/17 scores; actual score method and marker coverage stated |
| iNKT slides 10–11 | Global/tissue DEG dotplot beside signed gene table | Response bars or selected dotplot summary could replace original display | Gene-row Ctrl/T2 dotplots paired with the same genes' Wilcoxon scores, log2FC, P and BH FDR |
| iNKT slides 12–26 | Highlighted cluster/tissue, DEG table/dotplot, followed by pathway table | Crosswalk summaries and broad heatmaps dominated | Explicit compared-cell UMAP and gene panels; source-named pathway table with up/down FDR and actual drivers |
| iNKT slides 27–28 | Local vs global/tissue information | Significance-only statements risked implying an interaction | Global-pass flag in exported DEG table; existing discovery effect panels retained with their limits |
| iNKT slide 29 | PAGER DEG interaction network | Not the same as pathway overlap or PAGA | Marked not reproduced; no substitute network asserted |
| iNKT slides 30–32 | Trajectory/marker interpretation | DPT or median-shift arrows could be confused with velocity | Existing continuous state maps retained as supplementary analyses; no velocity claim |
| iNKT slide 33 | Marker-linked PAGER/Reactome records | No exact same-library replication | Not declared reproduced; current complete-membership ORA remains explicitly separate |
| Blood Fig. 3A/B | Sorted NK, t-SNE, K-means and marker colors | Our UMAP/Leiden is a different calculation | Retain UMAP and current C labels; do not relabel axes or invent NK identities |
| Blood Fig. 3C | Cluster fractions within condition and condition fractions within cluster | Different denominators scattered across pages | Paired 100% stacked bars per tissue, shared 0–100 logic, actual n values, stable blue/red condition palette |
| Blood Fig. 3D | log2FC vs −log10 nominal P volcano; K1+K2 pooled | Scope/threshold differences were easy to miss | Same plot axes and directional colors, explicit tissue×cluster comparisons; original project FDR/FC color rule retained and labeled |
| Blood Fig. 3E | Single enrichment-score column with driver-gene callouts, up-DEG enrichment | Multi-cluster signed GSEA NES heatmap was used as the main analogue | ORA score strips with −ln(P), original 17-term order and current genes; additional full-library top-10 pages; GSEA remains separate |
| Blood Fig. 3F | Four-set state×tissue Venn, up genes at FC≥2 and nominal P<.05 | Main display used UpSet | Four-ellipse Venn with all 15 exact regions from saved memberships; paper-rule and FDR sensitivity shown separately |
| Blood Fig. 3G/H | qPCR and degranulation module/functional context | Different assays cannot become equivalent by matching graphics | Not claimed as replicated; current marker scores are not qPCR, killing or degranulation measurements |
| Blood other main figures | Flow/functional/perturbation, external cohorts and signaling follow-up | Those experiments were not performed on this iNKT dataset | No corresponding experimental evidence or significance bars manufactured |
| Current R01–R13/discovery | Additional questions requested in later meetings | Some main pages relegated relevant UMAPs to attachments | Actual three-condition embedding comparison included; separate GSEA/state/discovery pages retained |

## Quantitative checks

- All 48 Table S3 `Enrichment score` values equal `−ln(P-value)` to within 0.0001 (published rounding). This is a direct numerical inference from the provided table. It is not NES or −log10(FDR).
- Our matching score is computed only by transforming the already saved nominal ORA P value. The original full-library BH FDR stays attached; star = FDR≤0.05. No BH correction is recomputed on displayed subsets.
- All 17 Figure 3E terms are displayed regardless of significance. The additional top-10 display is deterministically ordered by saved FDR, P and term name. Full source tables retain everything not displayed.
- Our ORA input threshold is the supplement's nominal P≤0.05 and linear FC≥1.5. Figure 3F uses the paper's separate FC≥2 and P<.05 rule.
- Paper uses Partek gene-specific analysis and its pathway library; local analysis uses Wilcoxon and frozen KEGG_2019_Mouse. Shared graphics/score scales do not mean identical analysis or comparable cell identities.
- Venn membership is taken from 20260830; it is not silently mixed with 20260905 DE. All 15 exact counts reconstruct each set and the union. Areas are schematic. All-four=0 is shown honestly.
- Both stacked-bar denominators are verified against the original saved table, with zero-cell groups labeled NA.
- Dotplot area represents detected fraction; color is mean log1p over all cells in the displayed group. Numerical tables show the exact same genes. Current DE display rule is FDR≤0.05 and |log2FC|≥0.25, not the legacy 1.2 cutoff.
- Volcano limits and capping are explicitly reported; uncapped source values and display flags are exported.

## Delivery/version policy

This is a new dated presentation built from the latest analyses plus explicitly labeled 20260830 overlap data. Earlier dated reports and their provenance hashes are preserved. PDF has vector plot/text elements; PPTX contains matching rendered slide images (not editable chart data); individual SVG/PNG figures, display tables, source hashes and the builder are supplied. The 20260906 R01–R13 report remains the complete item-by-item supplement.

Source-layout previews in `reference_audit/legacy_*` are picture/text-position inspection aids generated from PowerPoint XML, not an Office rendering; native source tables were inspected directly. They are not used as analytical result figures.
'''
    (OUT/'presentation_audit.md').write_text(text)

def package():
    prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
    for p in PAGES:
        s=prs.slides.add_slide(prs.slide_layouts[6]);s.shapes.add_picture(str(OUT/p['png']),0,0,width=prs.slide_width,height=prs.slide_height)
        s.notes_slide.notes_text_frame.text=p['title']+'\n'+p['reference']+'\n'+p['note']
    prs.save(OUT/'iNKT_reference_aligned_20260915.pptx')
    for start in range(0,len(PAGES),12):
        chunk=PAGES[start:start+12];sheet=Image.new('RGB',(1440,((len(chunk)+2)//3)*294),'#dddddd');draw=ImageDraw.Draw(sheet)
        for j,p in enumerate(chunk):
            im=Image.open(OUT/p['png']);im.thumbnail((470,265));x=(j%3)*480;y=(j//3)*294;sheet.paste(im,(x,y+22));draw.text((x+5,y+4),f'{p["page"]:02} {p["title"][:57]}',fill='black')
        sheet.save(OUT/'qa'/f'contact_{start+1:02}.jpg')
    (OUT/'slide_index.csv').write_text(pd.DataFrame(PAGES).to_csv(index=False))
    check(len(PdfReader(OUT/'iNKT_reference_aligned_20260915.pdf').pages)==len(PAGES),'PDF and slide index page counts agree')
    check(len(Presentation(OUT/'iNKT_reference_aligned_20260915.pptx').slides)==len(PAGES),'PPTX and PDF page counts agree')
    for p,h in SOURCES.items():check(hashlib.sha256((ROOT/p).read_bytes()).hexdigest()==h,f'input unchanged {p}')
    manifest={'source_sha256':SOURCES,'builder_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'pages':PAGES,'tables':EXPORTS,'DE_recomputed':False,'enrichment_recomputed':False,'embedding_recomputed':False,'checks_passed':len(CHECKS),'checks':CHECKS,'visual_review':'pending'}
    (OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    (OUT/'README.md').write_text('# iNKT reference-aligned presentation — 2026-09-15\n\n'
        '- [Revised PDF](iNKT_reference_aligned_20260915.pdf)\n- [Matching PPTX](iNKT_reference_aligned_20260915.pptx)\n- [Audit: original format, mismatch and correction](presentation_audit.md)\n- [Slide index](slide_index.csv)\n- [Displayed numerical tables](tables/)\n- [PNG and SVG pages](slides/)\n- [Sources and verification](manifest.json)\n\n'
        f'{len(PAGES)} slides. This is a presentation revision from existing results. No analytical statistics or embeddings were rerun. '
        'PPTX contains rendered slide images; SVG pages and source code support figure edits. Earlier dated reports remain available.\n')

def main():
    for name in ['slides','tables','qa']: (OUT/name).mkdir(parents=True,exist_ok=True)
    record(ROOT/'input/iNKT/iNKT.pptx');record(ROOT/'docs/references/blooda_adv-2024-014592-main.pdf');record(M/'sources/BLOODA_ADV-2024-014592-mmc2.pdf')
    a=ad.read_h5ad(record(M/'objects/scored_base.h5ad'));check(a.shape==(15532,10670),'current object expected cohort')
    with PdfPages(OUT/'iNKT_reference_aligned_20260915.pdf') as pdf:
        overview(pdf);qc_and_embedding(a,pdf)
        image_page(ROOT/'output/iNKT_reproduction_deck/20260825_UMAP_PC_sweep/umap_pc_sweep/figures/umap_pc_sweep_tissue_cluster.png','UMAP: 50, 100 and 200 principal components','Controlled sensitivity comparison; selected 50 PCs','Original iNKT slide 4; frozen 20260825 PC-sweep result.',pdf,'PC number varies; all other settings are fixed. Labels are fixed reference clusters, not separate reclustering outcomes.')
        marker_pages(a,pdf);frequency_pages(pdf)
        for uid,ref in [('global','Original iNKT slide 10'),*[(f'tissue__{t}','Original iNKT slide 11') for t in TISSUES],('cluster_tissue__C0__bone_marrow','Original iNKT slide 17 layout; current BM C0'),('cluster_tissue__C6__thymus','Original iNKT slide 13 layout; current Thy C6'),('cluster_tissue__C5-1__spleen','Original iNKT slide 19 layout; current Spl C5-1'),('cluster_tissue__C5-2__bone_marrow','Original iNKT cluster-DE layout; current BM C5-2')]:gene_page(a,uid,pdf,ref+'; current 20260905 DE.')
        legacy_pathway_pages(pdf);volcano_page(pdf);enrichment_pages(pdf);venn_pages(pdf);integration_page(pdf)
        image_page(M/'figures/13_PPT_KEGG_heat_shock_sensitivity.png','Additional analysis: signed GSEA and heat-shock sensitivity','Keep the signed NES heatmap as a separate analysis','20260905 existing GSEA; this is not the enrichment-score plot of Blood Fig. 3E.',pdf,'Red/blue represent signed NES; symbols use GSEA q. Hsp/Dnaj masking changes the ranking and tested gene sets.')
        image_page(M/'figures/09_full_feature_NK_state_sensitivity.png','Additional analysis: reference NK state coverage','Full-feature sensitivity for the requested state mapping','R02; original 20260905 full-feature state-score output.',pdf,'Adaptive state is not reliably scoreable. These continuous programs do not establish the paper NK cluster identities.')
        image_page(M/'figures/10_stable_bone_marrow.png','Additional analysis: bone marrow separately','Within-tissue clustering, condition and state maps','R03; 20260905 stability-selected bone marrow output.',pdf,'New within-tissue labels differ from original C# labels; use saved cell-overlap tables for correspondence.')
        image_page(C/'figures/20260830_IL4_CD94_condition_masks.png','IL-4 and CD94: condition-masked expression','Both conditions shown using the same saved embedding','20260830 condition-mask output; marker display requested in the earlier meeting.',pdf,'Complementarity is weak, not strict exclusivity. UMAP positions are not physical tissue locations or RNA velocity.')
        image_page(D/'figures/discovery_BM_C4_focal.png','Local information: bone marrow C4','Current candidate expression effects across sensitivity analyses','R13; 20260905 discovery results; relates to the local-vs-global question in original iNKT slides 27–28.',pdf,'Same-data sensitivity checks support a candidate, not an independent replicate or established mechanism.')
        image_page(D/'figures/discovery_CCT_driver_evidence.png','Local information: spleen C3 CCT/TriC program','Specific driver genes and current enrichment evidence','R13; 20260905 discovery output.',pdf,'Shared driver genes link several pathway names. Expression evidence does not directly measure protein-folding or killing function.')
        final_status(pdf)
    write_audit();package();print('COMPLETE',len(PAGES),'slides',len(CHECKS),'checks',flush=True)

if __name__=='__main__':main()
