"""Independent data-to-display checks for the dated reference-aligned presentation."""
from pathlib import Path
import hashlib,json,re
import numpy as np
import pandas as pd
from pptx import Presentation
from pypdf import PdfReader
import pypdfium2 as pdfium
ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'output/iNKT_presentation_alignment_20260915'
M=ROOT/'output/iNKT_meeting_followup_20260905'
C=ROOT/'output/iNKT_reproduction_deck/20260830_C5_paper_Fig3DEF_followup'
checks=[]
def check(ok,label):
    if not ok:raise AssertionError(label)
    checks.append(label)
manifest=json.loads((OUT/'manifest.json').read_text())
for name,digest in manifest['source_sha256'].items():check(hashlib.sha256((ROOT/name).read_bytes()).hexdigest()==digest,'original input unchanged: '+name)
for p in (OUT/'tables').glob('DEG_*.csv'):
    unit=p.stem[4:];display=pd.read_csv(p).set_index('gene');source=pd.read_csv(M/'de'/f'{unit}.csv.gz').set_index('gene').loc[display.index]
    for col in ['scores','logfoldchanges','pvals','pvals_adj','mean_log1p_expression_control','mean_log1p_expression_tumor','pct_expressing_control','pct_expressing_tumor']:
        check(np.allclose(display[col],source[col],rtol=1e-12,atol=0),'DEG display equals source '+unit+' '+col)
for p in (OUT/'tables').glob('Fig3C_*.csv'):
    d=pd.read_csv(p);nc=d.control_count.sum();nt=d.tumor_count.sum();den=d.control_count+d.tumor_count;valid=den>0
    check(np.allclose(d.control_cluster_frequency_pct,d.control_count/nc*100),'Control frequency reconstructed '+p.stem)
    check(np.allclose(d.tumor_cluster_frequency_pct,d.tumor_count/nt*100),'Tumor frequency reconstructed '+p.stem)
    check(np.allclose(d.loc[valid,'control_share_within_tissue_cluster_pct'],d.loc[valid,'control_count']/den[valid]*100),'Control share reconstructed '+p.stem)
    check(np.allclose(d.loc[valid,'tumor_share_within_tissue_cluster_pct'],d.loc[valid,'tumor_count']/den[valid]*100),'Tumor share reconstructed '+p.stem)
source=pd.read_csv(M/'tables/Blood_TableS3_all48_pathway_validation.csv')
for p in (OUT/'tables').glob('Fig3E_reference_*.csv'):
    unit=p.stem[len('Fig3E_reference_'):];d=pd.read_csv(p);original=source[source.unit_id.eq(unit)&source.sensitivity.eq('full')].set_index('paper_term').loc[d.paper_term]
    check(len(d)==17,'all reference terms '+unit)
    check(np.allclose(d.Tumor_ORA_p,original.Tumor_ORA_p),'ORA P unchanged '+unit)
    check(np.allclose(d.Tumor_ORA_fdr,original.Tumor_ORA_fdr),'ORA BH FDR unchanged '+unit)
    check(np.allclose(d.current_score,-np.log(d.Tumor_ORA_p)),'ORA score is -ln(P) '+unit)
for rule in ['paper','robust']:
    d=pd.read_csv(C/'tables'/f'20260830_Fig3F_C5split_BM_Spleen_membership_{rule}.csv');cols=[c for c in d if c.startswith('C5-')];shown=pd.read_csv(OUT/'tables'/f'Fig3F_exact_regions_{rule}.csv').set_index('bitmask')
    for mask in range(1,16):
        selected=np.ones(len(d),bool)
        for i,col in enumerate(cols):selected &= d[col].to_numpy(bool) if mask&(1<<i) else ~d[col].to_numpy(bool)
        check(int(selected.sum())==int(shown.loc[mask,'count']),f'Venn {rule} exact region {mask}')
check(len(PdfReader(OUT/'iNKT_reference_aligned_20260915.pdf').pages)==len(manifest['pages']),'PDF page count')
prs=Presentation(OUT/'iNKT_reference_aligned_20260915.pptx')
check(len(prs.slides)==len(manifest['pages']),'PPTX page count')
for page,slide in zip(manifest['pages'],prs.slides):
    check(len(slide.shapes)==1 and hashlib.sha256(slide.shapes[0].image.blob).hexdigest()==hashlib.sha256((OUT/page['png']).read_bytes()).hexdigest(),f'PPTX matches rendered page {page["page"]}')
# Check that actual vector PDF text stays within the physical page bounds.
doc=pdfium.PdfDocument(OUT/'iNKT_reference_aligned_20260915.pdf');overflow=[]
for i in range(len(doc)):
    page=doc[i];w,h=page.get_size();text=page.get_textpage()
    for j in range(text.count_chars()):
        char=text.get_text_range(j,1)
        if not char.strip():continue
        l,b,r,t=text.get_charbox(j)
        if min(l,b)<-1 or r>w+1 or t>h+1:overflow.append({'page':i+1,'char':char,'box':[l,b,r,t]})
check(not overflow,'all vector PDF text within page boundaries')
# Updated date navigation: every entry remains a live link to its declared source.
nav=json.loads((ROOT/'iNKT_by_date/manifest.json').read_text())
for e in nav['entries']:
    p=ROOT/'iNKT_by_date'/e['destination'];check(p.is_symlink() and p.exists() and p.resolve()==(ROOT/e['source']).resolve(),'date navigation '+e['destination'])
result={'status':'passed','checks_passed':len(checks),'checks':checks,'slides':len(prs.slides),'source_files_unchanged':True,'analyses_rerun':False,'visual_review':'pending','pdf_text_overflows':overflow}
(OUT/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:v for k,v in result.items() if k!='checks'},indent=2))
