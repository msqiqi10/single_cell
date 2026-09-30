"""Build Chinese R01-R13 response documents from previously completed analysis outputs.
No DE, integration, enrichment or velocity analyses are rerun by this script.
"""
from __future__ import annotations
import os
for k in ['OMP_NUM_THREADS','OPENBLAS_NUM_THREADS','MKL_NUM_THREADS','NUMEXPR_NUM_THREADS']:os.environ[k]='1'
os.environ['CUDA_VISIBLE_DEVICES']=''
import sys,json,hashlib,shutil,base64,html,zipfile,re
from pathlib import Path
import numpy as np,pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from fontTools import subset
from fontTools.ttLib import TTFont
sys.path.insert(0,'/tmp/inkt_revision_docx_tools')
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path('/home/zzz0054/bio3')
M=ROOT/'output/iNKT_meeting_followup_20260905';D=ROOT/'output/iNKT_discovery_20260905';OUT=ROOT/'output/iNKT_revision_responses_20260906'
TITLE='iNKT 修改意见逐条回应：分析状态、方法与主要结果'
REFINED='cluster_c5_split_20260830'
SOURCES={};TABLES=[]
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p):
    p=Path(p);SOURCES[str(p.relative_to(ROOT))]=sha(p);return pd.read_csv(p,low_memory=False)
def mt(name):return read(M/'tables'/name)
def dt(name):return read(D/'tables'/name)
def num(x,n=3):return '—' if pd.isna(x) else f'{x:.{n}f}'
def q(x):return '—' if pd.isna(x) else ('0（有限置换估计）' if x==0 else f'{x:.3g}')
def asset(p):
    p=Path(p);tag='meeting' if M in p.parents else ('discovery' if D in p.parents else 'report')
    dest=OUT/'assets'/f'{tag}__{p.name}'
    if p!=dest:shutil.copy2(p,dest)
    SOURCES[str(p.relative_to(ROOT))]=sha(p)
    return str(dest.relative_to(OUT))
def data(p):
    p=Path(p);tag='meeting' if M in p.parents else ('discovery' if D in p.parents else 'source')
    dest=OUT/'data'/f'{tag}__{p.name}';shutil.copy2(p,dest);SOURCES[str(p.relative_to(ROOT))]=sha(p)
    return {'name':p.name,'path':str(dest.relative_to(OUT)),'original':str(p),'sha256':sha(p)}
def table(label,headers,rows):return {'caption':label,'headers':headers,'rows':[[str(x) for x in row] for row in rows]}
def section(rid,title,status,done,method,result,pending,talk,tables,figures,files,meeting):
    return dict(id=rid,title=title,status=status,done=done,method=method,result=result,pending=pending,talk=talk,tables=tables,figures=[{'path':asset(p),'caption':cap} for p,cap in figures],files=[data(p) for p in files],meeting=meeting)
def plot_dotplot(d):
    genes='Tmsb10 Cd52 Hspa8 Dnaja1 Cish Il7r Il4 Klrd1 Gata3 Rorc'.split()
    groups=[g for g in d.group.unique() if g.startswith(('bone_marrow/C0/','spleen/C5-1/','thymus/C6/'))]
    fig,ax=plt.subplots(figsize=(10,4.5),layout='constrained')
    for r in d[d.group.isin(groups)&d.gene.isin(genes)].itertuples():
        im=ax.scatter(genes.index(r.gene),groups.index(r.group),s=200*r.fraction_positive,c=r.mean_log1p,cmap='viridis',vmin=0,vmax=3.5)
    ax.set_xticks(range(len(genes)),genes,rotation=35,ha='right',fontsize=10)
    ax.set_yticks(range(len(groups)),[g.replace('bone_marrow','BM').replace('spleen','Spl').replace('thymus','Thy') for g in groups],fontsize=9);ax.invert_yaxis()
    ax.set_title('Selected reporting examples: Ctrl versus T2\nDot area = fraction detected; color = mean log1p expression',fontsize=12)
    fig.colorbar(im,ax=ax,label='Mean log1p expression')
    path=OUT/'assets/R06_dotplot_examples.png';fig.savefig(path,dpi=190);plt.close(fig);return path

def make_sections():
    emb=mt('embedding_metrics.csv');cov=mt('state_signature_coverage.csv').set_index('signature');full=mt('full_feature_score_coverage.csv').set_index('signature');stab=mt('selected_stability_resolution.csv')
    counts=mt('IL4_CD94_counts_by_cluster_tissue_condition.csv');freq=mt('frequency_reconciliation_1.csv');legacy=mt('legacy_best_response_candidates_NOT_identity.csv');de=mt('de_status_original.csv');paper=mt('Blood_TableS3_all48_pathway_validation.csv');ppt=mt('PPT_pathway_validation_every_row.csv');dot=mt('meeting_driver_dotplot_values.csv');priority=dt('discovery_priority_all4_FDR05.csv');focal=dt('discovery_BM_C4_focal_genes.csv');path=dt('discovery_functional_priority.csv');short=dt('discovery_shortlist.csv');allcan=dt('gene_candidates_all_sensitivities.csv');qc=dt('QC_matching_audit.csv');redundancy=mt('KEGG_top10_driver_redundancy.csv');display=mt('KEGG_display_selection_all_significant.csv')
    gsea=pd.concat([read(p) for p in sorted((M/'enrichment').glob('*__full__KEGG_Mouse_2019__GSEA.csv.gz'))],ignore_index=True)
    ss=[]
    em=emb[emb.scope.eq('all_tissues')]
    t=table('校正前后定量指标（邻域纯度越低表示对应标签更混合）',['表示','群数','同组织邻居比例','同条件邻居比例','与原 C5 拆分标签 ARI'],[[r.method,r.n_clusters,num(r.tissue_neighbor_purity),num(r.condition_neighbor_purity),num(r.ARI_to_original_refined)] for r in em.itertuples()])
    ss.append(section('R01','Batch effect removal 前后比较','已做校正尝试；技术 batch 解释未解决','已生成原始、按 tissue 校正、按 sample 校正三套表示和 UMAP，并计算混合程度、分群变化及 state score 的邻域一致性。','Harmony 使用相同的前 50 个 PC，theta=2、最多 20 轮；之后重建 20 邻居图、UMAP（min_dist=0.5）和 Leiden（resolution=0.5）。三套均保留原计数；DE 沿用未整合的表达。','按 tissue 校正后，同组织邻居比例由 0.872 降至 0.504；与原 refined 标签的 ARI 从原表示重算的 0.868 降至 0.214。按 sample 校正后，同条件邻居比例降至 0.501。组织混合增强，同时原生物分群和条件结构也明显改变。','没有独立技术 batch 字段；sample 同时编码 tissue 和 condition。因此尚不能识别被移除的成分中哪些属于技术偏差，也不能宣称已证明完成纯技术 batch removal。','我们已经做了 batch removal 对照，但组织混合改善伴随分群结构改变，目前把它作为敏感性结果。',[t],[(M/'figures/18_integration_evaluation.png','图 R01：整合前后混合与分群指标。原始及两套校正 UMAP 在随附文件中。')],[M/'tables/embedding_metrics.csv',M/'figures/04_original.png',M/'figures/04_harmony_tissue.png',M/'figures/04_harmony_sample.png'],'10:14–10:18'))
    five=['NK_resting','NK_adaptive','NK_activated','NK_type_I_IFN','NK_cytokine']
    t=table('五种参考程序的 marker 覆盖',['程序','主分析使用/参考数','完整特征使用数','当前判定'],[[k,f'{cov.loc[k,"n_used"]}/{cov.loc[k,"n_original"]}',full.loc[k,'n_used'],'不足以评分' if k=='NK_adaptive' else ('仅完整特征敏感性可评分，信号稀疏' if k=='NK_resting' else '已评分')] for k in five])
    ss.append(section('R02','参考 marker/signature 的 cluster 状态映射','部分完成；五状态分类未确立','已生成参考状态、iNKT1/2/17、cycling 等连续分数图，以及组织/条件/群汇总；还回到相同 15,532 个细胞的完整 32,285 个表达特征检查覆盖。','参考 NK marker 来自会议所示 Dufva 2023 图中面板；iNKT 对照面板来自 Krovi 2020。人鼠映射采用可唯一匹配的 MGI 同源关系；至少 3 个基因且覆盖≥40% 才评分。Scanpy score_genes 使用表达量匹配的对照基因（ctrl_size=50、n_bins=25）。完整特征敏感性另要求 marker 至少在 10 个保留细胞中检出。','主过滤集可评分五种 NK 程序中的 3 种；完整特征可评分 4 种。C1 的 type I IFN 分数较高，骨髓和脾脏的 Ctrl/T2 两组均可见；C5-1 偏 iNKT2 程序，C5-2 偏 iNKT17。它们是相对 marker 程序，不是自动分类标签。','Adaptive 只有 1/4 个 marker 可用；resting 补入的 Gzmk 仅 17 个细胞检出，不能据此做可靠五分类。成熟度及方向也未确立；不同程序的分值不能当概率或直接横向比较。','已有状态映射支持 IFN 相关群和 C5 两个亚群的程序差异，但尚不能把所有细胞可靠地分成论文五种状态。',[t],[(M/'figures/09_full_feature_NK_state_sensitivity.png','图 R02：完整特征敏感性下的 NK 状态分数。每个面板色阶独立；adaptive 留空表示覆盖不足。')],[M/'tables/state_signature_coverage.csv',M/'tables/full_feature_score_coverage.csv',M/'tables/full_feature_marker_coverage_audit.csv',M/'tables/state_score_group_summary.csv',M/'figures/01_reference_state_maps.png',M/'figures/02_iNKT_and_function_maps.png'],'09:23–09:28；10:04–10:06；10:17–10:18'))
    t=table('按跨种子稳定性选择的敏感性分群',['组织','细胞数','分辨率','群数','最小 ARI'],[[r.tissue,{'bone_marrow':6110,'spleen':6793,'thymus':2629}[r.tissue],r.resolution,r.n_clusters_seed0,num(r.minimum_ARI)] for r in stab.itertuples()])
    ss.append(section('R03','单组织重分析，骨髓优先','已完成计算与稳定性比较','骨髓、脾脏、胸腺均独立重算 HVG/PCA/邻居图/UMAP/Leiden；保存 marker、状态图、与原群的细胞交集表及新的组内 DE。','从既有 log1p 表达出发，在每个组织内重新选 3,000 HVG 和 50 PC；20 邻居。比较 resolution=0.3/0.5/0.8、seed=0/17/42 的分群。保留原先 0.5 版本；另将最小两两 ARI 最大的分辨率作为事后敏感性选择。','骨髓在 resolution=0.5 时最小 ARI=0.297，降至 0.3 后形成 5 群且最小 ARI=0.958；脾脏在 0.3 时为 4 群、ARI=0.990；胸腺选 0.5，为 6 群、ARI=0.702。较粗分群更稳定，但稳定性本身不能证明生物身份。','单组织分析已经做完；尚未得到一套可直接等同于论文五状态的标签。新的 S#/L# 与原 C# 仅通过细胞交集对应，不直接按编号认同。','骨髓单独分析可以得到较稳定的较粗分群；这有助于复核局部信号，但还需结合 marker 解释身份。',[t],[(M/'figures/10_stable_bone_marrow.png','图 R03：骨髓稳定性选择后的分群、条件及三个可评分 NK 程序。')],[M/'tables/within_tissue_cluster_stability.csv',M/'tables/selected_stability_resolution.csv',M/'tables/bone_marrow_stable_old_new_crosswalk.csv',M/'tables/bone_marrow_stable_top30_markers.csv',M/'figures/10_stable_spleen.png',M/'figures/10_stable_thymus.png'],'10:16–10:18'))
    t=table('Velocity 输入核查',['输入/输出','当前情况'],[['表达计数 counts','已有'],['spliced / unspliced 层','当前对象没有'],['可生成上述计数的 BAM/FASTQ','已搜索 input/iNKT，未发现'],['RNA velocity、velocity graph、方向图','未计算'],['已有 DPT / 中位数位移图','属于其他分析，不计为 velocity 完成']])
    ss.append(section('R04','RNA velocity 映射到 clusters','未运行；输入不足','已检查分析对象中的 layer 和 input/iNKT 原始文件类型，保存缺失输入审计。','目前实际执行的方法是输入审计；没有运行 velocity 估计、速度图或其投影。后续计算需要与本次细胞匹配的 spliced/unspliced，或 BAM/FASTQ 及相应参考注释。','当前对象只含 counts 层；未找到 spliced/unspliced 或可据以生成它们的原始比对/测序文件。因此没有可汇报的真实 velocity 方向或动态 marker 结果。','缺少估计 RNA 剪接动力学所需的数据。此项仍待补输入；DPT、状态高低或 UMAP 位移都不能补足该输入缺口。','Velocity 还没跑，原因是缺少 spliced/unspliced 或可生成这些计数的原始数据；现有箭头只表示位置变化。',[t],[],[M/'tables/input_missing_tasks.json'],'09:24–09:28；10:01–10:03；10:08–10:13'))
    chosen=legacy[legacy.legacy_unit.isin(['c5_bone_marrow','c2_spleen','c6_spleen','c1_thymus'])]
    t=table('同组织中 DEG 响应最相似的候选对应（不代表同一身份）',['旧 PPT 单元','当前候选群','旧基因数','同方向比例','logFC Spearman'],[[r.legacy_unit,r.current_unit.replace('cluster_tissue__','').replace('__',' / '),r.n_measured,f'{r.direction_concordance:.0%}',num(r.spearman_logFC)] for r in chosen.itertuples()])
    ss.append(section('R05','旧 PPT 与当前 DEG/群对应核对','DEG 核对已做；群身份对应部分完成','旧 PPT 提取的 DEG 按相应组织与各当前群逐项比较，保留旧/新 logFC、FDR、检测状态及方向；最高效应相关者单独列作响应候选。','每个同组织候选群使用完整 DE 表查找旧基因；比较符号一致率和 logFC 的 Spearman 相关。选择最高相关者用于展示，因此该最优匹配具有选择效应。','代表性对应中，旧 c5_bone_marrow 与当前 BM C0 的 61 个基因全部同向，Spearman=0.708；旧 c6_spleen 与当前脾 C5-1 的 29 个基因全部同向，Spearman=0.846。方向信息大体保留，效应大小和通过显著性阈值的数量并不完全一致。','当前得到的是表达响应相似度，不是经过细胞级历史标签验证的身份对应。尚不能断言旧 c5 就是当前 C0；此处保留部分完成状态。','旧 PPT 的不少基因方向在当前数据中保留，但旧新群对应仍只能作为候选，不能按编号直接认同。',[t],[(M/'figures/16_PPT_DEG_response_concordance.png','图 R05：旧基因在各同组织最佳响应候选中的方向一致率；最佳匹配不构成身份验证。')],[M/'tables/legacy_DEG_signed_validation_all_candidates.csv',M/'tables/legacy_DEG_response_similarity.csv',M/'tables/legacy_best_response_candidates_NOT_identity.csv'],'09:30–09:41'))
    examples=['cluster_tissue__C0__bone_marrow','cluster_tissue__C5-1__spleen','cluster_tissue__C6__thymus']
    de_rows=de[de.unit_id.isin(examples)]
    t=table('主图示例的条件内 DE 数量（FDR≤0.05 且 |log2FC|≥0.25）',['群','Ctrl / T2 细胞','T2 上调数','T2 下调数'],[[r.unit_id.replace('cluster_tissue__','').replace('__',' / '),f'{r.n_control} / {r.n_tumor}',int(r.n_up_fdr05_fc025),int(r.n_down_fdr05_fc025)] for r in de_rows.itertuples()])
    ss.append(section('R06','有方向 DEG 表与 Ctrl/T2 dotplot','已完成；稀疏比较有标记','31 个预设原始比较中 26 个可计算 DE，21 个满足每条件≥20 细胞并进入主通路分析。完整原始 dotplot 有 19 个基因×12 个条件群；下图为便于汇报选取的子集。','使用全部 10,670 个检测基因的 raw log1p 表达，Wilcoxon、tie correction、基因层面 BH。方向统一为 T2 相对 Ctrl；主阈值 FDR≤0.05 且 |log2FC|≥0.25。dot 大小为表达>0 的细胞比例，颜色为群内所有细胞的平均 log1p（包含零表达细胞）。','例如骨髓 C0 的 Tmsb10 log2FC=+0.465、Cd52=+0.346，而 Hspa8=-0.494、Dnaja1=-0.552；这类变化可发生在检测比例本来就很高的基因上，因此必须同时看表达量和检测率。','5 个比较因任一条件不足 2 个细胞无法计算；另有 5 个虽可计算但低于每条件 20 个细胞，仅保留探索性结果。完整基因表已提供，主图只选择代表基因而非展示所有基因。','我们已补齐方向、效应、显著性和表达比例；dotplot 显示的是表达变化及覆盖，不能用细胞数替代动物重复。',[t],[(plot_dotplot(dot),'图 R06：从原始 dotplot 数值表重绘的汇报子集；完整 19×12 图及数值表随附。')],[M/'tables/de_status_original.csv',M/'tables/meeting_driver_dotplot_values.csv',M/'figures/15_PPT_and_paper_driver_dotplot.png']+[M/f'de/{uid}.csv.gz' for uid in examples],'09:41–09:42'))
    f=ppt[ppt.sensitivity.eq('full')&ppt.strict_min20];stats=f.groupby('status').ppt_row.nunique()
    names={'tested_current_membership':'已按当前成员检验','source_library_unavailable':'来源旧库不可用','term_unresolved':'条目尚未解析','below_minimum_set_size':'测量成员不足最小集合要求'}
    t=table('130 条旧 PPT 通路记录的状态（完整基因、满足细胞数的比较）',['状态','原 PPT 记录数'],[[names[k],int(stats.get(k,0))] for k in names])
    ss.append(section('R07','旧 PPT 的 KEGG/Reactome 等通路验证','部分完成；所有旧记录已登记','旧 PPT 的 130 条记录均纳入核查，其中 87 条来自 Reactome_2021、37 条来自 KEGG_2021_HUMAN，其余 6 条来自 WikiPathway/Spike/NCI。120 条有当前完整成员的可计算检验。','解析原始通路 ID，获取可用 KEGG 或官方 Reactome 完整成员并严格映射到测量鼠基因；使用测量基因背景做 ORA。分别保留旧 PPT 兼容的混合方向 DEG 规则（FDR≤0.05、|log2FC|≥1.2）和当前有方向的规则。旧 PPT 指定 Reactome 反应构成预指定检验家族。','Heat-shock/伴侣蛋白相关反应仍反复出现。例如当前胸腺 C6 的 ATP hydrolysis by HSP70，在 T2 上调 DEG 中 ORA FDR=0.00129，在下调 DEG 中 FDR=0.0113；不同成员方向不同，不能把它解释为整个反应统一激活或抑制。','6 条因旧来源库不可用、2 条未解析、2 条测量成员太少而未完成检验。即使可计算，当前成员/版本及阈值也与旧库不同，因此不能称为旧分析逐数值复现。','旧 PPT 的通路已逐项登记并验证可计算部分；许多结果仍由少数伴侣蛋白驱动，不能按通路数量计算独立发现。',[t],[(M/'figures/17_PPT_shared_driver_matrix.png','图 R07：旧 PPT KEGG 名称所共享的驱动基因；用于解释重复信息，不表示表达方向。')],[M/'tables/PPT_pathway_validation_every_row.csv',M/'tables/PPT_pathway_driver_genes_signed.csv',M/'tables/legacy_reactome_full_membership_audit.csv'],'09:52–09:58'))
    terms=['TNF signaling pathway','NF-kappa B signaling pathway','JAK-STAT signaling pathway','IL-17 signaling pathway','Th17 cell differentiation','Cytokine-cytokine receptor interaction','Antigen processing and presentation']
    rr=[]
    for term in terms:
        z=gsea[gsea.Term.eq(term)];sig=z[z['FDR q-val']<=.05]
        rr.append([term,q(z['FDR q-val'].min()),str(len(sig)),'Ctrl 方向：整体与脾组织' if len(sig) else '未获 GSEA q≤0.05 支持'])
    t=table('21 个主比较中关注通路的完整基因 GSEA 结果',['通路','最小 q','通过比较数','主要判定'],rr)
    ss.append(section('R08','Blood Figure 3E KEGG 与驱动基因验证','主图已核查；补表部分受库覆盖限制','Figure 3E 的 17 条通路均核查；补表 S3 的 48 条扩展到 21 个比较×完整/去 heat-shock 两版本，共 2,016 条记录。另核对 S5 的 23 个基因及会议关注基因，并制作原图并排对照。','参照论文的上调 DEG 门槛做 ORA（nominal P≤0.05、线性 FC≥1.5，即 log2FC≥log2(1.5)），同时另做有方向 GSEA（完整 Wilcoxon 排名、5,000 次 gene-set permutations）。本地 Wilcoxon 与冻结 KEGG_2019_Mouse 不同于论文的统计模型和库，因此是方法对照，不是精确数值复刻。','TNF、NF-κB、JAK–STAT、IL-17、Th17 等在 21 个主比较中均未达到 GSEA q≤0.05。Antigen processing/presentation 在整体及脾组织呈 Ctrl 方向富集；胸腺 C6 的 Cish 为 log2FC=-0.818、FDR=0.000133，骨髓 C0 则 +0.231、FDR=0.0711。结果不支持这些轴在所有群统一朝论文方向变化。','S3 中 3 条不在冻结 2019 库、1 条测量成员太少，故 44/48 条可检验。论文主图中 JAK–STAT 和 SLE 本身 FDR 分别约 0.0747、0.0599；出现在图上不等于通过 0.05。ORA 与 GSEA 分别解释，不能把一种方法的阴性自动推广到另一种。','论文关注通路已经在我们数据中系统核查；结果存在不支持和方向不同，不能写成统一复现论文机制。',[t],[(M/'figures/11_Blood_Fig3E_side_by_side.png','图 R08：论文 Figure 3E 原图与本数据指定组织/群的 GSEA 对照。数值不是同一种富集统计量。')],[M/'tables/Blood_TableS3_all48_pathway_validation.csv',M/'tables/Blood_reference_driver_gene_validation.csv',M/'figures/12_Blood_paper_threshold_ORA.png',M/'figures/14_Blood_driver_genes_signed.png'],'09:45–09:48；09:57–09:58；10:06–10:08'))
    red=redundancy[redundancy.unit_id.eq('global')&redundancy.term_a.eq('Oxidative phosphorylation')&redundancy.term_b.isin(['Alzheimer disease','Parkinson disease','Thermogenesis'])]
    t=table('同一整体比较内的 leading-edge 重复例子',['通路 A','通路 B','驱动基因 Jaccard'],[[r.term_a,r.term_b,num(r.leading_edge_jaccard)] for r in red.itertuples()])
    ss.append(section('R09','缩小 scope、说明 top 10 与驱动重复','已完成比较范围及显示规则核查','固定组织×群内 T2/Ctrl 比较，同时保留组织和全局汇总。完整显著 GSEA 表、top 10 选择标记及两两 leading-edge 重叠表均已生成；发现阶段进一步检查更具体的 Reactome/KEGG 程序。','主结果先按 GSEA q≤0.05 筛选，再按 q 升序、NES 降序排列，取前 10 用于显示，不按人为大类合并。用 shared/(union) 的 Jaccard 衡量 leading-edge 重复，完整表保留未进 top 10 的条目。','21 个主比较合计有 145 条显著通路×比较记录，涉及 '+str(display.Term.nunique())+' 个不同名称；这不是 145 个独立机制。整体比较中 OxPhos 与 Alzheimer/Parkinson/Thermogenesis 的驱动重叠均较高。更具体的候选转向脾 C3 的 CCT/TriC 和骨髓 C0 的剪接基因子集，见 R13。','通路细化及重复检查已完成；尚不能仅凭更具体的名称确认机制。疾病命名的富集条目尤其需要先看实际驱动基因。','Top 10 只是展示规则；我们已检查完整结果，并把共享驱动基因的重复名称与更具体的候选分开说明。',[t],[],[M/'tables/KEGG_display_selection_all_significant.csv',M/'tables/KEGG_top10_driver_redundancy.csv',D/'tables/pathway_exact_driver_groups.csv',D/'tables/discovery_functional_priority.csv'],'09:48–09:57'))
    rows=[]
    for label,pat in [('骨髓 C0 / Spliceosome','Spliceosome'),('脾 C3 / CCT 微管蛋白折叠','Formation Of Tubulin')]:
        r=path[path.term.str.contains(pat,regex=False)].iloc[0];rows.append([label,q(r.FDR),q(r.no_heat_FDR),'GSEA q=0.654 / 0.801，未显著' if pat=='Spliceosome' else '这里为 Reactome ORA；未新增该库 GSEA'])
    t=table('去 heat-shock 后仍有的候选证据（ORA 与 GSEA 分开）',['组织/群/程序','完整 ORA FDR','去 Hsp/Dnaj ORA FDR','其他方法核查'],rows)
    ss.append(section('R10','Mask heat-shock 后的敏感性分析','已完成','对全部 21 个主比较各跑完整与去 heat-shock 两版本，形成 42 个 KEGG GSEA 结果，并补对应 ORA；新发现筛选也保留同样敏感性。','将名称匹配 Hsp* 或 Dnaj* 的基因同时从 DEG/query、排名、通路成员和测量背景移除，保持细胞集合不变；比较相同组织/群、相同方向的结果。','去 heat-shock 并未令会议关注的免疫通路普遍获得支持。新筛选中骨髓 C0 的 Spliceosome 子集 ORA 仍支持，但 GSEA 不显著；脾 C3 CCT/TriC 相关 ORA 也保留。旧 PPT 曾出现 Spliceosome，但当时主要由 HSPA 基因驱动。','已完成所定义的 Hsp/Dnaj 排除。CCT 等其他伴侣蛋白未被该正则排除，因此“去 heat-shock 后保留”不能等同于与全部应激/蛋白折叠过程无关，也不能把剩余信号自动视为更真实。','去掉 Hsp/Dnaj 后仍有部分具体基因程序值得追踪，但没有得到普遍复现论文免疫通路的结果。',[t],[(M/'figures/13_PPT_KEGG_heat_shock_sensitivity.png','图 R10：旧 PPT 相关 KEGG 在完整/去 heat-shock 排名中的 NES 与显著性对照。')],[M/'tables/enrichment_ranking_audit.csv',M/'tables/PPT_pathway_validation_every_row.csv',D/'tables/discovery_functional_priority.csv'],'09:55–09:58'))
    selected=counts[(counts.tissue.eq('bone_marrow')&counts.cluster.eq('C0'))|(counts.tissue.eq('spleen')&counts.cluster.isin(['C5-1','C5-2']))]
    t=table('逐群阳性计数示例（完整表包含全部原群）',['组织/群','条件','总数','Il4+','Klrd1+','双阳性'],[[r.tissue+'/'+r.cluster,r.condition,r.n_cells,r.Il4_positive,r.Klrd1_positive,r.both_positive] for r in selected.itertuples()])
    ss.append(section('R11','IL4/CD94 的逐群计数、互补分布与位置','已完成描述性分析','按 all/三个组织×9 个原群（含 C5-1/C5-2）×2 条件生成 72 行计数表；补 108 行 all-cells/Il4+/Klrd1+ 中位坐标结果、带 n/N 标注的表达 UMAP 和位移图。','以原 counts>0 定义 Il4 或 Klrd1（CD94 对应基因）阳性；报告双阳性/单阳性/双阴性及检测 phi、表达 Spearman。位置分析在固定原 UMAP 上计算各条件阳性细胞的二维中位数，箭头仅在两组各≥5 个阳性细胞时显示。','骨髓 C0 的 Il4 检测率由 354/2533=13.98% 至 286/2090=13.68%；Klrd1 由 42.24% 至 44.45%。检测 phi 为 -0.112/-0.121，提示较弱负关联而非完全互斥。该群 Il4+ 与 Klrd1+ 的中位位移分别为 0.249、0.301 个 UMAP 坐标单位。','没有独立动物层面的互补性或空间位移检验；UMAP 距离没有物理单位，受嵌入影响。不能根据这些箭头宣称迁移、分化方向或 velocity。','数量与位置都已补齐；骨髓 C0 显示较弱的 Il4/Klrd1 负关联，但变化并不支持强互斥或实际迁移的结论。',[t],[(M/'figures/06_counts_bone_marrow.png','图 R11：骨髓按条件显示 Il4/Klrd1 表达；每群标注阳性细胞数/总细胞数。')],[M/'tables/IL4_CD94_counts_by_cluster_tissue_condition.csv',M/'tables/IL4_CD94_fixed_UMAP_median_shifts.csv',M/'figures/07_median_shifts_bone_marrow.png',M/'figures/06_counts_spleen.png',M/'figures/06_counts_thymus.png'],'09:04–09:08'))
    fr=freq[(freq.tissue.eq('bone_marrow')&freq.cluster.eq('C0'))|(freq.tissue.eq('spleen')&freq.cluster.eq('C5-1'))]
    t=table('计数、不同分母及频率差的实例',['组织/群','Ctrl / T2 数量','Ctrl 群频率','T2 群频率','差值（百分点）'],[[r.tissue+'/'+r.cluster,f'{r.control_count} / {r.tumor_count}',f'{r.control_cluster_frequency_pct:.2f}%',f'{r.tumor_cluster_frequency_pct:.2f}%',f'{r.tumor_minus_control_percentage_points:+.3f}'] for r in fr.itertuples()])
    ss.append(section('R12','比例图、分母、基线与公式','已完成；全部 27 个组织×群已核对','补充每组织 Ctrl/T2 总数、每群计数、两种百分比和频率差值；统一条件颜色，逐项核对符号。','群频率差 Δpp=100×[n(T2,群,组织)/N(T2,组织)−n(Ctrl,群,组织)/N(Ctrl,组织)]。群内 T2 share=100×n(T2,群)/[n(T2,群)+n(Ctrl,群)]。两者分母不同；share 的零差值基线是该组织总体 T2 share，而非固定 50%。','骨髓 C0 虽然 Ctrl 数量 2533 多于 T2 的 2090，但群频率从 74.96% 至 76.53%，差值 +1.566 pp；它的 T2 share=45.21%，高于骨髓背景 44.70%。脾 C5-1 频率由 6.98% 至 10.23%，差值 +3.250 pp。所有 27 个群的符号与各自基线一致。','计算和图示已完成；这些是捕获细胞中的相对比例，不代表器官内绝对细胞量增加。','原图看似矛盾来自分母不同，公式与实际计数核对后是一致的；脾 C5-1 相对频率增加，但不能据此称绝对扩增。',[t],[(M/'figures/08_frequency_denominators.png','图 R12：两种分母和频率差的并排解释；正文提供骨髓 C0 与脾 C5-1 的可直接汇报数值。')],[M/'tables/frequency_reconciliation_0.csv',M/'tables/frequency_reconciliation_1.csv'],'09:01–09:02；09:12–09:21'))
    t=table('骨髓 C4 的六个局部候选（T2 相对 Ctrl）',['基因','原群 log2FC','原群 FDR','QC 后 log2FC','QC 后 FDR','骨髓整体 FDR'],[[r.gene,num(r.original_log2FC),q(r.original_FDR),num(r.QC_log2FC),q(r.QC_FDR),q(r.tissue_FDR)] for r in focal.itertuples()])
    il7=allcan[allcan.unit_id.eq('cluster_tissue__C5-1__spleen')&allcan.gene.eq('Il7r')].iloc[0]
    ss.append(section('R13','当前 pipeline 中局部新信号的筛选与解释','已完成探索性分析；独立验证未完成','沿用当前 iNKT 原群和组织重聚类，做全基因及完整 KEGG/Reactome ORA 筛选。审计 35 个对照，29 个 QC 匹配后仍各≥20 细胞并重算 DE；按重聚类、匹配及 5 次抽样检查候选敏感性。','在各比较内，按 log UMI、log 检测基因数、线粒体比例的条件盲四分位联合分层，层内等量选择 Ctrl/T2。先要求原群 FDR≤0.05、|log2FC|≥0.25、检测率≥10%，再检查稳定群映射（原群≥50%细胞）、效应同向、足够细胞及 QC 平衡。27 条局部非主导家族记录中，事后以四项 DE FDR 均≤0.05、五次抽样同号优先，得到 '+str(len(priority))+' 条记录（'+str(priority.gene.nunique())+' 个不同基因）。','首要候选是骨髓 C4：Tcf7/Itga4/Emb 下调，Ly6a/S100a6/Cdca4 上调，四项 DE 敏感性均支持。Tcf7、Emb、Ly6a 在骨髓整体比较未达 FDR 0.05，提示局部信息在平均结果中不明显。脾 C3 的 CCT/TriC 相关基因下降也获 ORA 支持。骨髓 C0 剪接基因子集有 ORA 支持但 GSEA 未显著；旧候选 Il7r 在 QC 匹配后 FDR='+num(il7.QC_FDR)+'，已降级。','没有可识别独立动物重复，无法完成动物层面的重复验证或因果机制检验。重聚类和抽样仍来自同一数据；局部与整体显著性不同不构成正式交互作用证据。部分候选已在旧 PPT 出现，未见旧表也不等于文献首创。','当前最值得继续验证的是骨髓 C4 的局部表达状态及脾 C3 的 CCT 程序；这是本数据生成的候选发现，尚不是已证实机制。',[t],[(D/'figures/discovery_BM_C4_focal.png','图 R13a：骨髓 C4 六个候选在原群、重聚类、QC 匹配及重聚类+QC 中的效应。'),(D/'figures/discovery_CCT_driver_evidence.png','图 R13b：脾 C3 的 CCT/TriC 驱动基因在四项敏感性中的表达效应；星号为基因 FDR≤0.05。')],[D/'tables/discovery_priority_all4_FDR05.csv',D/'tables/discovery_BM_C4_focal_genes.csv',D/'tables/discovery_functional_priority.csv',D/'tables/QC_matching_audit.csv',D/'tables/candidate_cell_subsampling_summary.csv',D/'tables/discovery_pathway_driver_evidence.csv',D/'tables/gene_candidates_all_sensitivities.csv'],'09:42–09:44；09:57；10:18；用户后续明确要求'))
    return ss

COMMON=[
 '范围：按已确认 R01–R13 逐条回应，分析结果来自 2026-09-05 的 meeting_followup 与 discovery 两轮输出；本次只读取、核对、重排结果并制作汇报图，不重新估计 DE/富集/轨迹。',
 '数据：15,532 个保留细胞，10,670 个主分析基因。骨髓 Ctrl/T2=3379/2731，脾脏=3422/3371，胸腺=1372/1257。原分群采用 C0–C7，C5 拆为 C5-1/C5-2；新组织分群用 L#/S#。',
 '统计口径：基因效应方向统一为 T2 相对 Ctrl。DE 的 BH FDR、ORA 的 BH FDR、GSEA 的置换 q 值属于不同检验。显著性和基因数随比较、阈值、库版本变化；未检出、不可计算与无显著证据分别标记。',
 '共同限制：每组织×条件只有一个 sample 标签，缺少可识别动物重复及独立技术 batch。p/FDR 是探索性细胞层面统计；整合、重聚类、匹配、抽样一致性不能代替动物重复，也不能建立因果机制。',
 '来源关系：旧 PPT 是 R05/R07 的比较对象；Blood Advances 2025 Figure 3E 是 R08 的对象。R02 的五种 NK 程序使用会议后半提及的另一篇参考（Dufva 2023）面板，R04 为 RNA velocity 方法，不能把所有图都归到 Blood Figure 3E。'
]

def html_table(t):
    return '<div class="table-caption">'+html.escape(t['caption'])+'</div><table><thead><tr>'+''.join('<th>'+html.escape(x)+'</th>' for x in t['headers'])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+html.escape(x)+'</td>' for x in row)+'</tr>' for row in t['rows'])+'</tbody></table>'
def make_font(text):
    opts=subset.Options();opts.layout_features=['*'];opts.name_IDs=['*'];opts.name_legacy=True;opts.name_languages=['*']
    font=TTFont(OUT/'assets/NotoSansCJKsc-Regular.otf');s=subset.Subsetter(options=opts);s.populate(text=text);s.subset(font);font.save(OUT/'assets/report_zh_subset.otf');font.close()
    return base64.b64encode((OUT/'assets/report_zh_subset.otf').read_bytes()).decode()

def create_html(sections):
    toc='<table><thead><tr><th>编号 / 修改意见</th><th>是否完成</th></tr></thead><tbody>'+''.join(f'<tr><td><a href="#{s["id"]}">{s["id"]} {html.escape(s["title"])}</a></td><td>{html.escape(s["status"])}</td></tr>' for s in sections)+'</tbody></table>'
    body='<header><div class="eyebrow">会议修改意见 · 逐条回应 · 2026-09-06</div><h1>'+TITLE+'</h1><p>每项固定列出：是否分析、实际方法、主要结果、未完成原因、图表及原始数据。</p></header><section id="overview"><h2>总览与共同口径</h2>'+toc+''.join('<p>'+html.escape(x)+'</p>' for x in COMMON)+'</section>'
    for s in sections:
        body+=f'<section class="revision" id="{s["id"]}"><div class="eyebrow">会议出处 {s["meeting"]}</div><h2>{s["id"]}　{html.escape(s["title"])}</h2><div class="status">{html.escape(s["status"])}</div>'
        for label,key in [('已做分析','done'),('采用的方法','method'),('主要结果','result')]:body+='<p><strong>'+label+'：</strong>'+html.escape(s[key])+'</p>'
        for t in s['tables']:body+=html_table(t)
        for fig in s['figures']:
            encoded=base64.b64encode((OUT/fig['path']).read_bytes()).decode();body+='<figure><a href="'+fig['path']+'"><img src="data:image/png;base64,'+encoded+'" alt="'+html.escape(fig['caption'])+'"></a><figcaption>'+html.escape(fig['caption'])+'</figcaption></figure>'
        body+='<p><strong>尚未完成及原因 / 解释边界：</strong>'+html.escape(s['pending'])+'</p><div class="takeaway"><strong>汇报可用表述：</strong>'+html.escape(s['talk'])+'</div>'
        body+=f'<p class="print-source">配套数值表：data/{s["id"]}_report_table_1.csv；完整数据与补图链接见同名 HTML/Word 的 {s["id"]} 条。</p>'
        body+='<details open><summary>原始数据与补充图（随附 data/ 文件夹）</summary><ul>'+''.join('<li><a href="'+f['path']+'">'+html.escape(f['name'])+'</a></li>' for f in s['files'])+'</ul></details></section>'
    body+='<section class="revision"><h2>额外项目与核验记录</h2><p>肝/肾 toxicity 是会议另提项目，未开展数据分析：本地未找到邮件中的 12 个 PMID 清单，无法确认目标文章及数据集。此项不混入 R01–R13。</p><p>既有会议分析验证：42 项单元测试、12,980 项产物检查通过。发现分析验证：48 项单元测试、3,647 项产物检查通过。两轮单元测试有重叠，不能将数量相加作为独立测试数。本次文档的链接、数值、图片及导出核查另见 verification.json。</p><p>文档与底层表均保留不支持、部分可计算和输入不足结果；分析来源与 SHA256 见 source_manifest.json。中文字体来自 Noto CJK，许可证随附。</p></section>'
    font=make_font((OUT/'iNKT_revision_responses_20260906.md').read_text()+TITLE+' '.join(COMMON)+json.dumps(sections,ensure_ascii=False)+re.sub('<[^>]+>','',body)+ '额外项目与核验记录肝肾毒性文章数据集重复单元测试产物检查链接数值图片导出核查中文字体来源许可证随附')
    css='''@font-face{font-family:ReportZH;src:url(data:font/otf;base64,FONT) format('opentype')}*{box-sizing:border-box}html{scroll-behavior:smooth}body{font-family:ReportZH,'Microsoft YaHei',sans-serif;color:#17263b;background:#edf2f7;line-height:1.65;margin:0 auto;max-width:1060px;padding:30px}header,section{background:white;padding:30px 38px;margin:0 0 24px;border:1px solid #dbe3ed;border-radius:8px}h1{font-size:28px;line-height:1.4;margin:10px 0}h2{font-size:22px;margin:6px 0 14px;line-height:1.45}.eyebrow{font-size:12px;color:#60738b}.status{border-left:4px solid #246b92;background:#eff7fb;padding:8px 12px;font-weight:bold}p{margin:12px 0}a{color:#126b9b;text-decoration:none;overflow-wrap:anywhere}table{width:100%;border-collapse:collapse;font-size:13px;line-height:1.5;margin:8px 0 15px;table-layout:auto}th,td{border:1px solid #d7e0e9;padding:7px 8px;text-align:left;overflow-wrap:anywhere}th{background:#e9f0f7}tr:nth-child(even){background:#f7f9fc}.table-caption{font-size:13px;font-weight:bold;margin-top:15px}figure{margin:16px 0;break-inside:avoid}figure img{width:100%;max-height:540px;object-fit:contain}figcaption{font-size:12px;color:#5b687a}.takeaway{background:#edf7f0;border-left:4px solid #3b8355;padding:10px 12px;margin-top:14px}details{font-size:12px;margin-top:13px;color:#536276}details ul{columns:2;padding-left:20px}li{overflow-wrap:anywhere}strong{font-weight:bold}.print-source{display:none}@media print{@page{size:A4;margin:14mm 15mm}body{background:white;max-width:none;padding:0;font-size:9.6pt;line-height:1.55}header,section{padding:0;border:0;border-radius:0;margin:0}header{margin-bottom:7mm}h1{font-size:20pt}h2{font-size:16pt;break-after:avoid}.revision{break-before:page}.eyebrow{font-size:8pt}.status{padding:5px 9px}p{margin:6px 0}table{font-size:8.3pt}th,td{padding:4px 5px}.table-caption{font-size:9pt}figure{margin:8px 0}figure img{max-height:90mm}#R08 figure img{max-height:78mm}figcaption{font-size:8pt}.takeaway{padding:6px 9px;break-inside:avoid}details{display:none}.print-source{display:block;font-size:7.5pt;color:#65748b}.takeaway{margin-top:7px}details ul{columns:2;margin:4px 0}summary{font-size:8pt}a{color:#126b9b}tr{break-inside:avoid}}'''.replace('FONT',font)
    (OUT/'iNKT_revision_responses_20260906.html').write_text('<!doctype html><html lang="zh-CN"><head><meta charset="UTF-8"><title>'+TITLE+'</title><style>'+css+'</style></head><body>'+body+'</body></html>')

def hyperlink(p,label,target):
    rid=p.part.relate_to(target,RT.HYPERLINK,is_external=True);h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid);run=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'126B9B');pr.append(color);run.append(pr);t=OxmlElement('w:t');t.text=label;run.append(t);h.append(run);p._p.append(h)
def doc_table(doc,t):
    p=doc.add_paragraph(t['caption']);p.runs[0].bold=True
    tab=doc.add_table(rows=1,cols=len(t['headers']));tab.style='Light Shading Accent 1'
    for c,x in zip(tab.rows[0].cells,t['headers']):c.text=x
    for row in t['rows']:
        for c,x in zip(tab.add_row().cells,row):c.text=x
    for row in tab.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                for r in para.runs:r.font.size=Pt(8.5)
    return tab

def create_docx(sections):
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.left_margin=sec.right_margin=Inches(.64);sec.top_margin=sec.bottom_margin=Inches(.65)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st=doc.styles[name];st.font.name='Microsoft YaHei';st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    doc.styles['Normal'].font.size=Pt(10);doc.styles['Normal'].paragraph_format.space_after=Pt(6);doc.styles['Normal'].paragraph_format.line_spacing=1.15
    doc.styles['Heading 1'].font.size=Pt(16);doc.styles['Heading 2'].font.size=Pt(12)
    doc.core_properties.title=TITLE;doc.core_properties.subject='R01–R13 修改意见逐条回应';doc.core_properties.author='iNKT analysis project'
    doc.add_heading(TITLE,0);doc.add_paragraph('2026-09-06｜会议修改意见逐条回应',style='Subtitle')
    doc_table(doc,table('13 条修改意见的当前状态',['编号','事项','状态'],[[s['id'],s['title'],s['status']] for s in sections]))
    doc.add_heading('共同数据与统计口径',1)
    for p in COMMON:doc.add_paragraph(p)
    for s in sections:
        doc.add_page_break();doc.add_heading(s['id']+'　'+s['title'],1);doc.add_paragraph('会议出处：'+s['meeting']);p=doc.add_paragraph();p.add_run('当前状态：').bold=True;p.add_run(s['status'])
        for label,key in [('已做分析','done'),('采用的方法','method'),('主要结果','result')]:
            p=doc.add_paragraph();p.add_run(label+'：').bold=True;p.add_run(s[key])
        for t in s['tables']:doc_table(doc,t)
        for fig in s['figures']:
            from PIL import Image
            im=Image.open(OUT/fig['path']);w,h=im.size;width=min(6.8,3.7*w/h)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(OUT/fig['path']),width=Inches(width));doc.add_paragraph(fig['caption'],style='Caption')
        for label,key in [('尚未完成及原因 / 解释边界','pending'),('汇报可用表述','talk')]:
            p=doc.add_paragraph();p.add_run(label+'：').bold=True;p.add_run(s[key])
        doc.add_paragraph('原始数据与补充图（解压资料包后可打开 data/ 内文件）：')
        for f in s['files']:
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);hyperlink(p,f['name'],f['path']);
            for r in p.runs:r.font.size=Pt(8)
    doc.add_page_break();doc.add_heading('额外项目与核验记录',1)
    doc.add_paragraph('肝/肾 toxicity：未开展数据分析。本地未找到会议邮件中的 12 个 PMID 清单，无法确认研究和数据集。该项作为额外工作单列。')
    doc.add_paragraph('既有会议分析：42 项单元测试、12,980 项产物检查通过。发现分析：48 项单元测试、3,647 项产物检查通过。测试存在重叠，不合并计数。本次文档另检查全部 R 编号、数据数值、图片和文件链接，结果见 verification.json。')
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.add_run('iNKT 修改意见逐条回应　|　')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    doc.save(OUT/'iNKT_revision_responses_20260906.docx')

def create_md(sections):
    lines=['# '+TITLE,'','2026-09-06。按 R01–R13 固定编号组织；详细图表、Word、PDF 和离线 HTML 在同目录。','']+COMMON+['']
    for s in sections:
        lines+=['## '+s['id']+' '+s['title'],'','**当前状态：'+s['status']+'**','', '会议出处：'+s['meeting'],'']
        for label,key in [('已做分析','done'),('采用的方法','method'),('主要结果','result')]:lines+=['**'+label+'：** '+s[key],'']
        for t in s['tables']:
            lines += [t['caption'],'','| '+' | '.join(t['headers'])+' |','|'+'|'.join(['---']*len(t['headers']))+'|']+['| '+' | '.join(row)+' |' for row in t['rows']]+['']
        for fig in s['figures']:lines+=['!['+fig['caption']+']('+fig['path']+')','']
        lines+=['**尚未完成及原因 / 解释边界：** '+s['pending'],'','**汇报可用表述：** '+s['talk'],'','数据与补充图：','']+['- ['+f['name']+']('+f['path']+')；原路径：'+f['original'] for f in s['files']]+['']
    lines+=['## 额外 toxicity 项目','','未分析：缺少会议邮件中的 12 个 PMID 清单，无法确认文章和数据集。该项单列，不混入 R01–R13。','', '## 核验记录','','既有分析测试记录来自两个源目录的 verification.json；本次文档核查另见同目录 verification.json。完整文件及 SHA256 见 source_manifest.json。','']
    (OUT/'iNKT_revision_responses_20260906.md').write_text('\n'.join(lines))

def main():
    for n in ['assets','data','logs']:(OUT/n).mkdir(parents=True,exist_ok=True)
    sections=make_sections();assert [s['id'] for s in sections]==[f'R{i:02d}' for i in range(1,14)]
    (OUT/'response_sections.json').write_text(json.dumps(sections,ensure_ascii=False,indent=2))
    pd.DataFrame([{k:s[k] for k in ['id','title','status','done','method','result','pending','talk','meeting']} for s in sections]).to_csv(OUT/'R01_R13_response_matrix.csv',index=False,encoding='utf-8-sig')
    for s in sections:
        for i,t in enumerate(s['tables'],1):pd.DataFrame(t['rows'],columns=t['headers']).to_csv(OUT/f'data/{s["id"]}_report_table_{i}.csv',index=False,encoding='utf-8-sig')
    create_md(sections);create_html(sections);create_docx(sections)
    extra=[ROOT/'TODOs.txt',ROOT/'docs/inkt_revision_checklist_20260906.md',M/'verification.json',D/'verification.json',Path(__file__)]
    for p in extra:SOURCES[str(p.relative_to(ROOT))]=sha(p)
    manifest={'created_date':'2026-09-06','scope':'report existing R01-R13 analysis; no analysis recomputation','source_SHA256':SOURCES,'sections':len(sections),'main_figures':sum(len(s['figures']) for s in sections),'tables':sum(len(s['tables']) for s in sections),'font_source':'https://github.com/notofonts/noto-cjk/tree/main/Sans','font_license':'assets/NotoSansCJK-LICENSE.txt','new_reporting_only_figure':'R06 subset dotplot from existing values','GPU_used':False}
    (OUT/'source_manifest.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False));print(json.dumps({k:v for k,v in manifest.items() if k!='source_SHA256'},ensure_ascii=False,indent=2),flush=True)
if __name__=='__main__':main()
