"""Readable meeting deliverables from completed, audited analysis outputs."""
from run_inkt_meeting_followup import *
from run_inkt_meeting_pathways import libraries, ora_table, clean_rank, matchterm, KEGG_ID_NAMES
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import importlib.metadata as metadata
import textwrap

ANCHORS=['global','tissue__bone_marrow','cluster_tissue__C0__bone_marrow','cluster_tissue__C5-1__bone_marrow','cluster_tissue__C5-2__bone_marrow','tissue__spleen','cluster_tissue__C3__spleen','cluster_tissue__C5-1__spleen','cluster_tissue__C5-2__spleen','tissue__thymus','cluster_tissue__C6__thymus']
def short(s):return s.replace('cluster_tissue__','').replace('tissue__','').replace('__',' / ').replace('bone_marrow','BM').replace('spleen','Spl').replace('thymus','Thy')
def heatmap(ax,values,rows,cols,title,cmap='RdBu_r',limit=2,stars=None):
    cm=plt.get_cmap(cmap).copy();cm.set_bad('#dddddd')
    v=np.asarray(values,float);im=ax.imshow(v,cmap=cm,aspect='auto',vmin=-limit if cmap=='RdBu_r' else 0,vmax=limit)
    ax.set_yticks(range(len(rows)),rows,fontsize=8);ax.set_xticks(range(len(cols)),[short(c) for c in cols],rotation=55,ha='right',fontsize=8);ax.set_title(title,fontsize=11)
    if stars is not None:
        for y,x in zip(*np.where(np.asarray(stars,float)<=.05)):ax.text(x,y,'*',ha='center',va='center',fontsize=10,color='black')
    return im

def stable_pathways():
    lib=libraries()['KEGG_Mouse_2019'];status=pd.read_csv(OUT/'tables/de_status_stable.csv');frames=[];genes=[]
    pg=pd.read_csv(OUT/'sources/Blood_TableS5.csv');gset=set(pg['Gene name'])|set(PAPER_GENES)
    for r in status[status.strict_min20&status.status.eq('computed')].itertuples():
        d=pd.read_csv(OUT/f'de/{r.unit_id}.csv.gz');q=d[d.gene.isin(gset)].copy();q['unit_id']=r.unit_id;genes.append(q)
        for sens in ['full','without_heat_shock']:
            clean,_=clean_rank(d,sens!='full')
            for rule in ['robust','paper']:
                q=ora_table(clean,lib,set(clean.gene),rule);q['unit_id']=r.unit_id;q['sensitivity']=sens;q['resolution']=r.resolution;frames.append(q)
    csv(pd.concat(frames,ignore_index=True),'stable_cluster_KEGG_ORA.csv.gz');csv(pd.concat(genes,ignore_index=True),'stable_cluster_paper_driver_DEG.csv')

def pathway_plots():
    d=pd.read_csv(OUT/'tables/Blood_TableS3_all48_pathway_validation.csv');q=d[d.sensitivity.eq('full')&d.in_Fig3E];terms=PAPER_TERMS
    nes=q.pivot(index='paper_term',columns='unit_id',values='GSEA_NES').reindex(index=terms,columns=ANCHORS);fdr=q.pivot(index='paper_term',columns='unit_id',values='GSEA_q').reindex_like(nes)
    fig,(left,right)=plt.subplots(1,2,figsize=(19,9),gridspec_kw={'width_ratios':[1,1.55]},layout='constrained')
    im=Image.open(OUT/'sources/Blood2025_Fig3_source_page.png');left.imshow(im);w,h=im.size;left.set_xlim(.448*w,.844*w);left.set_ylim(.405*h,.207*h);left.axis('off');left.set_title('Supplied Blood Advances paper, Fig. 3E\nOriginal up-DEG enrichment; NK/CML',fontsize=12)
    v=heatmap(right,nes,terms,ANCHORS,'Our iNKT T2 - Ctrl: signed KEGG GSEA\nRed=T2 enriched; blue=Ctrl enriched; *=q<=0.05',stars=fdr,limit=3);fig.colorbar(v,ax=right,shrink=.7,label='NES, 5,000 gene-set permutations')
    fig.suptitle('Paper pathways mapped to explicit tissue / cluster contrasts\nDifferent cell type, disease context and analysis method; no equivalence of score magnitudes',fontsize=14);save(fig,'11_Blood_Fig3E_side_by_side')
    ora=q.pivot(index='paper_term',columns='unit_id',values='Tumor_ORA_fdr').reindex_like(nes);own=-np.log10(ora.clip(lower=1e-12));paper=q.drop_duplicates('paper_term').set_index('paper_term').reindex(terms)['paper_fdr']
    fig,(ax0,ax1)=plt.subplots(1,2,figsize=(15,9),gridspec_kw={'width_ratios':[.75,5]},layout='constrained')
    heatmap(ax0,-np.log10(paper.to_numpy()[:,None]),terms,['Paper S3'],'Paper\nBH FDR',cmap='Reds',limit=4,stars=paper.to_numpy()[:,None]);v=heatmap(ax1,own,terms,ANCHORS,'Our paper-threshold up-DEG ORA: p<=.05; linear FC>=1.5\n*=BH FDR<=.05; all eligible KEGG terms included in correction',cmap='Reds',limit=4,stars=ora);ax1.set_yticklabels([]);fig.colorbar(v,ax=ax1,label='-log10(BH FDR), capped at4');save(fig,'12_Blood_paper_threshold_ORA')
    # Fixed original PPT KEGG vocabulary; both directions preserved in the signed rank.
    terms=list(dict.fromkeys(KEGG_ID_NAMES.values()));cols=ANCHORS[2:5]+ANCHORS[6:9]+ANCHORS[10:];vals=[];qs=[];labels=[]
    for uid in cols:
        for sens in ['full','without_heat_shock']:
            t=pd.read_csv(OUT/f'enrichment/{uid}__{sens}__KEGG_Mouse_2019__GSEA.csv.gz').set_index('Term');vals.append([float(t.loc[x,'NES']) if x in t.index else np.nan for x in terms]);qs.append([float(t.loc[x,'FDR q-val']) if x in t.index else np.nan for x in terms]);labels.append(short(uid)+('\nfull' if sens=='full' else '\nno HSP'))
    fig,ax=plt.subplots(figsize=(15,8),layout='constrained');v=heatmap(ax,np.array(vals).T,terms,labels,'PPT-named KEGG pathways | heat-shock exclusion sensitivity\nRed=T2; blue=Ctrl; *=q<=.05; same cells and contrasts; Hsp*/Dnaj* removed from rank + sets',limit=3,stars=np.array(qs).T);fig.colorbar(v,ax=ax,label='GSEA NES');save(fig,'13_PPT_KEGG_heat_shock_sensitivity')

def driver_plots():
    d=pd.read_csv(OUT/'tables/Blood_reference_driver_gene_validation.csv');genes=list(dict.fromkeys(PAPER_GENES+pd.read_csv(OUT/'sources/Blood_TableS5.csv')['Gene name'].tolist()))
    val=d.pivot(index='gene',columns='unit_id',values='log2FC').reindex(index=genes,columns=ANCHORS);fdr=d.pivot(index='gene',columns='unit_id',values='fdr').reindex_like(val)
    fig,ax=plt.subplots(figsize=(13,11),layout='constrained');v=heatmap(ax,val,genes,ANCHORS,'Paper driver genes in our explicit contrasts\nRed=higher in T2; blue=higher in Ctrl; *=gene BH FDR<=.05; grey=filtered out',stars=fdr,limit=2);fig.colorbar(v,ax=ax,label='log2 fold change, clipped at +/-2');save(fig,'14_Blood_driver_genes_signed')
    a=ad.read_h5ad(OUT/'objects/scored_base.h5ad');gene_order='Il4 Klrd1 Tmsb10 Cd52 Hspa8 Dnaja1 Jun Fos Cish Socs1 Socs2 Tnf Tnfrsf1b Il7r Gzmb Prf1 Tbx21 Gata3 Rorc'.split();targets=[('bone_marrow','C0'),('bone_marrow','C5-1'),('bone_marrow','C5-2'),('spleen','C3'),('spleen','C5-1'),('thymus','C6')];rows=[];groups=[]
    for tissue,c in targets:
        for cond in ['Ctrl','T2']:
            m=a.obs.tissue.astype(str).eq(tissue)&a.obs[REFINED_CLUSTER_KEY].astype(str).eq(c)&a.obs.condition.astype(str).eq(cond);label=f'{tissue}/{c}/{cond} (n={m.sum()})';groups.append(label)
            for gene in gene_order:
                x=vec(a.raw[:,gene].X)[m];rows.append({'group':label,'gene':gene,'fraction_positive':float(np.mean(x>0)),'mean_log1p':float(np.mean(x)),'n':int(m.sum())})
    d=pd.DataFrame(rows);csv(d,'meeting_driver_dotplot_values.csv');fig,ax=plt.subplots(figsize=(16,8),layout='constrained')
    for r in d.itertuples():scat=ax.scatter(gene_order.index(r.gene),groups.index(r.group),s=220*r.fraction_positive,c=r.mean_log1p,cmap='viridis',vmin=0,vmax=float(d.mean_log1p.quantile(.98)),edgecolor='none')
    ax.set_xticks(range(len(gene_order)),gene_order,rotation=55,ha='right');ax.set_yticks(range(len(groups)),[short(x) for x in groups]);ax.invert_yaxis();fig.colorbar(scat,ax=ax,label='Mean log1p expression');ax.set_title('Requested dotplot: signed DEG tables accompany detection and expression\nDot area=positive fraction, 0-100%; a missing dot means zero detection')
    for v in [.1,.5,1.]:ax.scatter([],[],s=220*v,color='gray',label=f'{v:.0%}')
    ax.legend(title='Positive fraction',loc='upper left',bbox_to_anchor=(1.17,1),frameon=False);save(fig,'15_PPT_and_paper_driver_dotplot')
    sim=pd.read_csv(OUT/'tables/legacy_DEG_response_similarity.csv');g=sim[sim.strict_min20].sort_values(['legacy_unit','spearman_logFC','n_concordant_robust'],ascending=[True,False,False]).groupby('legacy_unit',sort=False).head(1);csv(g,'legacy_best_response_candidates_NOT_identity.csv')
    fig,ax=plt.subplots(figsize=(13,8),layout='constrained');ax.barh(range(len(g)),g.direction_concordance,color='#4b86b4');ax.set_yticks(range(len(g)),[f'{r.legacy_unit} → {short(r.current_unit)}' for r in g.itertuples()],fontsize=9)
    for i,r in enumerate(g.itertuples()):ax.text(min(.96,r.direction_concordance+.01),i,f'{r.n_concordant_robust}/{r.n_measured} same-direction DEGs',va='center',fontsize=8)
    ax.set_xlim(0,1.35);ax.set_xlabel('Fraction of measured old DEGs with the same sign');ax.set_title('PPT DEG validation: highest response correlation within matching tissue\nResponse resemblance is NOT evidence of identical cell populations');save(fig,'16_PPT_DEG_response_concordance')
    old=pd.read_csv(OLD/'legacy_ppt_pathway_tables_extracted.csv');old=old[old.SOURCE.str.startswith('KEGG')];pairs=[]
    for _,r in old.iterrows():
        m=re.search(r'hsa(\d{5})',r.NAME);term=KEGG_ID_NAMES.get(m.group(1),r.NAME) if m else r.NAME
        for gene in str(r.GENE_SYM).split(','):pairs.append((term,gene.strip()))
    matrix=pd.crosstab(pd.Series([x[0] for x in pairs],name='term'),pd.Series([x[1] for x in pairs],name='gene')).clip(upper=1);fig,ax=plt.subplots(figsize=(15,7),layout='constrained');v=heatmap(ax,matrix,matrix.index,matrix.columns,'Original PPT KEGG terms share many of the same driver genes\nUnion of displayed overlaps across slides; presence does not give expression direction',cmap='Blues',limit=1);save(fig,'17_PPT_shared_driver_matrix')

def integration_plot():
    d=pd.read_csv(OUT/'tables/embedding_metrics.csv').iloc[:3];fig,axes=plt.subplots(1,3,figsize=(14,4),layout='constrained')
    for ax,metric in zip(axes,['tissue_neighbor_purity','condition_neighbor_purity','ARI_to_original_refined']):
        ax.bar(range(3),d[metric],color=['#888888','#4878cf','#d65f5f']);ax.set_xticks(range(3),['Original graph','Harmony tissue','Harmony sample'],rotation=25,ha='right',fontsize=9);ax.set_ylim(0,1);ax.set_title(metric.replace('_',' '),fontsize=11)
        for i,v in enumerate(d[metric]):ax.text(i,v+.02,f'{v:.3f}',ha='center')
    fig.suptitle('Integration sensitivity: better mixing also changes biological/condition structure\nNo independently identified technical batch; no corrected counts used for differential expression',fontsize=12);save(fig,'18_integration_evaluation')

FIGURES=[
 ('18_integration_evaluation','Batch-removal attempt: quantify what integration removes'),
 ('04_original','Recomputed original graph'),('04_harmony_tissue','Harmony by tissue: integration sensitivity'),('04_harmony_sample','Harmony by sample: condition is confounded'),
 ('01_reference_state_maps','Reference NK marker panels on current clusters'),('09_full_feature_NK_state_sensitivity','Coverage check using all raw gene features'),('02_iNKT_and_function_maps','iNKT lineage and functional marker panels'),
 ('10_stable_bone_marrow','Bone marrow separately: stability-selected resolution0.3'),('10_stable_spleen','Spleen separately: stability-selected resolution0.3'),('10_stable_thymus','Thymus separately: stability-selected resolution0.5'),
 ('11_Blood_Fig3E_side_by_side','Paper Fig.3E beside our tissue / cluster results'),('12_Blood_paper_threshold_ORA','Validate the paper with its up-DEG threshold'),('14_Blood_driver_genes_signed','Validate the actual paper driver genes'),
 ('16_PPT_DEG_response_concordance','Original PPT: DEG directions and response resemblance'),('15_PPT_and_paper_driver_dotplot','Requested expression / detection dotplot'),('13_PPT_KEGG_heat_shock_sensitivity','Original PPT KEGG terms: full versus heat-shock exclusion'),('17_PPT_shared_driver_matrix','Why multiple PPT terms repeat the same signal'),
 ('06_counts_bone_marrow','IL4 / CD94: bone marrow positive / total counts'),('06_counts_spleen','IL4 / CD94: spleen positive / total counts'),('06_counts_thymus','IL4 / CD94: thymus positive / total counts'),('08_frequency_denominators','Frequency formulas and explicit denominators'),('07_median_shifts_all','Descriptive distribution shifts: these are NOT RNA velocity'),
]

def package():
    prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
    def add_slide(title,body=None,picture=None):
        s=prs.slides.add_slide(prs.slide_layouts[6]);box=s.shapes.add_textbox(Inches(.45),Inches(.2),Inches(15.1),Inches(.65));tf=box.text_frame;tf.text=title;tf.paragraphs[0].font.size=Pt(25);tf.paragraphs[0].font.bold=True
        if body:
            box=s.shapes.add_textbox(Inches(.65),Inches(1.2),Inches(14.7),Inches(6.7));tf=box.text_frame;tf.word_wrap=True
            for i,line in enumerate(body):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.size=Pt(23);p.space_after=Pt(18)
        if picture:
            im=Image.open(picture);w,h=im.size;ratio=min(15.2/w,7.5/h);wi,hi=w*ratio,h*ratio;s.shapes.add_picture(str(picture),Inches((16-wi)/2),Inches(1+(7.5-hi)/2),width=Inches(wi),height=Inches(hi))
        ft=s.shapes.add_textbox(Inches(.45),Inches(8.65),Inches(15),Inches(.25));ft.text_frame.text=f'2026-09-05 | iNKT meeting follow-up | {len(prs.slides)} | Exploratory cell-level results; no identified animal replicates';ft.text_frame.paragraphs[0].font.size=Pt(9)
    overview=['Meeting scope: integration comparison, reference-state mapping, tissue reanalysis, PPT and paper validation, heat-shock sensitivity, counts and denominators.','15,532 retained cells; original C5-1 / C5-2 labels preserved. No biological-replicate metadata was inferred.','Completed results and negative findings are shown together. Full signed DEG, pathway and driver tables are in tables/ and de/.','RNA velocity requires spliced/unspliced counts or BAM/FASTQ; these inputs are missing. The toxicity PMID inventory is also missing.']
    add_slide('Meeting follow-up: executed analyses and remaining input gaps',overview)
    with PdfPages(OUT/'meeting_results_20260905.pdf') as pdf:
        fig,ax=plt.subplots(figsize=(16,9));ax.axis('off');ax.text(.04,.9,'iNKT meeting follow-up | 2026-09-05',fontsize=24,weight='bold',transform=ax.transAxes)
        for i,t in enumerate(overview):ax.text(.04,.73-i*.17,textwrap.fill(t,110),fontsize=16,transform=ax.transAxes,va='top',linespacing=1.5)
        pdf.savefig(fig,bbox_inches='tight');plt.close(fig)
        for name,title in FIGURES:
            path=OUT/'figures'/f'{name}.png'
            if not path.exists():raise FileNotFoundError(path)
            add_slide(title,picture=path);fig,ax=plt.subplots(figsize=(16,9));ax.imshow(Image.open(path));ax.axis('off');ax.set_title(title,fontsize=14,pad=12);fig.text(.02,.015,'Exploratory cell-level comparisons | source tables, parameters and limitations in README.md',fontsize=9);pdf.savefig(fig,bbox_inches='tight');plt.close(fig)
    add_slide('Pending inputs: velocity and the extra toxicity project',['Velocity: supply matching spliced/unspliced matrices or BAM/FASTQ with reference annotation. Counts-only matrices cannot determine RNA velocity.','Current DPT and IL4/CD94 median-shift arrows are not velocity and do not establish developmental direction.','Provide the emailed liver/kidney toxicity PMID list to continue that separate meeting task.','Batch interpretation: sample/tissue integration was tested. Real technical batch and animal/pooling metadata are still needed for stronger inference.'])
    prs.save(OUT/'meeting_results_20260905.pptx')


def write_readme():
    d=pd.read_csv(OUT/'tables/Blood_TableS3_all48_pathway_validation.csv');full=d[d.sensitivity.eq('full')];gene=pd.read_csv(OUT/'tables/Blood_reference_driver_gene_validation.csv');counts=pd.read_csv(OUT/'tables/de_status_original.csv');sim=pd.read_csv(OUT/'tables/legacy_DEG_response_similarity.csv');novel=pd.read_csv(OUT/'tables/cluster_signals_diluted_in_tissue.csv');validation=pd.read_csv(OUT/'tables/PPT_pathway_validation_every_row.csv')
    terms=['TNF signaling pathway','NF-kappa B signaling pathway','JAK-STAT signaling pathway','IL-17 signaling pathway','Th17 cell differentiation','Cytokine-cytokine receptor interaction','Antigen processing and presentation']
    lines=[]
    for term in terms:
        q=full[full.paper_term.eq(term)]
        if len(q):
            sig=q[q.GSEA_q<=.05];hits='；'.join(f'{short(r.unit_id)} ({r.GSEA_NES:+.2f}, q={r.GSEA_q:.3g})' for r in sig.itertuples()) or '无 q≤0.05 的对照';lines.append(f'| {term} | {q.GSEA_q.min():.4g} | {hits} |')
        else:
            vals=[]
            for r in counts[counts.strict_min20&counts.status.eq('computed')].itertuples():
                t=pd.read_csv(OUT/f'enrichment/{r.unit_id}__full__KEGG_Mouse_2019__GSEA.csv.gz');t=t[t.Term.eq(term)];vals.extend([{'unit':r.unit_id,'q':float(x['FDR q-val']),'NES':float(x['NES'])} for x in t.to_dict('records')])
            hits='；'.join(f'{short(x["unit"])} ({x["NES"]:+.2f}, q={x["q"]:.3g})' for x in vals if x['q']<=.05) or '无 q≤0.05 的对照';lines.append(f'| {term} | {min(x["q"] for x in vals):.4g} | {hits} |')
    cs=gene[gene.gene.eq('Cish')&gene.unit_id.isin(['cluster_tissue__C0__bone_marrow','cluster_tissue__C3__spleen','cluster_tissue__C6__thymus'])];ct='；'.join(f'{short(r.unit_id)} log2FC={r.log2FC:+.3f}, FDR={r.fdr:.3g}' for r in cs.itertuples())
    statuses=[
      ('Batch effect removal 尝试','已执行，有解释限制','Harmony 按 tissue、按 sample 各一版，与原图比较；没有独立技术 batch 字段。','figures/18_integration_evaluation.png'),
      ('五种参考状态映射','已执行；adaptive 无法可靠评分','当前过滤集可评分3种；完整原始特征敏感性可评分4种，resting 的 Gzmk 仅17细胞检出。','tables/full_feature_score_coverage.csv'),
      ('单组织分析，骨髓优先','已完成','三个组织分别重新选HVG/PCA/邻居/聚类；保留0.5分辨率以及按种子稳定性选择的分群。','tables/selected_stability_resolution.csv'),
      ('论文 velocity 叠加到 clusters','输入阻塞，未完成','缺少 spliced/unspliced 或 BAM/FASTQ；没有用DPT或位移箭头替代。','tables/input_missing_tasks.json'),
      ('PPT DEG 与通路验证','已完成可检验条目；部分旧库无法复现','全部原始DEG和130条通路记录逐行核查，KEGG和可解析Reactome重算；其余旧库保留未解决状态与驱动基因结果。','tables/PPT_pathway_validation_every_row.csv'),
      ('Blood Figure3E KEGG和驱动基因验证','已完成跨数据集验证','17条主图通路与S5的23基因已检验；S3全部48条逐项核查，库缺失/过小条目保留未检验状态。原图并排。','figures/11_Blood_Fig3E_side_by_side.png'),
      ('缩小scope、解释top10和重复驱动基因','已完成','固定组织×cluster对照，完整统计表；top10仅显示顺序，所有显著term和重叠驱动基因另存。','tables/KEGG_top10_driver_redundancy.csv'),
      ('Heat-shock masking','已完成','对所有≥20细胞/条件的原始对照，从排名、基因集、背景同时去掉Hsp*/Dnaj*再计算。','figures/13_PPT_KEGG_heat_shock_sensitivity.png'),
      ('IL4/CD94计数和位置变化','已完成','按组织、条件、cluster含C5-1/C5-2报告阳性/总数；位置箭头只表示描述性中位数位移。','tables/IL4_CD94_counts_by_cluster_tissue_condition.csv'),
      ('频率和分母说明','已完成','补组织×条件总数，计数/百分比/公式；区分condition内frequency与cluster内condition share。','figures/08_frequency_denominators.png'),
      ('cluster内新增而汇总稀释的信号','已完成探索性筛选','提供signed effect、tissue/global对照和dotplot；未将显著性不同当作交互作用显著。','tables/cluster_signals_diluted_in_tissue.csv'),
      ('额外肝/肾toxicity项目','输入阻塞，未完成','缺少会议所说邮件PMID清单。','tables/input_missing_tasks.json'),
    ]
    checklist=pd.DataFrame(statuses,columns=['meeting_task','status','detail','artifact']);csv(checklist,'meeting_task_status.csv')
    text='''# 2026-09-05 会议事项执行结果

这次按会议清单实际运行分析。现有数据能做的分析已完成；velocity 和额外 toxicity 项目仍缺输入，不能算完成。

- 汇报：[结果 PDF](meeting_results_20260905.pdf) · [可编辑 PPTX](meeting_results_20260905.pptx)
- 完整结果：[任务状态](tables/meeting_task_status.csv) · [签名覆盖](tables/full_feature_marker_coverage_audit.csv) · [原分群DE状态](tables/de_status_original.csv) · [稳定分群DE状态](tables/de_status_stable.csv)

## 会议任务逐项对应

| 事项 | 状态 | 实际内容 | 结果 |
|---|---|---|---|
'''+''.join(f'| {a} | {b} | {c} | [查看]({p}) |\n' for a,b,c,p in statuses)+f'''
## 当前可以得出的结果

**整合确实让组织混合，但还不能称为去除了已识别的技术 batch。** 邻居同组织比例从0.872降到0.504（tissue Harmony）；与原 refined cluster 的ARI从原图重算的0.868降到0.214。sample Harmony把同条件邻居比例从0.602降到0.501。sample同时编码条件和组织，不能从这些结果区分被消除的是技术偏差还是生物差异。因此保留原结果做表达差异检验，Harmony作为对照结果。所有校正对象在objects/中，不覆盖20260830结果。

**单组织分析已完成，但细分群存在稳定性差异。** 骨髓0.5分辨率跨种子最小ARI=0.297；脾脏=0.544。补充按最小ARI最大选择的分辨率：骨髓0.3（5群，最小ARI=0.958）、脾脏0.3（4群，0.990）、胸腺0.5（6群，0.702）。这是看到稳定性结果后的敏感性选择，不是预注册结果。保留原0.5版本，不把新的L#/S#等同于旧C#。

**参考状态只能作为连续marker程序。** 原10,670基因集支持activated、type I IFN、cytokine；全32,285表达特征、同15,532细胞的敏感性分析补出resting（Gzmk仅17细胞）。adaptive仅Lag3可用且稀疏，无法可靠评分。Ncam1在原始六样本中也为零。严格一对一人鼠同源映射排除模糊关系；完整覆盖损失在表中。没有强制五分类，也不能把未评分解释为该状态不存在。Cycling、iNKT17等评分对基因过滤/对照基因池敏感，需要连同覆盖表读。

**论文的KEGG应区分方法与方向。** 官方补充方法p6给出nominal P≤0.05、线性FC≥1.5；补充表S3明确基于K1+K2的上调基因做富集。这里按log2FC≥log2(1.5)和同阈值重算上/下调ORA，并另外做有方向GSEA。原文用Partek基因特异模型和其KEGG库；本地是Wilcoxon与冻结KEGG_2019_Mouse，不能称为精确重现。原文S3上调列表202基因、背景6763基因；本地背景10,670基因。原文JAK–STAT FDR=0.0747、SLE=0.0599，虽然画在主图，均未通过0.05。

以下是新跑的组织×cluster及汇总对照，不是旧的pooled-cluster GSEA：

| 通路 | 最小GSEA q | q≤0.05对照（NES正=T2、负=Ctrl） |
|---|---:|---|
'''+ '\n'.join(lines)+f'''

Cish具体方向：{ct}。这不支持不分组织/cluster地宣布复现论文Cish统一上调机制。

**PPT验证同时保留DEG方向与通路成员。** 所有旧DEG在对应组织的各现有cluster中比较，并提供效应相关性；最高相关者只表示响应相似，不是细胞身份对应。Reactome使用当前官方完整reaction成员及严格同源映射；同一stable ID去重后检验。旧Reactome_2021、WikiPathway、Spike、NCI等不是同一版本/数据库，不足以复现的条目明确标记；绝不把PPT中已命中的几个基因当作完整通路来重算P值。热休克相关pathway名称的重复驱动在图17和冗余表中展示。

## 计算与解释边界

- 输入：20260830已选定UMAP/C5分裂对象，15,532细胞×10,670基因。原始6个表达矩阵仅用于额外marker覆盖核查；Multiplexing Capture不作为RNA表达。
- 整合：Harmony0.0.10，theta=2，最多20轮，同一50PC表示；20邻居，UMAP min_dist=0.5，Leiden0.5。组织重分析重新选3000 HVG和50PC；分辨率0.3/0.5/0.8，种子0/17/42。
- DE：完整raw log1p表达的Wilcoxon、tie correction、基因层面BH；T2相对Ctrl；主阈值FDR≤0.05且|log2FC|≥0.25。原分群共{int(counts.status.eq('computed').sum())}个可计算对照，其中{int((counts.strict_min20&counts.status.eq('computed')).sum())}个满足每条件≥20细胞并进入主通路检验；小群不冒充可靠阴性。
- GSEA：完整有方向Wilcoxon排名，KEGG集合5–500测量基因，5,000次gene-set permutations，seed=20260905。旧库大写符号先映射回原鼠基因拼写。q=0只表示有限置换下的估计为0；小群中大量零值/并列排名见ranking audit，解释需谨慎。
- ORA：hypergeometric，集合3–500基因；对每一完整合格库、每方向、每对照做BH。原PPT指定Reactome反应构成单独预指定检验家族，不能当作全Reactome发现。PPT兼容规则为旧混合方向DEG union（FDR≤0.05、|log2FC|≥1.2）；这种ORA富集没有激活方向。
- Heat-shock敏感性：Hsp*/Dnaj*同时从rank、集合、背景和query中移除；原结果全部保留。未去除任意细胞，也不保证暴露出的通路更可信。
- 单组织稳定分群也已重算DE、KEGG ORA和论文基因表：[stable ORA](tables/stable_cluster_KEGG_ORA.csv.gz)。它们是分群敏感性，不参与原cluster的预设GSEA网格。
- 仅有每组织×条件1个sample标签，没有可识别动物replicate；这些是探索性细胞层面比较，不能将细胞数当生物重复。跨对照没有统一多重检验校正。
- 共筛出{len(novel)}条cluster内明显、汇总后减弱/未过阈值的gene×contrast记录。这不是正式cluster×condition interaction检验。
- 所有中位数位移箭头和既有DPT都不是RNA velocity；不从中推断发育方向。

## 尚需补充的输入

1. **Velocity**：匹配本次细胞的spliced/unspliced计数，或BAM/FASTQ及参考注释路径。当前对象只有counts层；搜索input/iNKT未发现速度所需文件。拿到后才能计算真实velocity并叠加cluster/state图。
2. **Toxicity**：会议邮件中的肝/肾毒性PMID清单或本机文件路径。
3. **解释batch及生物重复**：动物/池化/技术批次对应信息；这不影响已完成的整合敏感性对照，但决定能否做真正的技术batch与replicate推断。

## 来源与复现

- [会议记录](../../TODOs.txt)，时间点与任务对应见[已核对清单](../iNKT_next_steps_review_20260905/next_steps.md)。本文档中的会议内容是任务证据，执行范围来自用户“先完成会议里交代的事项”。
- [原PPT](../../input/iNKT/iNKT.pptx)，旧表提取及每页出处保留在PPT验证CSV的slide列。
- [提供的Blood论文](../../docs/blooda_adv-2024-014592-main.pdf)；[官方全文](https://pmc.ncbi.nlm.nih.gov/articles/PMC11869968/)；[官方补充材料](https://www.ebi.ac.uk/europepmc/webservices/rest/PMC11869968/supplementaryFiles)，本地sources/含S1–S8原Excel与补充方法PDF。
- [Dufva 2023 NK状态论文](https://www.sciencedirect.com/science/article/pii/S1074761323004909)；[Krovi 2020小鼠iNKT参考](https://www.nature.com/articles/s41467-020-20073-8)。用论文图中marker面板，不宣称使用完整作者分类器。
- [scVelo官方输入与方法说明](https://scvelo.readthedocs.io/en/stable/VelocityBasics.html)：需要spliced/unspliced计数，速度计算及投射不同于表达UMAP。
- [MGI同源映射](https://www.informatics.jax.org/downloads/reports/HOM_MouseHumanSequence.rpt)；Reactome官方API响应逐条冻结在sources/reactome_api/。
- 源库差异：S3中3条未能在冻结2019库精确匹配（Lipid and atherosclerosis、Viral protein interaction with cytokine and cytokine receptor、Chemical carcinogenesis - DNA adducts）；Caffeine metabolism测得成员太少，不强行检验。PPT中hsa04213未出现在冻结2021 Human库，不能用hsa04211替代。
- 入口脚本：notebooks/scripts/iNKT/run_inkt_meeting_followup.py、run_inkt_meeting_pathways.py、run_inkt_meeting_refinement.py、build_inkt_meeting_report.py。环境版本和文件SHA256见manifest.json。主流程入口为仓库根目录bash脚本；长任务均用tmux、CPU运行。
- [测试日志](logs/tests.log)：42项unittest通过；最终产物检查见[verification.json](verification.json)。
'''
    (OUT/'README.md').write_text(text)

def manifest():
    versions={p:metadata.version(p) for p in ['scanpy','anndata','gseapy','harmonypy','pypdf','pypdfium2','python-pptx','numpy','pandas','scipy','matplotlib']}
    files={str(p.relative_to(OUT)):sha256_file(p) for p in OUT.rglob('*') if p.is_file() and p.name not in ['manifest.json'] and not p.name.endswith('.log')}
    script_paths=set((ROOT/'notebooks/scripts/iNKT').glob('*meeting*.py'))|{ROOT/'notebooks/scripts/iNKT/run_inkt_c5_paper_followup.py',ROOT/'notebooks/scripts/iNKT/inkt_palette.py'}
    scripts={str(p.relative_to(ROOT)):sha256_file(p) for p in sorted(script_paths)}
    external=[ROOT/'TODOs.txt',ROOT/'input/iNKT/iNKT.pptx',ROOT/'docs/blooda_adv-2024-014592-main.pdf',OLD/'legacy_ppt_de_tables_extracted.csv',OLD/'legacy_ppt_pathway_tables_extracted.csv',BASE/'tables/20260830_KEGG_2019_Mouse.gmt']
    external_hashes={str(p.relative_to(ROOT)):sha256_file(p) for p in external}
    js(OUT/'manifest.json',{'created_UTC':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime()),'source_input':str(INPUT),'source_sha256':sha256_file(INPUT),'versions':versions,'scripts_sha256':scripts,'external_inputs_sha256':external_hashes,'artifacts_sha256':files,'GPU_used':False})
    (OUT/'sources/analysis_environment_versions.txt').write_text('\n'.join(f'{k}=={v}' for k,v in versions.items())+'\n')

if __name__=='__main__':
    log('report: stable ORA');stable_pathways();log('report: figures');pathway_plots();driver_plots();integration_plot();log('report: package');from plot_inkt_meeting_frequency import main as frequency_plot;frequency_plot();package();write_readme();manifest();js(OUT/'logs/report.completed.json',{'status':'complete'})
