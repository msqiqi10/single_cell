#!/usr/bin/env python3
"""Analyze curated iNKT signatures and selected biomarkers.

This script connects the single-column gene lists in
``input/iNKT/iNKT gene list for scRNAseq.xlsx`` with the processed iNKT
AnnData object.  It deliberately does not modify or rewrite the H5AD file.

The compute tasks write to separate directories and can therefore run
concurrently in separate tmux panes/processes::

    python analyze_inkt_gene_signatures.py --task gene-sets
    python analyze_inkt_gene_signatures.py --task signature-scores --threads 4
    python analyze_inkt_gene_signatures.py --task standardized-legacy --threads 4
    python analyze_inkt_gene_signatures.py --task marker-validation --threads 4

Use ``--validate-only`` first for a fast, backed-mode input/schema check.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
import unicodedata
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable, Sequence

import anndata as ad
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
import scanpy as sc
import scipy.sparse as sp
import seaborn as sns

try:
    from threadpoolctl import threadpool_limits
except ImportError:  # pragma: no cover - pulled in by scikit-learn in the project env
    threadpool_limits = None


ROOT = Path(__file__).resolve().parents[3]
DEFAULT_H5AD = ROOT / "output/iNKT_scanpy_tutorial_run/inkt_scanpy_tutorial_processed.h5ad"
DEFAULT_GENE_LIST = ROOT / "input/iNKT/iNKT gene list for scRNAseq.xlsx"
DEFAULT_OUT_DIR = ROOT / "output/iNKT_xlsx_reanalysis"
DEFAULT_CLUSTER_KEY = "leiden_res_0_5"
LEGACY_SUBTYPE_ORDER = ("inkt1", "inkt2", "inkt17")


# Only high-confidence mouse-symbol aliases/typos are changed automatically.
# Ambiguous items (for example Slam4 and Opct) are retained and reported as
# unresolved if they cannot be found in the dataset.
GENE_ALIASES = {
    "SEPP1": "Selenop",
    "CD103": "Itgae",
    "CD244": "Cd244a",
    "SLAM7": "Slamf7",
    "CKD6": "Cdk6",
    "HNMPA1": "Hnrnpa1",
    "DFNA5": "Gsdme",
    "ITGAB1": "Itgb1",
    "ITGAB4": "Itgb4",
    "ITGAB5": "Itgb5",
    "IL23RB": "Il23r",
}
AMBIGUOUS_UNMAPPED = ["Slam4", "Opct"]


DEFAULT_MARKERS = [
    "Xcl1",
    "Il4",
    "Ifng",
    "Rorc",
    "Il17a",
    "Il17f",
    "Gzma",
    "Gzmb",
    "Prf1",
    "Fasl",
    "Klf2",
    "Zbtb16",
]


@dataclass(frozen=True)
class SignatureDescriptor:
    sheet: str
    set_id: str
    label: str
    subtype: str
    source: str


def json_safe(value):
    """Convert numpy/pandas values into strict-JSON-compatible values."""

    if value is None or value is pd.NA:
        return None
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        return None if not math.isfinite(float(value)) else float(value)
    if isinstance(value, (np.bool_,)):
        return bool(value)
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, dict):
        return {str(k): json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_safe(v) for v in value]
    return value


def write_json(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(json_safe(payload), indent=2, sort_keys=True, allow_nan=False) + "\n",
        encoding="utf-8",
    )


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_gene_token(value: object) -> str | None:
    """Trim workbook artifacts while preserving the submitted gene symbol."""

    if value is None or isinstance(value, (float, np.floating)) and np.isnan(value):
        return None
    text = unicodedata.normalize("NFKC", str(value)).replace("\u00a0", " ").strip()
    text = re.sub(r"\s+", "", text)
    return text or None


def describe_signature_sheet(sheet: str) -> SignatureDescriptor | None:
    """Classify the 12 single-column signature sheets in the workbook."""

    lower = " ".join(sheet.strip().lower().split())
    if lower in {"inkt1", "inkt2", "inkt17"}:
        subtype = lower
        return SignatureDescriptor(sheet, f"{subtype}_literature", sheet, subtype, "literature")

    match = re.fullmatch(r"my list (inkt1|inkt2|inkt17)", lower)
    if match:
        subtype = match.group(1)
        return SignatureDescriptor(sheet, f"{subtype}_inhouse", sheet, subtype, "inhouse")

    match = re.match(r"(inkt1|inkt2|inkt17) \(wang et al\.,? 2022\)", lower)
    if match is None:
        match = re.fullmatch(r"wang\s*2022\s+(inkt1|inkt2|inkt17)", lower)
    if match:
        subtype = match.group(1)
        return SignatureDescriptor(sheet, f"{subtype}_wang_2022", sheet, subtype, "wang_2022")

    if lower.startswith("circulatory"):
        return SignatureDescriptor(
            sheet,
            "circulatory_wang_2022",
            "Circulatory markers",
            "circulatory",
            "wang_2022",
        )
    if lower.startswith("direct tcr activation"):
        return SignatureDescriptor(
            sheet,
            "direct_tcr_activation_curated",
            "Direct TCR activation",
            "direct_tcr_activation",
            "curated",
        )
    if lower.startswith("residency"):
        return SignatureDescriptor(
            sheet,
            "residency_wang_2022",
            "Residency markers",
            "residency",
            "wang_2022",
        )
    return None


def build_casefold_lookup(var_names: Iterable[str]) -> dict[str, str]:
    lookup: dict[str, str] = {}
    for name in map(str, var_names):
        lookup.setdefault(name.upper(), name)
    return lookup


def load_signature_members(
    workbook_path: Path,
    dataset_var_names: Sequence[str],
) -> tuple[pd.DataFrame, list[SignatureDescriptor]]:
    """Load, normalize, alias-map, and dataset-match all signature genes."""

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    descriptors = [
        descriptor
        for sheet in workbook.sheetnames
        if (descriptor := describe_signature_sheet(sheet)) is not None
    ]
    if len(descriptors) != 12:
        found = [descriptor.sheet for descriptor in descriptors]
        raise ValueError(f"Expected 12 signature sheets, found {len(descriptors)}: {found}")

    present = build_casefold_lookup(dataset_var_names)
    rows: list[dict] = []
    for descriptor in descriptors:
        seen: set[str] = set()
        worksheet = workbook[descriptor.sheet]
        for row_number, row in enumerate(worksheet.iter_rows(min_col=1, max_col=1, values_only=True), 1):
            original = row[0]
            cleaned = clean_gene_token(original)
            if cleaned is None:
                continue
            alias_target = GENE_ALIASES.get(cleaned.upper())
            canonical = alias_target or cleaned
            dataset_gene = present.get(canonical.upper())
            dedupe_key = (dataset_gene or canonical).upper()
            duplicate = dedupe_key in seen
            seen.add(dedupe_key)
            if duplicate:
                status = "duplicate_after_normalization"
            elif dataset_gene is None:
                status = "not_in_dataset"
            else:
                status = "matched"
            rows.append(
                {
                    **asdict(descriptor),
                    "row_number": row_number,
                    "gene_original": str(original),
                    "gene_clean": cleaned,
                    "gene_canonical": canonical,
                    "alias_applied": alias_target is not None,
                    "dataset_gene": dataset_gene,
                    "status": status,
                }
            )
    workbook.close()
    members = pd.DataFrame(rows)
    return members, descriptors


def genes_by_set(members: pd.DataFrame, *, measured_only: bool) -> dict[str, list[str]]:
    gene_column = "dataset_gene" if measured_only else "gene_canonical"
    status_mask = members["status"].eq("matched") if measured_only else ~members["status"].eq(
        "duplicate_after_normalization"
    )
    result: dict[str, list[str]] = {}
    for set_id, frame in members.loc[status_mask].groupby("set_id", sort=False):
        result[str(set_id)] = list(dict.fromkeys(frame[gene_column].dropna().astype(str)))
    return result


def signature_qc(members: pd.DataFrame, descriptors: Sequence[SignatureDescriptor]) -> pd.DataFrame:
    rows = []
    for descriptor in descriptors:
        frame = members.loc[members["set_id"].eq(descriptor.set_id)]
        matched = frame.loc[frame["status"].eq("matched"), "dataset_gene"].dropna().unique()
        canonical = frame.loc[
            ~frame["status"].eq("duplicate_after_normalization"), "gene_canonical"
        ].unique()
        unresolved = frame.loc[frame["status"].eq("not_in_dataset"), "gene_canonical"].unique()
        rows.append(
            {
                **asdict(descriptor),
                "n_rows": int(len(frame)),
                "n_canonical_unique": int(len(canonical)),
                "n_matched_unique": int(len(matched)),
                "match_fraction": float(len(matched) / len(canonical)) if len(canonical) else np.nan,
                "n_aliases_applied": int(frame["alias_applied"].sum()),
                "n_duplicate_rows": int(frame["status"].eq("duplicate_after_normalization").sum()),
                "unresolved_genes": ";".join(sorted(map(str, unresolved), key=str.upper)),
            }
        )
    return pd.DataFrame(rows)


def pairwise_overlaps(
    descriptors: Sequence[SignatureDescriptor],
    gene_sets: dict[str, list[str]],
    scope: str,
) -> pd.DataFrame:
    descriptor_by_id = {descriptor.set_id: descriptor for descriptor in descriptors}
    rows = []
    for index, set_a in enumerate(gene_sets):
        genes_a = set(map(str.upper, gene_sets[set_a]))
        for set_b in list(gene_sets)[index:]:
            genes_b = set(map(str.upper, gene_sets[set_b]))
            intersection = genes_a & genes_b
            union = genes_a | genes_b
            rows.append(
                {
                    "scope": scope,
                    "set_a": set_a,
                    "set_b": set_b,
                    "subtype_a": descriptor_by_id[set_a].subtype,
                    "source_a": descriptor_by_id[set_a].source,
                    "subtype_b": descriptor_by_id[set_b].subtype,
                    "source_b": descriptor_by_id[set_b].source,
                    "n_a": len(genes_a),
                    "n_b": len(genes_b),
                    "n_intersection": len(intersection),
                    "n_union": len(union),
                    "jaccard": len(intersection) / len(union) if union else np.nan,
                    "intersection_genes_upper": ";".join(sorted(intersection)),
                }
            )
    return pd.DataFrame(rows)


def subtype_membership_table(
    descriptors: Sequence[SignatureDescriptor],
    canonical_sets: dict[str, list[str]],
) -> pd.DataFrame:
    """Make an UpSet-compatible membership table for the 3x3 subtype sets."""

    rows = []
    for subtype in ("inkt1", "inkt2", "inkt17"):
        selected = [descriptor for descriptor in descriptors if descriptor.subtype == subtype]
        representative: dict[str, str] = {}
        for descriptor in selected:
            for gene in canonical_sets[descriptor.set_id]:
                representative.setdefault(gene.upper(), gene)
        union = [representative[key] for key in sorted(representative)]
        upper_sets = {
            descriptor.source: set(map(str.upper, canonical_sets[descriptor.set_id]))
            for descriptor in selected
        }
        for gene in union:
            rows.append(
                {
                    "subtype": subtype,
                    "gene": gene,
                    **{source: gene.upper() in upper_sets.get(source, set()) for source in (
                        "literature",
                        "inhouse",
                        "wang_2022",
                    )},
                }
            )
    return pd.DataFrame(rows)


def plot_jaccard(overlaps: pd.DataFrame, output_path: Path) -> None:
    nine_ids = [
        f"{subtype}_{source}"
        for subtype in ("inkt1", "inkt2", "inkt17")
        for source in ("literature", "inhouse", "wang_2022")
    ]
    frame = overlaps.loc[
        overlaps["scope"].eq("canonical_all")
        & overlaps["set_a"].isin(nine_ids)
        & overlaps["set_b"].isin(nine_ids)
    ]
    matrix = pd.DataFrame(np.eye(len(nine_ids)), index=nine_ids, columns=nine_ids)
    for row in frame.itertuples(index=False):
        matrix.loc[row.set_a, row.set_b] = row.jaccard
        matrix.loc[row.set_b, row.set_a] = row.jaccard
    plt.figure(figsize=(10, 8))
    sns.heatmap(matrix, cmap="viridis", vmin=0, vmax=1, annot=True, fmt=".2f", square=True)
    plt.title("iNKT subtype signature overlap (canonical genes)")
    plt.tight_layout()
    plt.savefig(output_path, dpi=180, bbox_inches="tight")
    plt.close()


def plot_membership_counts(membership: pd.DataFrame, output_path: Path) -> None:
    sources = ["literature", "inhouse", "wang_2022"]
    pattern_counts = (
        membership.groupby(["subtype", *sources], observed=True)
        .size()
        .rename("n_genes")
        .reset_index()
    )
    pattern_counts["membership"] = pattern_counts.apply(
        lambda row: " & ".join(source for source in sources if bool(row[source])) or "none", axis=1
    )
    pivot = pattern_counts.pivot_table(
        index="membership", columns="subtype", values="n_genes", fill_value=0
    )
    pivot = pivot.reindex(columns=["inkt1", "inkt2", "inkt17"], fill_value=0)
    pivot.plot(kind="barh", figsize=(9, max(4, 0.45 * len(pivot))), width=0.8)
    plt.xlabel("Canonical genes")
    plt.ylabel("Source membership")
    plt.title("iNKT1/2/17 signature membership patterns")
    plt.legend(title="Subtype")
    plt.tight_layout()
    plt.savefig(output_path, dpi=180, bbox_inches="tight")
    plt.close()


def validate_inputs(args: argparse.Namespace) -> dict:
    for path in (args.h5ad, args.gene_list):
        if not path.is_file():
            raise FileNotFoundError(path)

    adata = sc.read_h5ad(args.h5ad, backed="r")
    try:
        required_obs = {"sample", "condition", "tissue", args.cluster_key}
        missing_obs = sorted(required_obs - set(adata.obs.columns))
        if missing_obs:
            raise KeyError(f"H5AD is missing obs columns: {missing_obs}")
        var_names = list(map(str, adata.raw.var_names if adata.raw is not None else adata.var_names))
        members, descriptors = load_signature_members(args.gene_list, var_names)
        qc = signature_qc(members, descriptors)
        payload = {
            "h5ad": str(args.h5ad),
            "shape": [int(adata.n_obs), int(adata.n_vars)],
            "raw_shape": None if adata.raw is None else [int(adata.raw.n_obs), int(adata.raw.n_vars)],
            "layers": list(adata.layers.keys()),
            "required_obs": sorted(required_obs),
            "condition_values": sorted(adata.obs["condition"].astype(str).unique()),
            "tissue_values": sorted(adata.obs["tissue"].astype(str).unique()),
            "cluster_values": sorted(adata.obs[args.cluster_key].astype(str).unique()),
            "signature_sheet_count": len(descriptors),
            "signature_sets": qc[["set_id", "n_canonical_unique", "n_matched_unique"]].to_dict(
                orient="records"
            ),
        }
    finally:
        adata.file.close()
    return payload


def run_gene_sets(args: argparse.Namespace) -> dict:
    output_dir = args.out_dir / "gene_sets"
    output_dir.mkdir(parents=True, exist_ok=True)

    adata = sc.read_h5ad(args.h5ad, backed="r")
    try:
        var_names = list(map(str, adata.raw.var_names if adata.raw is not None else adata.var_names))
    finally:
        adata.file.close()

    members, descriptors = load_signature_members(args.gene_list, var_names)
    qc = signature_qc(members, descriptors)
    canonical_sets = genes_by_set(members, measured_only=False)
    measured_sets = genes_by_set(members, measured_only=True)
    overlaps = pd.concat(
        [
            pairwise_overlaps(descriptors, canonical_sets, "canonical_all"),
            pairwise_overlaps(descriptors, measured_sets, "measured_in_h5ad"),
        ],
        ignore_index=True,
    )
    membership = subtype_membership_table(descriptors, canonical_sets)

    members.to_csv(output_dir / "signature_gene_members.csv", index=False)
    qc.to_csv(output_dir / "signature_qc.csv", index=False)
    overlaps.to_csv(output_dir / "pairwise_signature_overlap.csv", index=False)
    membership.to_csv(output_dir / "subtype_gene_membership.csv", index=False)

    with (output_dir / "signatures_measured.gmt").open("w", encoding="utf-8") as handle:
        for descriptor in descriptors:
            genes = measured_sets[descriptor.set_id]
            handle.write("\t".join([descriptor.set_id, descriptor.sheet, *genes]) + "\n")

    write_json(
        output_dir / "signature_gene_sets.json",
        {
            descriptor.set_id: {
                **asdict(descriptor),
                "canonical_genes": canonical_sets[descriptor.set_id],
                "measured_genes": measured_sets[descriptor.set_id],
            }
            for descriptor in descriptors
        },
    )
    plot_jaccard(overlaps, output_dir / "signature_jaccard_heatmap.png")
    plot_membership_counts(membership, output_dir / "subtype_signature_membership.png")

    summary = {
        "task": "gene-sets",
        "input_gene_list": str(args.gene_list),
        "input_h5ad": str(args.h5ad),
        "n_signature_sets": len(descriptors),
        "n_subtype_canonical_unique": int(
            membership.loc[:, "gene"].str.upper().nunique()
        ),
        "all_signature_canonical_unique": int(
            members.loc[
                ~members["status"].eq("duplicate_after_normalization"), "gene_canonical"
            ].str.upper().nunique()
        ),
        "ambiguous_symbols_not_automatically_remapped": AMBIGUOUS_UNMAPPED,
        "outputs": sorted(path.name for path in output_dir.iterdir()),
    }
    write_json(output_dir / "summary.json", summary)
    return summary


def add_module_scores(
    adata: ad.AnnData,
    descriptors: Sequence[SignatureDescriptor],
    measured_sets: dict[str, list[str]],
    *,
    seed: int,
) -> list[str]:
    score_columns = []
    gene_pool = list(map(str, adata.raw.var_names if adata.raw is not None else adata.var_names))
    for offset, descriptor in enumerate(descriptors):
        genes = measured_sets[descriptor.set_id]
        if len(genes) < 2:
            raise ValueError(f"Too few measured genes for {descriptor.set_id}: {genes}")
        score_column = f"score__{descriptor.set_id}"
        sc.tl.score_genes(
            adata,
            gene_list=genes,
            gene_pool=gene_pool,
            ctrl_size=max(50, len(genes)),
            n_bins=25,
            score_name=score_column,
            random_state=seed + offset,
            use_raw=adata.raw is not None,
        )
        score_columns.append(score_column)
    return score_columns


def summarize_scores(
    obs: pd.DataFrame,
    score_columns: Sequence[str],
    descriptor_by_column: dict[str, SignatureDescriptor],
    cluster_key: str,
) -> pd.DataFrame:
    grouping_specs = [
        ("global", ["condition"]),
        ("tissue", ["tissue", "condition"]),
        ("cluster", [cluster_key, "condition"]),
        ("tissue_cluster", ["tissue", cluster_key, "condition"]),
    ]
    rows = []
    for stratum_type, keys in grouping_specs:
        for key_values, positions in obs.groupby(keys, observed=True, sort=True).indices.items():
            if not isinstance(key_values, tuple):
                key_values = (key_values,)
            labels = dict(zip(keys, map(str, key_values), strict=True))
            for column in score_columns:
                values = obs.iloc[positions][column].to_numpy(dtype=float)
                descriptor = descriptor_by_column[column]
                rows.append(
                    {
                        "stratum_type": stratum_type,
                        "tissue": labels.get("tissue"),
                        "cluster": labels.get(cluster_key),
                        "condition": labels["condition"],
                        **asdict(descriptor),
                        "n_cells": int(len(values)),
                        "mean_score": float(np.mean(values)),
                        "median_score": float(np.median(values)),
                        "std_score": float(np.std(values, ddof=1)) if len(values) > 1 else np.nan,
                    }
                )
    return pd.DataFrame(rows)


def condition_effects(
    summary: pd.DataFrame,
    value_columns: Sequence[str],
) -> pd.DataFrame:
    """Create descriptive T2-minus-Ctrl effects from a grouped summary."""

    measurement_columns = set(value_columns) | {"std_score"}
    id_columns = [
        column
        for column in summary.columns
        if column not in {"condition", "n_cells", *measurement_columns}
    ]
    rows = []
    for labels, frame in summary.groupby(id_columns, dropna=False, observed=True, sort=False):
        if not isinstance(labels, tuple):
            labels = (labels,)
        by_condition = frame.set_index("condition")
        if not {"Ctrl", "T2"}.issubset(by_condition.index):
            continue
        row = dict(zip(id_columns, labels, strict=True))
        row["n_ctrl"] = int(by_condition.loc["Ctrl", "n_cells"])
        row["n_t2"] = int(by_condition.loc["T2", "n_cells"])
        for value in value_columns:
            ctrl = float(by_condition.loc["Ctrl", value])
            t2 = float(by_condition.loc["T2", value])
            row[f"ctrl_{value}"] = ctrl
            row[f"t2_{value}"] = t2
            row[f"delta_{value}_t2_minus_ctrl"] = t2 - ctrl
        row["inference_note"] = "descriptive cell-level effect; no biological-replicate p-value"
        rows.append(row)
    return pd.DataFrame(rows)


def make_cluster_evidence(
    obs: pd.DataFrame,
    descriptors: Sequence[SignatureDescriptor],
    cluster_key: str,
    set_sizes: dict[str, int],
) -> pd.DataFrame:
    rows = []
    subtype_descriptors = [
        descriptor for descriptor in descriptors if descriptor.subtype in {"inkt1", "inkt2", "inkt17"}
    ]
    for source in ("literature", "inhouse", "wang_2022"):
        source_descriptors = [d for d in subtype_descriptors if d.source == source]
        for cluster, positions in obs.groupby(cluster_key, observed=True).indices.items():
            means = {
                descriptor.subtype: float(
                    obs.iloc[positions][f"score__{descriptor.set_id}"].mean()
                )
                for descriptor in source_descriptors
            }
            ordered = sorted(means.items(), key=lambda item: item[1], reverse=True)
            winner, winning_score = ordered[0]
            margin = winning_score - ordered[1][1]
            for descriptor in source_descriptors:
                rows.append(
                    {
                        "cluster": str(cluster),
                        "source": source,
                        "subtype": descriptor.subtype,
                        "set_id": descriptor.set_id,
                        "n_measured_genes": set_sizes[descriptor.set_id],
                        "mean_score": means[descriptor.subtype],
                        "winning_subtype": winner,
                        "winning_margin": margin,
                        "is_winner": descriptor.subtype == winner,
                    }
                )
    return pd.DataFrame(rows)


def make_cell_subtype_calls(
    obs: pd.DataFrame,
    descriptors: Sequence[SignatureDescriptor],
) -> pd.DataFrame:
    frames = []
    for source in ("literature", "inhouse", "wang_2022"):
        selected = [
            descriptor
            for descriptor in descriptors
            if descriptor.source == source and descriptor.subtype in {"inkt1", "inkt2", "inkt17"}
        ]
        columns = [f"score__{descriptor.set_id}" for descriptor in selected]
        values = obs[columns].to_numpy(dtype=float)
        order = np.argsort(values, axis=1)
        winner_index = order[:, -1]
        runner_up_index = order[:, -2]
        subtypes = np.asarray([descriptor.subtype for descriptor in selected])
        frames.append(
            pd.DataFrame(
                {
                    "cell_id": obs.index.astype(str),
                    "source": source,
                    "subtype_call": subtypes[winner_index],
                    "winning_score": values[np.arange(len(values)), winner_index],
                    "margin_to_second": values[np.arange(len(values)), winner_index]
                    - values[np.arange(len(values)), runner_up_index],
                }
            )
        )
    return pd.concat(frames, ignore_index=True)


def plot_score_summaries(
    summary: pd.DataFrame,
    effects: pd.DataFrame,
    evidence: pd.DataFrame,
    output_dir: Path,
) -> None:
    cluster = summary.loc[
        summary["stratum_type"].eq("cluster") & summary["condition"].eq("Ctrl")
    ].copy()
    t2 = summary.loc[
        summary["stratum_type"].eq("cluster") & summary["condition"].eq("T2")
    ].copy()
    combined = pd.concat(
        [cluster.assign(label=lambda frame: "Ctrl c" + frame["cluster"].astype(str)),
         t2.assign(label=lambda frame: "T2 c" + frame["cluster"].astype(str))]
    )
    matrix = combined.pivot(index="set_id", columns="label", values="mean_score")
    plt.figure(figsize=(max(10, 0.55 * len(matrix.columns)), 7))
    sns.heatmap(matrix, cmap="vlag", center=0)
    plt.title("Mean signature scores by condition and current cluster")
    plt.tight_layout()
    plt.savefig(output_dir / "signature_scores_by_cluster_condition.png", dpi=180, bbox_inches="tight")
    plt.close()

    tissue_effects = effects.loc[effects["stratum_type"].isin(["global", "tissue"])].copy()
    tissue_effects["stratum"] = tissue_effects["tissue"].fillna("global")
    effect_matrix = tissue_effects.pivot(
        index="set_id", columns="stratum", values="delta_mean_score_t2_minus_ctrl"
    )
    ordered_columns = [column for column in ["global", "bone_marrow", "spleen", "thymus"] if column in effect_matrix]
    plt.figure(figsize=(7, 7))
    sns.heatmap(effect_matrix[ordered_columns], cmap="vlag", center=0, annot=True, fmt=".2f")
    plt.title("Signature score effect: T2 minus Ctrl")
    plt.tight_layout()
    plt.savefig(output_dir / "signature_t2_ctrl_effect_heatmap.png", dpi=180, bbox_inches="tight")
    plt.close()

    figure, axes = plt.subplots(1, 3, figsize=(18, 5), sharey=True)
    for axis, source in zip(axes, ("literature", "inhouse", "wang_2022"), strict=True):
        frame = evidence.loc[evidence["source"].eq(source)]
        matrix = frame.pivot(index="subtype", columns="cluster", values="mean_score").reindex(
            ["inkt1", "inkt2", "inkt17"]
        )
        sns.heatmap(
            matrix,
            cmap="vlag",
            center=0,
            annot=True,
            fmt=".2f",
            ax=axis,
            cbar=True,
            cbar_kws={"label": "raw score_genes"},
        )
        axis.set_title(source)
        axis.set_xlabel("current cluster")
    axes[0].set_ylabel("subtype")
    figure.suptitle("Raw score_genes sensitivity across sources (independent panel scales)")
    figure.tight_layout()
    figure.savefig(output_dir / "cluster_subtype_evidence.png", dpi=180, bbox_inches="tight")
    plt.close(figure)


def run_signature_scores(args: argparse.Namespace) -> dict:
    output_dir = args.out_dir / "signature_scores"
    output_dir.mkdir(parents=True, exist_ok=True)
    adata = sc.read_h5ad(args.h5ad)
    var_names = list(map(str, adata.raw.var_names if adata.raw is not None else adata.var_names))
    members, descriptors = load_signature_members(args.gene_list, var_names)
    measured_sets = genes_by_set(members, measured_only=True)
    score_columns = add_module_scores(adata, descriptors, measured_sets, seed=args.seed)
    descriptor_by_column = {f"score__{descriptor.set_id}": descriptor for descriptor in descriptors}

    meta_columns = ["sample", "condition", "tissue", args.cluster_key]
    scores = adata.obs[meta_columns + score_columns].copy()
    scores.index.name = "cell_id"
    scores.to_csv(output_dir / "signature_scores_per_cell.csv.gz", compression="gzip")
    summary = summarize_scores(scores, score_columns, descriptor_by_column, args.cluster_key)
    effects = condition_effects(summary, ["mean_score", "median_score"])
    evidence = make_cluster_evidence(
        scores,
        descriptors,
        args.cluster_key,
        {set_id: len(genes) for set_id, genes in measured_sets.items()},
    )
    calls = make_cell_subtype_calls(scores, descriptors)

    summary.to_csv(output_dir / "signature_score_summary.csv", index=False)
    effects.to_csv(output_dir / "signature_t2_vs_ctrl_effects.csv", index=False)
    evidence.to_csv(output_dir / "cluster_subtype_evidence.csv", index=False)
    calls.to_csv(output_dir / "cell_subtype_calls.csv.gz", index=False, compression="gzip")
    plot_score_summaries(summary, effects, evidence, output_dir)

    task_summary = {
        "task": "signature-scores",
        "n_cells": int(adata.n_obs),
        "n_sets_scored": len(score_columns),
        "score_method": "scanpy.tl.score_genes on raw log-normalized expression",
        "random_seed": args.seed,
        "cluster_key": args.cluster_key,
        "h5ad_was_modified": False,
        "outputs": sorted(path.name for path in output_dir.iterdir()),
    }
    write_json(output_dir / "summary.json", task_summary)
    return task_summary


def expression_matrix(adata: ad.AnnData, genes: Sequence[str]) -> tuple[np.ndarray, np.ndarray]:
    source = adata.raw if adata.raw is not None else adata
    matrix = source[:, list(genes)].X
    if sp.issparse(matrix):
        matrix = matrix.toarray()
    log_expression = np.asarray(matrix, dtype=np.float64)
    normalized_expression = np.expm1(log_expression)
    return log_expression, normalized_expression


def standardized_mean_z_scores(
    log_expression: np.ndarray,
    gene_names: Sequence[str],
    genes_by_subtype: dict[str, Sequence[str]],
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Put unequal legacy subtype lists on one explicitly shared score scale.

    First, every gene is z-scored across all cells.  Each subtype receives the
    unweighted mean of its measured gene z-scores.  Finally, each resulting
    signature score is centered and scaled across all cells.  The last step
    makes cluster means comparable in standard-deviation units without forcing
    the three biological gene lists to have equal lengths.
    """

    values = np.asarray(log_expression, dtype=np.float64)
    if values.ndim != 2:
        raise ValueError("log_expression must be a two-dimensional matrix")
    if values.shape[1] != len(gene_names):
        raise ValueError("gene_names length does not match expression columns")
    if not np.isfinite(values).all():
        raise ValueError("log_expression contains non-finite values")

    lookup = {str(gene).upper(): index for index, gene in enumerate(gene_names)}
    gene_mean = values.mean(axis=0)
    gene_std = values.std(axis=0, ddof=0)
    variable = gene_std > 0
    safe_std = gene_std.copy()
    safe_std[~variable] = 1.0
    gene_z = (values - gene_mean) / safe_std

    score_columns: dict[str, np.ndarray] = {}
    signature_rows = []
    used_by_any: set[int] = set()
    for subtype in LEGACY_SUBTYPE_ORDER:
        requested = list(genes_by_subtype[subtype])
        indices = [lookup[str(gene).upper()] for gene in requested if str(gene).upper() in lookup]
        variable_indices = [index for index in indices if variable[index]]
        if not variable_indices:
            raise ValueError(f"No variable measured genes for legacy {subtype}")
        used_by_any.update(variable_indices)
        mean_gene_z = gene_z[:, variable_indices].mean(axis=1)
        score_mean = float(mean_gene_z.mean())
        score_std = float(mean_gene_z.std(ddof=0))
        if not math.isfinite(score_std) or score_std <= 0:
            raise ValueError(f"Degenerate standardized score for legacy {subtype}")
        standardized = (mean_gene_z - score_mean) / score_std
        score_columns[subtype] = standardized
        signature_rows.append(
            {
                "subtype": subtype,
                "n_input_measured_genes": len(requested),
                "n_mapped_genes": len(indices),
                "n_variable_genes": len(variable_indices),
                "pre_standardization_mean": score_mean,
                "pre_standardization_std": score_std,
                "standardized_mean": float(standardized.mean()),
                "standardized_std": float(standardized.std(ddof=0)),
            }
        )

    gene_rows = [
        {
            "gene": str(gene),
            "mean_log1p_expression": float(gene_mean[index]),
            "std_log1p_expression": float(gene_std[index]),
            "variable_for_scoring": bool(index in used_by_any),
        }
        for index, gene in enumerate(gene_names)
    ]
    return (
        pd.DataFrame(score_columns),
        pd.DataFrame(gene_rows),
        pd.DataFrame(signature_rows),
    )


def summarize_standardized_scores(
    metadata: pd.DataFrame,
    scores: pd.DataFrame,
    cluster_key: str,
) -> pd.DataFrame:
    frame = pd.concat([metadata.reset_index(drop=True), scores.reset_index(drop=True)], axis=1)
    rows = []
    grouping_specs = [
        ("cluster", [cluster_key]),
        ("cluster_condition", [cluster_key, "condition"]),
    ]
    for stratum_type, keys in grouping_specs:
        for key_values, positions in frame.groupby(keys, observed=True, sort=True).indices.items():
            if not isinstance(key_values, tuple):
                key_values = (key_values,)
            labels = dict(zip(keys, map(str, key_values), strict=True))
            for subtype in LEGACY_SUBTYPE_ORDER:
                values = frame.iloc[positions][subtype].to_numpy(dtype=float)
                rows.append(
                    {
                        "stratum_type": stratum_type,
                        "cluster": labels[cluster_key],
                        "condition": labels.get("condition"),
                        "subtype": subtype,
                        "n_cells": len(values),
                        "mean_standardized_score": float(values.mean()),
                        "median_standardized_score": float(np.median(values)),
                        "std_standardized_score": (
                            float(values.std(ddof=1)) if len(values) > 1 else np.nan
                        ),
                    }
                )
    return pd.DataFrame(rows)


def bootstrap_standardized_cluster_calls(
    metadata: pd.DataFrame,
    scores: pd.DataFrame,
    cluster_key: str,
    *,
    n_bootstrap: int,
    seed: int,
    batch_size: int = 100,
    strata_key: str | None = None,
) -> pd.DataFrame:
    """Call a subtype only when cell-bootstrap evidence clears zero and runner-up.

    These intervals describe cell-level stability inside this dataset.  They do
    not replace biological-replicate uncertainty, which cannot be estimated
    from the available one-sample-per-tissue-condition design.
    """

    if n_bootstrap < 1:
        raise ValueError("n_bootstrap must be >= 1")
    frame = pd.concat([metadata.reset_index(drop=True), scores.reset_index(drop=True)], axis=1)
    rng = np.random.default_rng(seed)
    rows = []
    for cluster, positions in frame.groupby(cluster_key, observed=True, sort=True).indices.items():
        cluster_frame = frame.iloc[positions]
        values = cluster_frame[list(LEGACY_SUBTYPE_ORDER)].to_numpy(dtype=float)
        means = values.mean(axis=0)
        order = np.argsort(means)[::-1]
        winner_index = int(order[0])
        runner_index = int(order[1])
        bootstrap_means = np.empty((n_bootstrap, len(LEGACY_SUBTYPE_ORDER)), dtype=np.float64)
        if strata_key is None:
            strata = [np.arange(len(values), dtype=int)]
        else:
            if strata_key not in cluster_frame.columns:
                raise KeyError(f"Bootstrap strata column is missing: {strata_key}")
            strata = [
                np.asarray(indices, dtype=int)
                for indices in cluster_frame.reset_index(drop=True)
                .groupby(strata_key, observed=True, sort=True)
                .indices.values()
            ]
        for start in range(0, n_bootstrap, batch_size):
            stop = min(start + batch_size, n_bootstrap)
            batch_total = np.zeros((stop - start, len(LEGACY_SUBTYPE_ORDER)), dtype=np.float64)
            for stratum_positions in strata:
                draws = rng.integers(
                    0,
                    len(stratum_positions),
                    size=(stop - start, len(stratum_positions)),
                )
                batch_total += values[stratum_positions[draws]].sum(axis=1)
            bootstrap_means[start:stop] = batch_total / len(values)
        winner_samples = bootstrap_means[:, winner_index]
        competitor_samples = np.max(
            np.delete(bootstrap_means, winner_index, axis=1),
            axis=1,
        )
        margin_samples = winner_samples - competitor_samples
        winner_ci_low, winner_ci_high = np.quantile(winner_samples, [0.025, 0.975])
        margin_ci_low, margin_ci_high = np.quantile(margin_samples, [0.025, 0.975])
        positive = bool(winner_ci_low > 0)
        separated = bool(margin_ci_low > 0)
        winner = LEGACY_SUBTYPE_ORDER[winner_index]
        runner_up = LEGACY_SUBTYPE_ORDER[runner_index]
        row = {
            "cluster": str(cluster),
            "n_cells": len(values),
            "winner": winner,
            "runner_up": runner_up,
            "winner_mean": float(means[winner_index]),
            "winner_ci_low": float(winner_ci_low),
            "winner_ci_high": float(winner_ci_high),
            "margin_to_runner_up": float(means[winner_index] - means[runner_index]),
            "margin_ci_low": float(margin_ci_low),
            "margin_ci_high": float(margin_ci_high),
            "bootstrap_fraction_winner_above_zero": float(np.mean(winner_samples > 0)),
            "bootstrap_fraction_margin_above_zero": float(np.mean(margin_samples > 0)),
            "supported_call": (
                f"legacy-list {winner}-enriched"
                if positive and separated
                else "mixed"
                if positive
                else "unclassified"
            ),
            "bootstrap_strata": strata_key or "none",
            "inference_scope": (
                "sample-stratified cell-bootstrap stability; not biological-replicate inference"
                if strata_key
                else "cell-bootstrap stability; not biological-replicate inference"
            ),
        }
        row.update(
            {
                f"mean_{subtype}": float(means[index])
                for index, subtype in enumerate(LEGACY_SUBTYPE_ORDER)
            }
        )
        rows.append(row)
    return pd.DataFrame(rows)


def _cluster_order(values: Iterable[object]) -> list[str]:
    labels = list(dict.fromkeys(map(str, values)))
    return sorted(labels, key=lambda value: (not value.isdigit(), int(value) if value.isdigit() else value))


def plot_standardized_legacy_scores(
    adata: ad.AnnData,
    metadata: pd.DataFrame,
    scores: pd.DataFrame,
    summary: pd.DataFrame,
    output_dir: Path,
    cluster_key: str,
) -> dict[str, float]:
    score_values = scores[list(LEGACY_SUBTYPE_ORDER)].to_numpy(dtype=float)
    score_limit = max(1.0, float(np.ceil(np.quantile(np.abs(score_values), 0.99) * 2) / 2))
    cluster_frame = summary.loc[summary["stratum_type"].eq("cluster")]
    cluster_matrix = cluster_frame.pivot(
        index="subtype", columns="cluster", values="mean_standardized_score"
    ).reindex(LEGACY_SUBTYPE_ORDER)
    cluster_order = _cluster_order(cluster_matrix.columns)
    cluster_matrix = cluster_matrix.reindex(columns=cluster_order)
    condition_frame = summary.loc[summary["stratum_type"].eq("cluster_condition")].copy()
    condition_frame["label"] = condition_frame["condition"] + " c" + condition_frame["cluster"]
    condition_matrix = condition_frame.pivot(
        index="subtype", columns="label", values="mean_standardized_score"
    ).reindex(LEGACY_SUBTYPE_ORDER)
    condition_columns = [
        f"{condition} c{cluster}"
        for condition in ("Ctrl", "T2")
        for cluster in cluster_order
        if f"{condition} c{cluster}" in condition_matrix.columns
    ]
    condition_matrix = condition_matrix.reindex(columns=condition_columns)
    aggregate_limit = max(
        0.5,
        float(
            np.ceil(
                max(
                    np.nanmax(np.abs(cluster_matrix.to_numpy(dtype=float))),
                    np.nanmax(np.abs(condition_matrix.to_numpy(dtype=float))),
                )
                * 4
            )
            / 4
        ),
    )

    plt.figure(figsize=(10, 3.5))
    sns.heatmap(
        cluster_matrix,
        cmap="vlag",
        center=0,
        vmin=-aggregate_limit,
        vmax=aggregate_limit,
        annot=True,
        fmt=".2f",
        cbar_kws={"label": "standardized score"},
    )
    plt.xlabel("current cluster")
    plt.ylabel("legacy literature list")
    plt.title("Unified legacy-list subtype scores by current cluster")
    plt.tight_layout()
    plt.savefig(output_dir / "standardized_cluster_heatmap.png", dpi=180, bbox_inches="tight")
    plt.close()

    plt.figure(figsize=(13, 3.5))
    sns.heatmap(
        condition_matrix,
        cmap="vlag",
        center=0,
        vmin=-aggregate_limit,
        vmax=aggregate_limit,
        annot=True,
        fmt=".2f",
        cbar_kws={"label": "standardized score"},
    )
    plt.xlabel("condition and current cluster")
    plt.ylabel("legacy literature list")
    plt.title("Unified legacy-list scores by condition and current cluster")
    plt.tight_layout()
    plt.savefig(
        output_dir / "standardized_cluster_condition_heatmap.png", dpi=180, bbox_inches="tight"
    )
    plt.close()

    umap = np.asarray(adata.obsm["X_umap"], dtype=float)
    combined_figure, combined_axes = plt.subplots(1, 3, figsize=(12, 4), sharex=True, sharey=True)
    for axis, subtype in zip(combined_axes, LEGACY_SUBTYPE_ORDER, strict=True):
        subtype_values = scores[subtype].to_numpy(dtype=float)
        order = np.argsort(subtype_values)
        points = axis.scatter(
            umap[order, 0],
            umap[order, 1],
            c=subtype_values[order],
            cmap="vlag",
            vmin=-score_limit,
            vmax=score_limit,
            s=2.5,
            linewidths=0,
            rasterized=True,
        )
        axis.set_title(subtype.upper())
        axis.set_xlabel("UMAP1")
        axis.set_ylabel("UMAP2")
        axis.set_xticks([])
        axis.set_yticks([])
    combined_figure.colorbar(points, ax=combined_axes, fraction=0.025, pad=0.02, label="standardized score")
    combined_figure.suptitle("Current cells scored with the three legacy literature lists on one scale")
    combined_figure.subplots_adjust(left=0.04, right=0.92, bottom=0.08, top=0.84, wspace=0.08)
    combined_figure.savefig(output_dir / "standardized_umaps_all.png", dpi=180, bbox_inches="tight")
    plt.close(combined_figure)

    plot_frame = pd.concat([metadata.reset_index(drop=True), scores.reset_index(drop=True)], axis=1)
    plot_frame["cluster_label"] = "c" + plot_frame[cluster_key].astype(str)
    cluster_labels = [f"c{cluster}" for cluster in cluster_order]
    for subtype in LEGACY_SUBTYPE_ORDER:
        figure, axes = plt.subplots(1, 2, figsize=(9, 4.2), gridspec_kw={"width_ratios": [1.0, 1.25]})
        subtype_values = scores[subtype].to_numpy(dtype=float)
        order = np.argsort(subtype_values)
        points = axes[0].scatter(
            umap[order, 0],
            umap[order, 1],
            c=subtype_values[order],
            cmap="vlag",
            vmin=-score_limit,
            vmax=score_limit,
            s=3.0,
            linewidths=0,
            rasterized=True,
        )
        axes[0].set_title("UMAP")
        axes[0].set_xticks([])
        axes[0].set_yticks([])
        figure.colorbar(points, ax=axes[0], fraction=0.046, pad=0.03)
        sns.violinplot(
            data=plot_frame,
            x="cluster_label",
            y=subtype,
            order=cluster_labels,
            cut=0,
            inner="quartile",
            linewidth=0.55,
            color="#73B7B8",
            ax=axes[1],
        )
        axes[1].axhline(0, color="#52616B", linewidth=0.8, linestyle="--")
        axes[1].set_title("Current cluster distributions")
        axes[1].set_xlabel("current cluster")
        axes[1].set_ylabel("standardized score")
        axes[1].set_ylim(-score_limit, score_limit)
        figure.suptitle(f"Legacy-list {subtype.upper()} score: unified method and scale")
        figure.tight_layout()
        figure.savefig(
            output_dir / f"standardized_{subtype}_umap_violin.png",
            dpi=180,
            bbox_inches="tight",
        )
        plt.close(figure)
    return {"cell_score_color_limit": score_limit, "aggregate_color_limit": aggregate_limit}


def run_standardized_legacy(args: argparse.Namespace) -> dict:
    output_dir = args.out_dir / "standardized_legacy"
    output_dir.mkdir(parents=True, exist_ok=True)
    adata = sc.read_h5ad(args.h5ad)
    var_names = list(map(str, adata.raw.var_names if adata.raw is not None else adata.var_names))
    members, descriptors = load_signature_members(args.gene_list, var_names)
    measured_sets = genes_by_set(members, measured_only=True)
    coverage = signature_qc(members, descriptors)
    selected = [
        descriptor
        for descriptor in descriptors
        if descriptor.source == "literature" and descriptor.subtype in LEGACY_SUBTYPE_ORDER
    ]
    selected_by_subtype = {descriptor.subtype: descriptor for descriptor in selected}
    if set(selected_by_subtype) != set(LEGACY_SUBTYPE_ORDER):
        raise ValueError("The workbook does not contain one literature list for each legacy subtype")
    genes_by_subtype = {
        subtype: measured_sets[selected_by_subtype[subtype].set_id]
        for subtype in LEGACY_SUBTYPE_ORDER
    }
    union_genes = list(
        dict.fromkeys(gene for subtype in LEGACY_SUBTYPE_ORDER for gene in genes_by_subtype[subtype])
    )
    log_expression, _ = expression_matrix(adata, union_genes)
    scores, gene_stats, signature_stats = standardized_mean_z_scores(
        log_expression, union_genes, genes_by_subtype
    )
    metadata = adata.obs[["sample", "condition", "tissue", args.cluster_key]].copy()
    metadata[args.cluster_key] = metadata[args.cluster_key].astype(str).to_numpy()
    scores.index = metadata.index
    per_cell = pd.concat([metadata, scores], axis=1)
    per_cell.index.name = "cell_id"
    summary = summarize_standardized_scores(metadata, scores, args.cluster_key)
    calls = bootstrap_standardized_cluster_calls(
        metadata,
        scores,
        args.cluster_key,
        n_bootstrap=args.bootstrap,
        seed=args.seed,
        strata_key="sample",
    )
    selected_ids = {descriptor.set_id for descriptor in selected}
    coverage = coverage.loc[coverage["set_id"].isin(selected_ids)].copy()
    coverage["scoring_role"] = "primary unified legacy-list comparison"

    per_cell.to_csv(output_dir / "standardized_scores_per_cell.csv.gz", compression="gzip")
    summary.to_csv(output_dir / "standardized_score_summary.csv", index=False)
    calls.to_csv(output_dir / "standardized_cluster_calls.csv", index=False)
    gene_stats.to_csv(output_dir / "gene_standardization.csv", index=False)
    signature_stats.to_csv(output_dir / "signature_standardization.csv", index=False)
    coverage.to_csv(output_dir / "legacy_literature_gene_coverage.csv", index=False)
    plot_limits = plot_standardized_legacy_scores(
        adata, metadata, scores, summary, output_dir, args.cluster_key
    )

    task_summary = {
        "task": "standardized-legacy",
        "n_cells": int(adata.n_obs),
        "primary_source": "literature sheets used by the legacy PPT",
        "subtypes": list(LEGACY_SUBTYPE_ORDER),
        "score_method": (
            "gene-wise z-score across all cells; unweighted mean within each measured literature "
            "list; signature-wise z-score across all cells"
        ),
        "bootstrap_replicates": args.bootstrap,
        "call_rule": (
            "legacy-list enrichment only when the 95% sample-stratified cell-bootstrap CI is "
            "above zero and the 95% CI for winner-minus-the-strongest-bootstrap-competitor is "
            "above zero; positive but unseparated candidates are mixed, and non-positive "
            "candidates are unclassified"
        ),
        "inference_note": (
            "bootstrap intervals quantify cell-level stability only; one biological sample per "
            "tissue-condition prevents biological-replicate inference"
        ),
        "plot_limits": plot_limits,
        "h5ad_was_modified": False,
        "outputs": sorted(path.name for path in output_dir.iterdir()),
    }
    write_json(output_dir / "summary.json", task_summary)
    return task_summary


def marker_expression_summary(
    obs: pd.DataFrame,
    log_expression: np.ndarray,
    normalized_expression: np.ndarray,
    genes: Sequence[str],
    cluster_key: str,
) -> pd.DataFrame:
    grouping_specs = [
        ("global", ["condition"]),
        ("tissue", ["tissue", "condition"]),
        ("cluster", [cluster_key, "condition"]),
        ("tissue_cluster", ["tissue", cluster_key, "condition"]),
    ]
    rows = []
    for stratum_type, keys in grouping_specs:
        for key_values, positions in obs.groupby(keys, observed=True, sort=True).indices.items():
            if not isinstance(key_values, tuple):
                key_values = (key_values,)
            labels = dict(zip(keys, map(str, key_values), strict=True))
            log_subset = log_expression[positions]
            norm_subset = normalized_expression[positions]
            mean_log = log_subset.mean(axis=0)
            mean_norm = norm_subset.mean(axis=0)
            fraction = (log_subset > 0).mean(axis=0)
            for index, gene in enumerate(genes):
                rows.append(
                    {
                        "stratum_type": stratum_type,
                        "tissue": labels.get("tissue"),
                        "cluster": labels.get(cluster_key),
                        "condition": labels["condition"],
                        "gene": gene,
                        "n_cells": int(len(positions)),
                        "mean_log1p_expression": float(mean_log[index]),
                        "mean_normalized_expression": float(mean_norm[index]),
                        "fraction_expressing": float(fraction[index]),
                    }
                )
    return pd.DataFrame(rows)


def marker_condition_effects(summary: pd.DataFrame) -> pd.DataFrame:
    effects = condition_effects(
        summary,
        ["mean_log1p_expression", "mean_normalized_expression", "fraction_expressing"],
    )
    pseudocount = 1e-3
    effects["log2fc_mean_normalized_t2_vs_ctrl"] = np.log2(
        (effects["t2_mean_normalized_expression"] + pseudocount)
        / (effects["ctrl_mean_normalized_expression"] + pseudocount)
    )
    return effects


def plot_marker_results(summary: pd.DataFrame, effects: pd.DataFrame, output_dir: Path) -> None:
    dot = summary.loc[summary["stratum_type"].eq("tissue")].copy()
    dot["condition_tissue"] = dot["condition"] + " " + dot["tissue"]
    x_order = [
        f"{condition} {tissue}"
        for tissue in ("bone_marrow", "spleen", "thymus")
        for condition in ("Ctrl", "T2")
    ]
    gene_order = list(dict.fromkeys(dot["gene"]))
    dot["x"] = pd.Categorical(dot["condition_tissue"], categories=x_order, ordered=True).codes
    dot["y"] = pd.Categorical(dot["gene"], categories=gene_order[::-1], ordered=True).codes
    plt.figure(figsize=(10, 7))
    scatter = plt.scatter(
        dot["x"],
        dot["y"],
        s=20 + 500 * dot["fraction_expressing"],
        c=dot["mean_log1p_expression"],
        cmap="magma",
        edgecolor="0.25",
        linewidth=0.3,
    )
    plt.xticks(range(len(x_order)), x_order, rotation=40, ha="right")
    plt.yticks(range(len(gene_order)), gene_order[::-1])
    plt.colorbar(scatter, label="Mean log1p expression")
    plt.title("Selected iNKT biomarkers: color=mean, size=fraction expressing")
    plt.tight_layout()
    plt.savefig(output_dir / "marker_expression_condition_tissue_dotplot.png", dpi=180, bbox_inches="tight")
    plt.close()

    effect = effects.loc[effects["stratum_type"].isin(["global", "tissue"])].copy()
    effect["stratum"] = effect["tissue"].fillna("global")
    matrix = effect.pivot(
        index="gene", columns="stratum", values="log2fc_mean_normalized_t2_vs_ctrl"
    )
    ordered_columns = [column for column in ["global", "bone_marrow", "spleen", "thymus"] if column in matrix]
    matrix = matrix.reindex(gene_order)
    plt.figure(figsize=(7, 7))
    sns.heatmap(matrix[ordered_columns], cmap="vlag", center=0, annot=True, fmt=".2f")
    plt.title("Biomarker descriptive effect: log2(T2/Ctrl)")
    plt.tight_layout()
    plt.savefig(output_dir / "marker_t2_ctrl_effect_heatmap.png", dpi=180, bbox_inches="tight")
    plt.close()


def run_marker_validation(args: argparse.Namespace) -> dict:
    output_dir = args.out_dir / "marker_validation"
    output_dir.mkdir(parents=True, exist_ok=True)
    adata = sc.read_h5ad(args.h5ad)
    source_names = list(map(str, adata.raw.var_names if adata.raw is not None else adata.var_names))
    lookup = build_casefold_lookup(source_names)
    requested = [clean_gene_token(gene) for gene in args.markers]
    genes = [lookup[gene.upper()] for gene in requested if gene and gene.upper() in lookup]
    missing = [gene for gene in requested if gene and gene.upper() not in lookup]
    if not genes:
        raise ValueError("None of the requested marker genes are present")

    log_expression, normalized_expression = expression_matrix(adata, genes)
    obs = adata.obs[["sample", "condition", "tissue", args.cluster_key]].copy()
    summary = marker_expression_summary(
        obs, log_expression, normalized_expression, genes, args.cluster_key
    )
    effects = marker_condition_effects(summary)
    summary.to_csv(output_dir / "marker_expression_summary.csv", index=False)
    effects.to_csv(output_dir / "marker_t2_vs_ctrl_effects.csv", index=False)
    plot_marker_results(summary, effects, output_dir)

    headline = effects.loc[
        effects["stratum_type"].isin(["global", "tissue"]),
        [
            "gene",
            "stratum_type",
            "tissue",
            "n_ctrl",
            "n_t2",
            "log2fc_mean_normalized_t2_vs_ctrl",
            "delta_fraction_expressing_t2_minus_ctrl",
        ],
    ].to_dict(orient="records")
    task_summary = {
        "task": "marker-validation",
        "n_cells": int(adata.n_obs),
        "markers_requested": requested,
        "markers_measured": genes,
        "markers_missing": missing,
        "effects_are_descriptive_only": True,
        "reason_no_cell_level_p_values": (
            "The H5AD does not expose independent animal IDs; cells are not biological replicates."
        ),
        "global_and_tissue_effects": headline,
        "h5ad_was_modified": False,
        "outputs": sorted(path.name for path in output_dir.iterdir()),
    }
    write_json(output_dir / "summary.json", task_summary)
    return task_summary


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--task",
        choices=(
            "gene-sets",
            "signature-scores",
            "standardized-legacy",
            "marker-validation",
            "all",
        ),
        default="all",
        help="Independent task; each writes to its own subdirectory.",
    )
    parser.add_argument("--h5ad", type=Path, default=DEFAULT_H5AD)
    parser.add_argument("--gene-list", type=Path, default=DEFAULT_GENE_LIST)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT_DIR)
    parser.add_argument("--cluster-key", default=DEFAULT_CLUSTER_KEY)
    parser.add_argument("--markers", nargs="+", default=DEFAULT_MARKERS)
    parser.add_argument("--seed", type=int, default=0)
    parser.add_argument(
        "--bootstrap",
        type=int,
        default=10000,
        help="Cell-bootstrap replicates for standardized legacy cluster calls.",
    )
    parser.add_argument("--threads", type=int, default=4)
    parser.add_argument(
        "--validate-only",
        action="store_true",
        help="Backed-mode schema/signature check only; performs no long computation or result writes.",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    args.h5ad = args.h5ad.resolve()
    args.gene_list = args.gene_list.resolve()
    args.out_dir = args.out_dir.resolve()
    if args.threads < 1:
        raise ValueError("--threads must be >= 1")
    if args.bootstrap < 1:
        raise ValueError("--bootstrap must be >= 1")

    validation = validate_inputs(args)
    if args.validate_only:
        print(json.dumps(json_safe(validation), indent=2, sort_keys=True, allow_nan=False))
        return 0

    tasks = (
        ["gene-sets", "signature-scores", "standardized-legacy", "marker-validation"]
        if args.task == "all"
        else [args.task]
    )
    context = threadpool_limits(limits=args.threads) if threadpool_limits is not None else None
    if context is None:
        results = [globals()[f"run_{task.replace('-', '_')}"](args) for task in tasks]
    else:
        with context:
            results = [globals()[f"run_{task.replace('-', '_')}"](args) for task in tasks]

    manifest = {
        "script": str(Path(__file__).resolve()),
        "task_requested": args.task,
        "threads": args.threads,
        "input_h5ad": str(args.h5ad),
        "input_h5ad_size_bytes": args.h5ad.stat().st_size,
        "input_gene_list": str(args.gene_list),
        "input_gene_list_sha256": sha256_file(args.gene_list),
        "validation": validation,
        "results": results,
        "python": sys.version,
        "scanpy": sc.__version__,
        "anndata": ad.__version__,
    }
    # Task-specific manifests avoid concurrent writes when tasks run in parallel.
    manifest_path = args.out_dir / f"manifest_{args.task}.json"
    write_json(manifest_path, manifest)
    print(json.dumps(json_safe(manifest), indent=2, sort_keys=True, allow_nan=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
