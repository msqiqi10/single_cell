"""OOXML element preview for visual QA (not a Microsoft Office rendering)."""
from pathlib import Path
import io
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from PIL import Image,ImageDraw
OUT=Path(__file__).resolve().parents[1];BASE=OUT/'package';QA=OUT/'qa';QA.mkdir(exist_ok=True)
def rgb(font,default='#222222'):
 try:return '#'+str(font.color.rgb)
 except (AttributeError,TypeError):return default

def main():
 prs=Presentation(BASE/'decks/iNKT_integrated_58_editable.pptx');W=prs.slide_width/914400;H=prs.slide_height/914400;selected=[1,2,4,10,11,13,16,22,23,27,28,38,47];thumbs=[]
 for num in selected:
  s=prs.slides[num-1];fig=plt.figure(figsize=(W,H));ax=fig.add_axes([0,0,1,1]);ax.set_xlim(0,W);ax.set_ylim(H,0);ax.axis('off')
  for shape in s.shapes:
   x,y,w,h=[v/914400 for v in [shape.left,shape.top,shape.width,shape.height]]
   if shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
    im=Image.open(io.BytesIO(shape.image.blob));ax.imshow(im,extent=[x,x+w,y+h,y],aspect='auto')
   elif shape.has_text_frame:
    yy=y
    for para in shape.text_frame.paragraphs:
     if not para.runs:yy+=.18;continue
     r=para.runs[0];font=r.font;size=font.size.pt if font.size else 11
     ha={PP_ALIGN.LEFT:'left',PP_ALIGN.CENTER:'center',PP_ALIGN.RIGHT:'right'}.get(para.alignment,'left');xx=x if ha=='left' else x+w/2 if ha=='center' else x+w
     ax.text(xx,yy,para.text,fontsize=size,ha=ha,va='top',color=rgb(font),weight='bold' if font.bold else 'normal',fontfamily='DejaVu Sans')
     yy+=size*1.19/72
   elif shape.has_table:
    t=shape.table;yy=y
    for row in t.rows:
     xx=x;rh=row.height/914400
     for col,cell in zip(t.columns,row.cells):
      cw=col.width/914400
      try:fill='#'+str(cell.fill.fore_color.rgb)
      except (AttributeError,TypeError):fill='white'
      ax.add_patch(Rectangle((xx,yy),cw,rh,facecolor=fill,edgecolor='#ddd',lw=.5))
      font=cell.text_frame.paragraphs[0].runs[0].font if cell.text_frame.paragraphs[0].runs else None;size=font.size.pt if font and font.size else 11
      # Matplotlib preview retains explicit source line breaks; Office wraps inside cells.
      import textwrap
      maxchars=max(8,int((cw-.1)*72/(size*.52)));text='\n'.join('\n'.join(textwrap.wrap(line,maxchars)) for line in cell.text.split('\n'))
      ax.text(xx+.06,yy+rh/2,text,fontsize=size,ha='left',va='center',color=rgb(font) if font else '#222',weight='bold' if font and font.bold else 'normal')
      xx+=cw
     yy+=rh
  p=QA/f'pptx_elements_{num:02}.png';fig.savefig(p,dpi=100);plt.close(fig);im=Image.open(p);im.thumbnail((480,270));tile=Image.new('RGB',(490,295),'#ddd');tile.paste(im,(5,5));ImageDraw.Draw(tile).text((8,277),f'OOXML preview — integrated slide {num}',fill='black');thumbs.append(tile)
 sheet=Image.new('RGB',(1470,295*((len(thumbs)+2)//3)),'white')
 for i,im in enumerate(thumbs):sheet.paste(im,(i%3*490,i//3*295))
 sheet.save(QA/'editable_element_contact.jpg')
if __name__=='__main__':main()
