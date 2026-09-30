"""Assemble a navigable, source-grounded handoff package and verified ZIP."""
from pathlib import Path
import json,html,hashlib,shutil,zipfile,datetime,sys,tempfile
import pandas as pd
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
ROOT=Path(__file__).resolve().parents[3];OUT=Path(__file__).resolve().parents[1];BASE=OUT/'package'
def sha(p):
 h=hashlib.sha256()
 with Path(p).open('rb') as f:
  for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
 return h.hexdigest()
def main():
 scenes=json.loads((BASE/'slides.json').read_text());byid={s['id']:s for s in scenes};index=pd.read_csv(BASE/'decks/iNKT_integrated_58_editable_index.csv')
 # A real edit/save/reopen check, performed only on a temporary copy.
 with tempfile.TemporaryDirectory() as td:
  prs=Presentation(BASE/'decks/iNKT_previous_37_editable.pptx');tx=next(s for s in prs.slides[0].shapes if s.has_text_frame);tx.text='EDITABLE_TITLE_SMOKE_TEST';table=next(s.table for s in prs.slides[0].shapes if s.has_table);table.cell(0,0).text='EDITABLE_TABLE_SMOKE_TEST';im=next(s for s in prs.slides[1].shapes if s.shape_type==MSO_SHAPE_TYPE.PICTURE);newleft=im.left+91440;im.left=newleft;tmp=Path(td)/'edit_test.pptx';prs.save(tmp);r=Presentation(tmp)
  assert any(s.has_text_frame and s.text=='EDITABLE_TITLE_SMOKE_TEST' for s in r.slides[0].shapes)
  assert next(s.table for s in r.slides[0].shapes if s.has_table).cell(0,0).text=='EDITABLE_TABLE_SMOKE_TEST'
  assert next(s for s in r.slides[1].shapes if s.shape_type==MSO_SHAPE_TYPE.PICTURE).left==newleft
 (BASE/'notes/editability_smoke_test.json').write_text(json.dumps({'status':'passed','title_edit_save_reopen':True,'table_cell_edit_save_reopen':True,'independent_panel_move_save_reopen':True,'delivered_files_modified_by_test':False},indent=2))
 # Copy the numeric files actually read by the previous builder; do not duplicate large H5ADs.
 sources=json.loads((BASE/'notes/previous_source_hashes.json').read_text());numeric_map=[]
 for p,h in sources.items():
  if p.endswith(('.csv','.csv.gz','.json')):
   target=BASE/'data/previous_source_tables'/p;target.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(ROOT/p,target);numeric_map.append({'original_path':p,'package_path':str(target.relative_to(BASE)),'sha256':h})
 (BASE/'notes/numeric_source_mapping.json').write_text(json.dumps(numeric_map,indent=2))
 # Useful extraction/build scripts for audit; standalone rebuilding uses code/rebuild_decks.py.
 for p in (OUT/'code').glob('*.py'):shutil.copy2(p,BASE/'code'/p.name)
 oldprefix={5:'marker_iNKT1',6:'marker_iNKT2',7:'marker_iNKT17',8:'Fig3C_bone_marrow',9:'Fig3C_spleen',10:'Fig3C_thymus',11:'DEG_global',12:'DEG_tissue__bone_marrow',13:'DEG_tissue__spleen',14:'DEG_tissue__thymus',15:'DEG_cluster_tissue__C0__bone_marrow',16:'DEG_cluster_tissue__C6__thymus',17:'DEG_cluster_tissue__C5-1__spleen',18:'DEG_cluster_tissue__C5-2__bone_marrow',19:'legacy_pathways_c1_thymus',20:'legacy_pathways_c5_bone_marrow',21:'legacy_pathways_c6_spleen',22:'Fig3D_',23:'Fig3E_reference_cluster_tissue__C0',24:'Fig3E_reference_cluster_tissue__C5-1',25:'Fig3E_reference_cluster_tissue__C6',26:'Fig3E_top10_cluster_tissue__C0',27:'Fig3E_top10_cluster_tissue__C5-1',28:'Fig3E_top10_cluster_tissue__C6',29:'Fig3F_'}
 newfiles={1:['cell_metadata_scores.csv.gz'],2:['GO_contrast_coverage.csv'],3:['module_gene_coverage.csv','cytotoxicity_sample_summary.csv'],4:['marker_dotplot_values.csv','cytotoxicity_contrasts.csv'],5:['marker_dotplot_values.csv','cytotoxicity_contrasts.csv'],6:['marker_dotplot_values.csv','cytotoxicity_contrasts.csv'],7:['GO_ORA_significant.csv','GO_ORA_all.csv.gz'],8:['GO_ORA_significant.csv','GO_ORA_all.csv.gz'],9:['GO_ORA_significant.csv','GO_ORA_all.csv.gz'],10:['GO_ORA_significant.csv','GO_ORA_all.csv.gz'],11:['GO_ORA_significant.csv','GO_ORA_all.csv.gz'],12:['GO_ORA_significant.csv','GO_ORA_all.csv.gz'],13:['trajectory_binned_curves.csv','trajectory_root_audit.csv','trajectory_correlations.csv'],14:['trajectory_binned_curves.csv','trajectory_root_audit.csv','trajectory_correlations.csv'],15:['trajectory_binned_curves.csv','trajectory_root_audit.csv','trajectory_correlations.csv'],16:['cytotoxicity_contrasts.csv','GO_cytotoxicity_targeted_audit.csv','focus_stable_cluster_crosswalk.csv'],17:['cytotoxic_gene_CC_membership.csv'],18:['cytotoxicity_QC_sensitivity.csv'],19:['cytotoxicity_integrated_summary.csv']}
 page_rows=[];cards=[];notes=['# 逐页资料与讲稿索引','', '前28页为主报告，后30页为保留全部源页的附录。图内坐标/图例属于独立SVG/PNG图形；标题、说明、表格为原生可编辑元素。','']
 for row in index.itertuples():
  s=byid[row.source_page];data=[]
  if s['series']=='old' and s['original_page'] in oldprefix:data=[str(p.relative_to(BASE)) for p in (BASE/'data/previous_display_tables').glob(oldprefix[s['original_page']]+'*.csv')]
  elif s['series']=='new':data=['data/current_analysis_tables/'+p for p in newfiles.get(s['original_page'],[]) if (BASE/'data/current_analysis_tables'/p).exists()]
  if s['id']=='custom_02':data=['data/current_analysis_tables/focus_cytotoxic_gene_DE.csv','data/current_analysis_tables/cytotoxicity_contrasts.csv','data/current_analysis_tables/GO_ORA_significant.csv']
  data.extend(str(p.relative_to(BASE)) for p in (BASE/'data').glob(s['id']+'_editable_table_*.csv'))
  panels=[e for e in s['elements'] if e['type']=='image'];native_tables=sum(e['type']=='table' for e in s['elements'])
  page_rows.append({'integrated_slide':row.slide,'section':row.section,'source_page':s['id'],'title':s['title'],'reference':s['reference'],'source_panels':';'.join(e['png'] for e in panels),'numeric_tables':';'.join(data),'native_tables':native_tables,'notes':s.get('notes_zh','')})
  links=[]
  for i,e in enumerate(panels,1):
   svg=f' · <a href="{html.escape(e["svg"])}">SVG</a>' if e.get('svg') else ''
   links.append(f'<li>Panel {i}: <a href="{html.escape(e["png"])}">PNG / original image</a>{svg}</li>')
  datalinks=''.join(f'<li><a href="{html.escape(p)}">{html.escape(Path(p).name)}</a></li>' for p in data)
  cards.append(f'<section id="slide-{row.slide}"><h2>{row.slide:02}. {html.escape(s["title"])}</h2><p>{row.section} · source {s["id"]} · {len(panels)} independent panels · {native_tables} native tables</p><img loading="lazy" src="{s["preview"]}" alt="Source layout preview"><p>{html.escape(s["reference"])}</p><p>{html.escape(s.get("notes_zh",""))}</p><details><summary>Copy-ready source panels</summary><ul>{"".join(links) or "<li>Text/table slide: edit directly in PPTX.</li>"}</ul></details><details><summary>Numeric source tables</summary><ul>{datalinks or "<li>See notes/numeric_source_mapping.json and original source reference.</li>"}</ul></details></section>')
  notes.extend([f'## {row.slide:02}. {s["title"]}',f'源页：{s["id"]}；{row.section}',s.get('notes_zh',''),s['reference'],s['note'],'图元素：'+', '.join(e['png'] for e in panels),'数值表：'+', '.join(data),''])
 pd.DataFrame(page_rows).to_csv(BASE/'notes/slide_materials_index.csv',index=False)
 (BASE/'notes/presenter_guide.zh.md').write_text('\n'.join(notes))
 header='''<!doctype html><html lang="zh"><meta charset="utf-8"><title>iNKT presentation materials</title><style>body{font:16px/1.6 system-ui,sans-serif;background:#f4f6f8;color:#1d2939;max-width:1250px;margin:30px auto;padding:0 24px}a{color:#12658c}header,section{background:white;padding:24px;border-radius:10px;margin:20px 0}section img{width:72%;height:auto;border:1px solid #eee}h1{font-size:30px}h2{font-size:22px}summary{cursor:pointer;font-weight:bold}li{margin:4px 0}.nav{display:flex;gap:20px;flex-wrap:wrap}</style><header><h1>iNKT 论文式汇报资料包 · 2026-09-20</h1><p>独立源图面板 + 原生可编辑标题、说明和表格。这里的预览仅供找图，不会作为整页图片插入PPT。</p><div class="nav"><a href="decks/iNKT_main_talk_28_editable.pptx">28页主报告 PPTX</a><a href="decks/iNKT_main_talk_28_editable.pdf">主报告 PDF</a><a href="decks/iNKT_integrated_58_editable.pptx">58页完整版 PPTX</a><a href="decks/iNKT_integrated_58_editable.pdf">完整版 PDF</a><a href="decks/iNKT_previous_37_editable.pptx">上次37页可编辑版</a><a href="decks/iNKT_cytotoxicity_GO_19_editable.pptx">新增19页可编辑版</a></div><p>前28页为主线，其余为附录。所有图表保留原统计量、比较范围和解释限制。源页预览保留原页码，整合文件另用连续页码。</p><p>编辑PPT后请在PowerPoint中导出PDF，使修改同步。包内PDF来自同一绘图源布局，未宣称经过Office排版渲染。</p></header>'''
 (BASE/'index.html').write_text(header+'\n'.join(cards)+'</html>')
 readme='''# iNKT 可编辑汇报资料包｜2026-09-20

建议先打开 **index.html**：按整合顺序浏览每页，直接找到该页的独立源图、SVG、PNG、数值表和论文出处。

## 四个现成版本

| 版本 | 用途 | PPTX | PDF |
|---|---|---|---|
| 28页主报告 | 按论文逻辑组织的主线 | [可编辑PPT](decks/iNKT_main_talk_28_editable.pptx) | [PDF](decks/iNKT_main_talk_28_editable.pdf) |
| 58页完整版 | 主报告28页 + 详细附录30页；原37页与新增19页全部保留，加2页整合说明 | [可编辑PPT](decks/iNKT_integrated_58_editable.pptx) | [PDF](decks/iNKT_integrated_58_editable.pdf) |
| 上次37页转换版 | 9月15日综合PDF，逐页还原为可编辑元素 | [可编辑PPT](decks/iNKT_previous_37_editable.pptx) | [PDF](decks/iNKT_previous_37_editable.pdf) |
| 新增19页转换版 | 9月19日细胞毒性与GO分析，独立使用 | [可编辑PPT](decks/iNKT_cytotoxicity_GO_19_editable.pptx) | [PDF](decks/iNKT_cytotoxicity_GO_19_editable.pdf) |

## Supervisor 可以怎样编辑

- **标题、正文、图题、结论和说明**：原生PowerPoint文本框，直接点击修改。
- **数值表/方法表**：原生PowerPoint表格，逐格修改；CSV备份在data/。
- **UMAP、dotplot、GO气泡图、曲线等**：每个图形面板独立插入，可单独移动、缩放、删除或替换。程序图提供SVG矢量源和PNG回退；已有原始图片直接使用源文件。
- **没有使用整张幻灯片截图作为PPT内容。** assets/page_previews仅用于找图与质量检查，不嵌入PPT。
- 图内散点、坐标和图例仍属于图形对象，不是原生Excel数据图表；如果要改统计数值或重新绘图，请使用对应数值表与绘图代码。支持SVG的PowerPoint可尝试“转换为形状”；不承诺每个点都是原生可编辑数据点。
- 编辑PPT后，在PowerPoint使用“导出PDF”同步修改。包内PDF由同一源图布局生成，不是Office渲染导出；实际Office字体替换、换行可能略有差异。

## 主报告组织逻辑

1. 数据与QC → 细胞状态/marker → 组织分层组成（原iNKT展示和Blood Figure 3的逻辑）。
2. 差异表达 → 杀伤模块 → 群特异与条件相关GO（PAGER Figure 5与Borra Figure 4c）。
3. GO细胞组分 → 局部拟时序 → QC敏感性 → 证据边界与结论。
4. 附录保留所有详细DE、旧通路核查、GSEA、Venn、整合敏感性和既有候选发现。

## 资料来源与实际结果

- 旧部分：9月15日37页综合报告，复用其原绘图代码和数值。
- 新部分：9月19日19页细胞毒性/GO结果。未重新做生物分析，未改变统计阈值或挑选新结论。
- 细胞杀伤核心信号：骨髓C4升高、脾脏C3小幅降低；扩展E1代表模块与杀伤相关GO未显著。只支持局部转录状态候选，不能写成已证明杀伤功能改变。
- 当前数据每组织×条件只有1个sample标签，没有独立动物重复；拟时序不是velocity或谱系验证。

## 文件导航

- `assets/panels/`：可直接插入PPT的142个独立图元素及其SVG版本。
- `data/`：展示用数值表、GO完整结果、源DE表与原生表格CSV；不重复打包大型表达矩阵。
- `notes/slide_materials_index.csv`：整合页码—源页—论文—图文件—数据表索引。
- `notes/panel_source_index.csv`：每个图元素的来源与类型。
- `notes/presenter_guide.zh.md`：逐页中文讲稿提示和出处；PPT备注页也有说明。
- `references/`：PAGER、Borra与Blood三篇参考论文。
- `slides.json`：独立元素布局与文本/表格数据。
- `code/rebuild_decks.py`：从本包内素材独立重建PPTX和配套PDF。
- `code/source_builder_*.py`：原始绘图代码快照，作为来源审计；完整重绘依赖原项目数据和环境，不是独立分析入口。
- `notes/verification.json`：文本/表格/图片来源与文件结构核对。
- `notes/editability_smoke_test.json`：修改标题、表格单元格、移动图形后保存并重新打开的测试结果。

## 重建

准备含python-pptx、pypdf、matplotlib的Python环境，在解压后的package目录运行：

```bash
python code/rebuild_decks.py
```

本项目使用现有`.venv`，长任务由tmux运行，GPU全部禁用。旧数据和旧报告不覆盖。
'''
 (BASE/'README.zh.md').write_text(readme)
 (BASE/'notes/reference_presentation_mapping.md').write_text('''# 参考论文与图形组织

| 来源 | 图形/问题 | 本资料包中的实现 |
|---|---|---|
| 原iNKT展示 | QC、marker、DEG+表达dotplot、群与组织 | 保留原37页中的数据与口径；文字表格转原生对象，图按panel导出 |
| Blood Advances Figure 3C | 两种细胞比例分母 | 按组织放置成对堆积柱状图，保留实际细胞数 |
| Blood Figure 3D/E/F | 火山、上调基因富集、四集合交集 | 保留log2FC/P坐标、−ln(P)评分与集合阈值；不把GSEA NES混作ORA评分 |
| PAGER-scFGA Figure 5 | 功能模块、GO细胞组分、拟时序曲线 | 杀伤核心与E1代表基因、CC注释矩阵、组织内DPT曲线；不冒充作者完整PPI网络 |
| Borra et al. §3.5/Figure 4c | 亚群功能富集矩阵 | 组织内cluster marker GO与条件上下调GO分开展示 |

PDF文件与出处保留在references/。iNKT数据不同于参考论文NK群体；没有复制其生物结论，也未将文档中的指令视为用户授权。
''')
 # Include dependency versions used for this build, without machine-specific credentials/URLs.
 from importlib.metadata import version
 (BASE/'code/requirements-presentation.txt').write_text('\n'.join(f'{x}=={version(x)}' for x in ['python-pptx','pypdf','matplotlib','Pillow','pandas','lxml'])+'\n')
 rootreadme='''# 2026-09-20｜论文式资料包与可编辑PPT

[打开资料包说明](package/README.zh.md) · [浏览源图与逐页资料](package/index.html)

- [28页主报告 PPTX](package/decks/iNKT_main_talk_28_editable.pptx) · [PDF](package/decks/iNKT_main_talk_28_editable.pdf)
- [58页完整整合 PPTX](package/decks/iNKT_integrated_58_editable.pptx) · [PDF](package/decks/iNKT_integrated_58_editable.pdf)
- [上次37页可编辑转换版](package/decks/iNKT_previous_37_editable.pptx)
- [新增19页可编辑版](package/decks/iNKT_cytotoxicity_GO_19_editable.pptx)
- [完整ZIP资料包](iNKT_editable_materials_20260920.zip)

旧版按9月15日37页综合报告处理，同时提供9月19日19页报告的可编辑版本。原统计结果不变；496个原生文本框、17张原生表格、142个独立图形元素，不使用整页截图。图内坐标与点属于独立SVG/PNG图形，修改数值可用配套数据重绘。

伴随PDF来自同一源图布局。Supervisor编辑PPT后，请由PowerPoint导出PDF以同步修改。包内提供每页源图、数据索引、讲稿、参考论文、布局与重建代码。
'''
 (OUT/'README.md').write_text(rootreadme)
 # Keep the dates navigation explicit about real files vs historical links.
 p=ROOT/'iNKT_by_date/README.md';s=p.read_text()
 if '2026-09-20/README.md' not in s:
  s=s.replace('## 从这里开始\n','## 从这里开始\n\n- [最新可编辑PPT与论文式资料包：2026-09-20](2026-09-20/README.md)\n');s=s.replace('| [shared](shared/README.md)','| [2026-09-20](2026-09-20/README.md) | 可编辑PPT：旧37页、新19页、整合58页与主报告28页；源图资料包 | 展示整合，真实文件 |\n| [shared](shared/README.md)');s=s.replace('2026-09-19目录直接保存本轮代码与结果','2026-09-19和2026-09-20目录直接保存代码与结果');p.write_text(s)
 # Validate local HTML href and image targets before packaging.
 import re
 text=(BASE/'index.html').read_text()
 links=re.findall(r'(?:href|src)="([^"]+)"',text)
 for link in links:assert (BASE/link).exists(),link
 manifest={str(p.relative_to(BASE)):sha(p) for p in BASE.rglob('*') if p.is_file() and '__pycache__' not in str(p) and p.name!='file_manifest.json'}
 (BASE/'file_manifest.json').write_text(json.dumps(manifest,indent=2))
 zip_path=OUT/'iNKT_editable_materials_20260920.zip'
 with zipfile.ZipFile(zip_path,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=4) as z:
  for p in BASE.rglob('*'):
   if p.is_file() and '__pycache__' not in str(p):z.write(p,'iNKT_editable_materials_20260920/'+str(p.relative_to(BASE)))
 with zipfile.ZipFile(zip_path) as z:
  assert z.testzip() is None
  for path,h in manifest.items():assert hashlib.sha256(z.read('iNKT_editable_materials_20260920/'+path)).hexdigest()==h,path
 result={'status':'complete','created_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'archive':zip_path.name,'archive_sha256':sha(zip_path),'archive_bytes':zip_path.stat().st_size,'files_hashed':len(manifest),'HTML_local_links_checked':len(links),'archive_integrity':'CRC and every-file SHA256 passed','native_edit_save_reopen':'passed','no_full_slide_pictures':True}
 (OUT/'delivery_verification.json').write_text(json.dumps(result,indent=2));print(json.dumps(result,indent=2),flush=True)
if __name__=='__main__':main()
